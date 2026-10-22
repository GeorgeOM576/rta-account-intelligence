"""Real Streamlit page projections share live SQLite approval state; no paid calls."""
from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from rta.models import REVIEW_FILTERS
from tests.test_approved_translation_integration import fixture, approved
from tests.test_app import navigate, by_label
from tests.test_review_workflow import app_for
from tests.test_translation import Offline
from tests.translation_workspace_helpers import seed_google
from tests.workflow_display_helpers import csv_columns, rendered_table, visible_and_backend_row

PAGES = ('Run Reconciliation', 'AI Research', 'Review & Approvals', 'Export Results')
COLUMNS = ['Master Identifier', 'Master Approved English Translation', 'Master Translation Status',
           'Matched Source Approved English Translation', 'Matched Source Translation Status', 'Match Basis',
           'Master Original Name', 'Matched Source Original Name']


@pytest.fixture(autouse=True)
def block_live_services(monkeypatch):
    from google.cloud import translate_v3
    def forbidden(*args, **kwargs):
        raise AssertionError('Live paid services forbidden')
    monkeypatch.setattr(translate_v3, 'TranslationServiceClient', forbidden)
    monkeypatch.setattr('requests.Session.post', forbidden)


def setup(tmp_path, monkeypatch):
    case, engine = fixture(tmp_path, [{'Name': 'Les Jardins de Provence', 'Segment': 'Verify'}],
                           [{'Name': 'Jardines de Provenza', 'Segment': 'SMB'}])
    approved(engine, 'master', 'The Gardens of Provence')
    approved(engine, 'source', 'The Gardens of Provence')
    app = app_for(case, monkeypatch)
    navigate(app, PAGES[0])
    by_label(app.button, 'Run automatic local matching').click().run()
    assert not app.exception
    return case, engine, app


def row(app, page):
    navigate(app, page)
    assert not app.exception
    if page != 'Export Results':
        return visible_and_backend_row(app)
    table = rendered_table(app)
    assert list(table.columns[:8]) == COLUMNS  # Export deliberately keeps its full schema.
    assert len(table) == 1
    return table.iloc[0]


def test_all_four_pages_and_download_share_current_values(tmp_path, monkeypatch):
    case, engine, app = setup(tmp_path, monkeypatch)
    import rta.export
    original, downloads = rta.export.workbook, []
    def capture(*args, **kwargs):
        data = original(*args, **kwargs);downloads.append(data)
        return data
    monkeypatch.setattr(rta.export, 'workbook', capture)
    first = row(app, PAGES[0])
    for page in PAGES:
        current = row(app, page)
        assert current[COLUMNS].to_dict() == first[COLUMNS].to_dict()
        assert current['Master Approved English Translation'] == 'The Gardens of Provence'
        assert current['Matched Source Approved English Translation'] == 'The Gardens of Provence'
        assert current['Master Translation Status'] == current['Matched Source Translation Status'] == 'Approved Translation'
        assert current['Master Original Name'] == 'Les Jardins de Provence'
        assert current['Matched Source Original Name'] == 'Jardines de Provenza'
        assert current['Match Basis'] == 'Master Translation ↔ Source Translation'
        assert current['Proposed Value'] == 'SMB' and current['Final Value'] == ''
        assert 'Review Status' in current and 'Evidence URL 1' in current
    book = load_workbook(BytesIO(downloads[-1]))
    sheet = book['Reconciliation Results']
    exported = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert {c: exported[c] for c in COLUMNS} == first[COLUMNS].to_dict()
    assert exported['Final Value'] is None
    assert {'Audit Trail', 'Source Conflicts', 'Master Translations', 'Source Translations'} <= set(book.sheetnames)
    assert sheet.freeze_panes == 'A2' and sheet.auto_filter.ref
    assert case[0].load(case[1])['master'][0]['Name'] == 'Les Jardins de Provence'


def test_fuzzy_suggestion_is_not_displayed_as_matched_source(tmp_path, monkeypatch):
    case, engine = fixture(tmp_path, [{'Name': 'Bluebird Studio', 'Segment': 'Verify'}],
                           [{'Name': 'Jardines de Provenza', 'Segment': 'SMB'}])
    approved(engine, 'source', 'The Gardens of Provence')
    app = app_for(case, monkeypatch)
    navigate(app, PAGES[0])
    by_label(app.button, 'Run automatic local matching').click().run()
    for page in PAGES:
        current = row(app, page)
        assert current['Suggested Account 1'] == 'Jardines de Provenza'
        assert current['Match Status'] == 'Fuzzy candidates'
        for column in ('Matched Source Original Name', 'Matched Source Approved English Translation',
                       'Matched Source Translation Status', 'Matched Account Name', 'Source Original Name'):
            assert current[column] == ''
        assert current['Final Value'] == ''


@pytest.mark.parametrize('status', ['Translation Needed', 'Translated — Review', 'Rejected Translation'])
def test_pending_rejected_master_and_matched_source_statuses(tmp_path, monkeypatch, status):
    # Original-name identity matches independently of any translation draft.
    case, engine = fixture(tmp_path, [{'Name': 'Banco de México', 'Segment': 'Verify'}],
                           [{'Name': 'Banco de México', 'Segment': 'SMB'}])
    seed_google(engine, {'Banco de México': 'es'})
    if status != 'Translation Needed':
        engine.process(['master', 'source'], Offline())
    if status == 'Rejected Translation':
        for entry in engine.records():
            engine.reject(entry, 'Reviewer', True)
    app = app_for(case, monkeypatch)
    navigate(app, PAGES[0])
    by_label(app.button, 'Run automatic local matching').click().run()
    for page in (PAGES[0], PAGES[2], PAGES[3]):
        current = row(app, page)
        assert current['Master Translation Status'] == current['Matched Source Translation Status'] == status
        assert current['Master Approved English Translation'] == current['Matched Source Approved English Translation'] == ''
        assert current['Match Basis'] == 'Master Original ↔ Source Original'
        assert current['Final Value'] == ''


def test_current_correction_revocation_and_stale_notice_across_pages(tmp_path, monkeypatch):
    case, engine, app = setup(tmp_path, monkeypatch)
    approved(engine, 'master', 'The Gardens of Provence International')
    for page in PAGES:
        current = row(app, page)
        assert current['Master Approved English Translation'] == 'The Gardens of Provence International'
        assert any('results are stale' in w.value for w in app.warning)
        assert current['Final Value'] == ''
    entry = next(r for r in engine.records() if r['dataset'] == 'master')
    engine.mark_not_required(entry, 'Reviewer', True)
    for page in PAGES:
        current = row(app, page)
        assert current['Master Approved English Translation'] == ''
        assert current['Master Translation Status'] == 'Not Required'
        assert current['Matched Source Approved English Translation'] == 'The Gardens of Provence'
        assert any('results are stale' in w.value for w in app.warning)


def test_workspace_isolation_even_with_copied_saved_context(tmp_path, monkeypatch):
    case, engine, app = setup(tmp_path, monkeypatch)
    store, project, _ = case
    other, token = store.create()
    payload = deepcopy(store.load(project))
    store.save(other, payload)  # intentionally copy stale approved display context, not approvals
    app = app_for((store, other, token), monkeypatch)
    for page in PAGES:
        current = row(app, page)
        assert current['Master Approved English Translation'] == current['Matched Source Approved English Translation'] == ''
        assert current['Master Translation Status'] == current['Matched Source Translation Status'] == ''


def test_review_manual_decision_and_six_filters_preserved(tmp_path, monkeypatch):
    case, engine, app = setup(tmp_path, monkeypatch)
    current = row(app, 'Review & Approvals')
    assert current['Final Value'] == ''
    assert by_label(app.selectbox, 'Review status').options == REVIEW_FILTERS
    by_label(app.selectbox, 'Decision').set_value('Manual').run()
    by_label(app.text_input, 'Reviewer name').set_value('Account reviewer')
    by_label(app.text_input, 'Custom verified value').set_value('Healthcare')
    by_label(app.checkbox, 'I have verified the evidence and the field value for this specific account.').check()
    by_label(app.button, 'Save review decision').click().run()
    assert not app.exception
    current = app.dataframe[-1].value.iloc[0]
    assert current['Final Value'] == 'Healthcare' and current['Review Status'] == 'Manually Verified'
    assert current['Master Approved English Translation'] == 'The Gardens of Provence'
    assert len(case[0].audit(case[1])) == 3  # two translation decisions plus account decision


def test_saved_research_staleness_survives_matching_rerun(tmp_path, monkeypatch):
    case, engine, app = setup(tmp_path, monkeypatch)
    payload = case[0].load(case[1])
    saved = payload['results'][0]
    saved['research_snapshot'] = {'approved_name_context': deepcopy(saved['account_context'])}
    case[0].save(case[1], payload)
    approved(engine, 'master', 'The Gardens of Provence International')
    navigate(app, PAGES[0])
    by_label(app.button, 'Run automatic local matching').click().run()
    for page in PAGES[1:]:
        navigate(app, page)
        assert any('Saved AI research used different' in w.value for w in app.warning)
        assert not any('Reconciliation results are stale' in w.value for w in app.warning)


def test_live_display_refresh_preserves_atomic_source_approval(tmp_path, monkeypatch):
    case, engine, app = setup(tmp_path, monkeypatch)
    # Refreshing a display-only status must not invalidate a genuine saved proposal version.
    approved(engine, 'master', 'The Gardens of Provence Updated')
    navigate(app, 'Review & Approvals')
    assert not app.exception
    by_label(app.radio, 'Top 3 Suggested Source Matches').set_value(0).run()
    by_label(app.text_input, 'Reviewer name').set_value('Account reviewer')
    by_label(app.checkbox, 'I have verified the evidence and the field value for this specific account.').check()
    by_label(app.button, 'Save review decision').click().run()
    assert not app.exception
    saved = case[0].load(case[1])['results'][0]
    assert saved['decision'] == 'Approve' and saved['final'] == 'SMB'
    current = row(app, 'Review & Approvals')
    assert current['Master Approved English Translation'] == 'The Gardens of Provence Updated'
    assert current['Matched Source Original Name'] == 'Jardines de Provenza'


def test_reviewer_selected_source_uses_its_own_translation(tmp_path, monkeypatch):
    from rta.models import Result
    from rta.review import suggestion_version
    case, engine = fixture(tmp_path, [{'Name': 'Les Jardins de Provence', 'Segment': 'Verify'}],
        [{'Name': 'Jardines de Provenza', 'Segment': 'SMB'},
         {'Name': 'Ateliers de la Lune', 'Segment': 'Enterprise'}])
    approved(engine, 'master', 'The Gardens of Provence')
    approved(engine, 'source', 'The Gardens of Provence', 0)
    approved(engine, 'source', 'Moon Workshops', 1)
    app = app_for(case, monkeypatch)
    navigate(app, PAGES[0])
    by_label(app.button, 'Run automatic local matching').click().run()
    payload = case[0].load(case[1])
    saved = payload['results'][0]
    saved['candidates'].append({'source_row': 1, 'account': 'Ateliers de la Lune', 'value': 'Enterprise', 'similarity': 30})
    case[0].save(case[1], payload)
    current = Result(**saved)
    case[0].review(case[1], 0, 'Approve', 'Reviewer', 'Independently verified selected Source account',
                   selected_source_row=1, selection_version=suggestion_version(current, payload), confirmed=True)
    for page in PAGES:
        current = row(app, page)
        assert current['Matched Source Original Name'] == 'Ateliers de la Lune'
        assert current['Matched Source Approved English Translation'] == 'Moon Workshops'
        assert current['Matched Source Translation Status'] == 'Approved Translation'
        assert current['Final Value'] == 'Enterprise'


def test_excluded_status_stays_current_without_reactivating_export_fields(tmp_path, monkeypatch):
    case, engine = fixture(tmp_path, [{'Name': 'Banco de México', 'Segment': 'Verify'}],
                           [{'Name': 'Banco de México', 'Segment': 'SMB'}])
    payload = case[0].load(case[1])
    payload['language_detection_settings'] = {'field_types': {
        'master': {'Name': 'classification'}, 'source': {'Name': 'classification'}}}
    case[0].save(case[1], payload)
    engine.detect({'master': ['Name'], 'source': ['Name']})
    assert all(entry['status'] == 'Excluded' for entry in engine.records())
    app = app_for(case, monkeypatch)
    navigate(app, PAGES[0])
    by_label(app.button, 'Run automatic local matching').click().run()
    import rta.export
    original, downloads = rta.export.workbook, []
    def capture(*args, **kwargs):
        assert kwargs['translations'] == []
        assert len(kwargs['current_translations']) == 2
        data = original(*args, **kwargs);downloads.append(data)
        return data
    monkeypatch.setattr(rta.export, 'workbook', capture)
    for page in (PAGES[0], PAGES[2], PAGES[3]):
        current = row(app, page)
        assert current['Master Translation Status'] == current['Matched Source Translation Status'] == 'Excluded'
        assert current['Master Approved English Translation'] == current['Matched Source Approved English Translation'] == ''
    book = load_workbook(BytesIO(downloads[-1]))
    sheet = book['Reconciliation Results']
    exported = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert exported['Master Translation Status'] == exported['Matched Source Translation Status'] == 'Excluded'
    assert book['Translation Details'].max_row == 1


def test_research_action_immediately_displays_new_supported_source(tmp_path, monkeypatch):
    case, engine = fixture(tmp_path, [{'Name': 'Les Jardins de Provence', 'Segment': 'Verify'}],
        [{'Name': 'Jardines de Provenza', 'Segment': 'SMB'},
         {'Name': 'Ateliers de la Lune', 'Segment': 'Enterprise'}])
    approved(engine, 'master', 'The Gardens of Provence')
    approved(engine, 'source', 'The Gardens of Provence', 0)
    approved(engine, 'source', 'Moon Workshops', 1)
    monkeypatch.setenv('TAVILY_API_KEY', 'mock-only')
    monkeypatch.setenv('OPENAI_API_KEY', 'mock-only')
    seen = []
    def finding(self, result, *args, **kwargs):
        seen.append(deepcopy(result.account_context))
        result.recommendation_status = 'Proposed'
        result.recommended_source_account = 'Ateliers de la Lune'
        result.suggested_matches = [{'source_row': 1, 'account': 'Ateliers de la Lune',
                                     'value': 'Enterprise', 'mapping_supported': True}]
        result.proposed = 'Enterprise'
        return result
    monkeypatch.setattr('rta.research.Researcher.investigate', finding)
    app = app_for(case, monkeypatch)
    navigate(app, PAGES[0])
    by_label(app.button, 'Run automatic local matching').click().run()
    navigate(app, 'AI Research')
    assert not seen
    by_label(app.checkbox, 'I authorize external research for this batch and the associated API charges.').check().run()
    by_label(app.button, 'Start / resume paid research').click().run()
    assert not app.exception and len(seen) == 1
    current = visible_and_backend_row(app)
    assert current['Matched Source Original Name'] == 'Ateliers de la Lune'
    assert current['Matched Source Approved English Translation'] == 'Moon Workshops'
    assert current['Proposed Value'] == 'Enterprise' and current['Final Value'] == ''
    assert seen[0]['master']['approved_english'] == 'The Gardens of Provence'


def test_reviewed_display_reuses_dataset_fingerprint_and_rejects_stale_source(tmp_path, monkeypatch):
    import rta.approved_names as module
    from rta.models import Mapping, Result
    from rta.storage import digest
    mapping = Mapping('Name', 'Name', 'Segment', 'Segment')
    rows = [{'Name': f'Company {i}', 'Segment': 'SMB'} for i in range(100)]
    version = digest({'source': rows, 'mapping': vars(mapping)})
    original, full_source_hashes = module.digest, []
    def counted(value):
        if isinstance(value, dict) and 'source' in value:
            full_source_hashes.append(1)
        return original(value)
    monkeypatch.setattr(module, 'digest', counted)
    names = module.ApprovedNames({'master': rows, 'source': rows}, mapping)
    results = [Result(i, r['Name'], 'Verify', selected_source={
        'source_row': i, 'account': r['Name'], 'dataset': version}) for i, r in enumerate(rows)]
    names.attach(results)
    assert len(full_source_hashes) == 1  # one dataset hash, never one per reviewed row
    assert all(r.account_context['matched_source']['original_name'] == f'Company {i}' for i, r in enumerate(results))
    changed = deepcopy(rows)
    changed[0]['Name'] = 'New uploaded company'
    module.ApprovedNames({'master': rows, 'source': changed}, mapping).attach(results)
    assert all(not r.account_context['matched_source'] for r in results)
    assert len(full_source_hashes) == 2
