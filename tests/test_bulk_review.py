from copy import deepcopy
from io import BytesIO
import sqlite3

import pytest
from openpyxl import load_workbook

from rta.export import workbook, summary
from rta.models import Result
from rta.storage import Store, digest
from tests.test_app import by_label, navigate
from tests.test_review_workflow import workspace, record, app_for
from tests.test_recommendations import setup, research, M


def test_bulk_approve_persistence_audit_and_export(workspace):
    store, project, _ = workspace
    before = store.load(project)
    saved, report = store.bulk_review(project, [1, 4, 8, 1], 'Approve', ' Reviewer ', ' Checked ', True)
    assert report == {'saved': [1, 4, 8], 'skipped': []}
    reopened = Store(store.path)
    assert reopened.load(project) == saved
    events = reopened.audit(project)
    assert [e['row_id'] for e in events] == [1, 4, 8]
    for row_id, value in [(1, 'Healthcare'), (4, 'SMB'), (8, 'Public Sector')]:
        r = record(reopened, project, row_id)
        original = before['results'][row_id]
        assert r.final == r.proposed == value
        assert r.review_status == 'Approved' and r.reviewer == 'Reviewer' and r.notes == 'Checked'
        for field in ['row_id', 'identifier', 'proposed', 'candidates', 'evidence', 'match_status', 'research_snapshot']:
            assert r.dict()[field] == original[field]
        event = next(e for e in events if e['row_id'] == row_id)
        assert event['approved_value'] == value and event['timestamp'] == r.reviewed_at
    assert saved['results'][2] == before['results'][2]
    results = [Result(**r) for r in saved['results']]
    assert summary(results)['Pending Approvals'] == 7
    book = load_workbook(BytesIO(workbook(results, events, [], 10)))
    sheet = book['Reconciliation Results']
    headers = [c.value for c in sheet[1]]
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    for row_id in report['saved']:
        assert rows[row_id]['Final Value'] == saved['results'][row_id]['final']
        assert rows[row_id]['Reviewer Decision'] == 'Approve'
    assert book['Audit Trail'].max_row == 4


def test_mixed_eligibility_and_reject(workspace):
    store, project, _ = workspace
    before = store.load(project)
    saved, report = store.bulk_review(project, [1, 2, 3, 7, 999], 'Approve', 'Reviewer', confirmed=True)
    assert report['saved'] == [1]
    assert [r['row_id'] for r in report['skipped']] == [2, 3, 7, 999]
    assert all(r['reason'] for r in report['skipped'])
    for row_id in [2, 3, 7]:
        assert saved['results'][row_id] == before['results'][row_id]
    saved, report = store.bulk_review(project, [0, 2, 3], 'Reject', 'Reviewer', 'Rejected')
    assert report['saved'] == [0, 2, 3]
    assert record(store, project, 0).final == ''
    assert all(record(store, project, i).review_status == 'Rejected' for i in [0, 2, 3])
    assert not record(store, project, 3).final


def test_overwrite_confirmation_and_history(workspace):
    store, project, _ = workspace
    store.review(project, 1, 'Manual', 'First', manual_value='Custom')
    before = store.load(project)
    for versions in [None, {1: 'stale'}]:
        with pytest.raises(ValueError, match='Confirm overwriting'):
            store.bulk_review(project, [4, 1], 'Reject', 'Second', overwrite_versions=versions)
        assert store.load(project) == before and len(store.audit(project)) == 1
    versions = {1: digest(before['results'][1])}
    store.bulk_review(project, [4, 1], 'Reject', 'Second', 'New evidence', overwrite_versions=versions)
    events = store.audit(project)
    assert events[-1]['previous_decision'] == 'Manual' and events[-1]['previous_final'] == 'Custom'
    assert events[-1]['reviewer'] == 'Second' and events[-1]['new_final'] == ''
    with pytest.raises(ValueError, match='Confirm overwriting'):
        store.bulk_review(project, [1], 'Reject', 'Third', overwrite_versions=versions)


@pytest.mark.parametrize('action,reviewer,confirmed', [('Manual', 'R', True), ('Approve', '', True), ('Approve', 'R', False)])
def test_invalid_batch_never_saves(workspace, action, reviewer, confirmed):
    store, project, _ = workspace
    before = store.load(project)
    with pytest.raises(ValueError):
        store.bulk_review(project, [1, 4], action, reviewer, confirmed=confirmed)
    assert store.load(project) == before and not store.audit(project)


def test_transaction_rolls_back_all_rows(workspace):
    store, project, _ = workspace
    before = store.load(project)
    with store.connect() as db:
        db.execute("CREATE TRIGGER fail_review BEFORE INSERT ON audit WHEN (SELECT COUNT(*) FROM audit) = 1 BEGIN SELECT RAISE(ABORT, 'test failure'); END")
    with pytest.raises(sqlite3.IntegrityError):
        store.bulk_review(project, [1, 4], 'Approve', 'R', confirmed=True)
    assert store.load(project) == before and not store.audit(project)


def test_bulk_ai_uses_existing_evidence_rules(setup):
    research(setup)
    store, project, _, source, r = setup
    records = [r.dict()]
    for row_id in range(1, 4):
        copy = deepcopy(r.dict())
        copy['row_id'] = row_id
        if row_id == 1:
            copy['research_snapshot']['retrieved_evidence'][0]['content'] = 'Unrelated text'
        elif row_id == 2:
            copy['proposed'] = 'SMB'
        else:
            copy['recommendation_status'] = 'Insufficient Evidence'
        records.append(copy)
    store.save(project, {'source': source.to_dicts(), 'mapping': vars(M), 'results': records, 'context': 'source-context'})
    payload, report = store.bulk_review(project, [0, 1, 2, 3], 'Approve', 'R', confirmed=True)
    assert report['saved'] == [0] and len(report['skipped']) == 3
    assert payload['results'][0]['review_status'] == 'Approved — AI Recommendation'
    assert payload['results'][0]['research_snapshot'] == records[0]['research_snapshot']
    assert len(store.audit(project)) == 1


def test_ui_filtered_selection_and_immediate_update(workspace, monkeypatch):
    app = app_for(workspace, monkeypatch)
    navigate(app, 'Review & Approvals')
    assert by_label(app.selectbox, 'Bulk decision').options == ['Approve', 'Reject']
    by_label(app.checkbox, 'Select All filtered accounts').check().run()
    assert any('Selected accounts: 10 of 10' in c.value for c in app.caption)
    by_label(app.text_input, 'Search accounts').set_value('Pioneer').run()
    # Filtering preserves the checked Pioneer row; hidden rows cannot be acted on.
    assert any('Selected accounts: 1 of 1' in c.value for c in app.caption)
    by_label(app.checkbox, 'Select All filtered accounts').check().run()
    by_label(app.text_input, 'Bulk reviewer name').set_value('Bulk UI Reviewer')
    by_label(app.text_area, 'Bulk reviewer notes').set_value('Verified separately')
    by_label(app.checkbox, 'I have verified the evidence and field value for every selected account.').check()
    by_label(app.button, 'Save Review Decisions').click().run()
    assert not app.exception
    assert any('saved for 1 accounts' in s.value for s in app.success)
    assert by_label(app.metric, 'Pending Approvals').value == '9'
    store, project, _ = workspace
    assert record(store, project, 4).final == 'SMB'
    assert not record(store, project, 1).decision
    by_label(app.selectbox, 'Review status').set_value('All').run()
    assert app.dataframe[-1].value.iloc[0]['Final Value'] == 'SMB'
    by_label(app.checkbox, 'Select All filtered accounts').check().run()
    by_label(app.selectbox, 'Bulk decision').set_value('Reject')
    by_label(app.button, 'Save Review Decisions').click().run()
    assert any('Confirm overwriting' in e.value for e in app.error)
    by_label(app.checkbox, 'I confirm overwriting the existing saved decisions for these selected accounts.').check()
    by_label(app.button, 'Save Review Decisions').click().run()
    assert not app.exception and record(store, project, 4).review_status == 'Rejected'
    assert len(store.audit(project)) == 2


def test_ui_mixed_batch_reports_skips_and_exports(workspace, monkeypatch):
    app = app_for(workspace, monkeypatch)
    navigate(app, 'Review & Approvals')
    by_label(app.checkbox, 'Select All filtered accounts').check().run()
    by_label(app.button, 'Show Account List').click().run()
    # Individual row checkboxes update the Select All state as well as the count.
    row_box = next(c for c in app.checkbox if c.label.startswith('Select row 2 '))
    row_box.uncheck().run()
    assert not by_label(app.checkbox, 'Select All filtered accounts').value
    row_box = next(c for c in app.checkbox if c.label.startswith('Select row 2 '))
    row_box.check().run()
    assert by_label(app.checkbox, 'Select All filtered accounts').value
    by_label(app.text_input, 'Bulk reviewer name').set_value('R')
    by_label(app.checkbox, 'I have verified the evidence and field value for every selected account.').check()
    by_label(app.button, 'Save Review Decisions').click().run()
    assert not app.exception
    assert len([w for w in app.warning if 'Bulk review:' in w.value]) == 1
    assert any('Selected 10 · Saved 4 · Skipped 6 · Warnings 6' in w.value for w in app.warning)
    assert by_label(app.metric, 'Pending Approvals').value == '6'
    store, project, _ = workspace
    assert len(store.audit(project)) == 4
    # Exercise the actual export page after the bulk checkpoint.
    import rta.export
    original = rta.export.workbook
    exports = []
    def capture(*args, **kwargs):
        data = original(*args, **kwargs)
        exports.append(data)
        return data
    monkeypatch.setattr(rta.export, 'workbook', capture)
    navigate(app, 'Export Results')
    book = load_workbook(BytesIO(exports[-1]))
    sheet = book['Reconciliation Results']
    headers = [c.value for c in sheet[1]]
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert [(rows[i]['Final Value'], rows[i]['Review Status']) for i in [1, 4, 8]] == [
        ('Healthcare', 'Approved'), ('SMB', 'Approved'), ('Public Sector', 'Approved')]
    assert book['Audit Trail'].max_row == 5
