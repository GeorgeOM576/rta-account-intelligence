"""Stale production-style detection rows and cached exports; no Google calls."""
from copy import deepcopy
import csv
from io import BytesIO, StringIO
import json
from unittest.mock import Mock

from openpyxl import load_workbook
import pytest

from rta.export import translation_csv, workbook
from rta.language_detection import VERSION, RULES_FINGERPRINT, runtime_info
from rta.models import Result
from rta.storage import Store, digest
from rta.translation import TranslationEngine, detection_options, revision, translation_eligible
from rta.translation_usage import GoogleTranslationUsage
from tests.test_app import by_label, navigate
from tests.test_review_workflow import app_for
from tests.test_translation import make_workspace, Offline

NAMES = ['  HARBOR HEALTH  ', 'Cedar Labs', 'Summit Holdings', 'Pioneer Data', 'Orchard Systems', 'Bluebird Studio']
OLD = {'  HARBOR HEALTH  ': ('vi', 'Translation Needed'), 'Cedar Labs': ('ca', 'Uncertain'),
       'Summit Holdings': ('en', 'Uncertain'), 'Pioneer Data': ('nl', 'Translation Needed'),
       'Orchard Systems': ('de', 'Translation Needed'), 'Bluebird Studio': ('de', 'Translation Needed')}


@pytest.fixture
def stale_workspace(tmp_path):
    master = [{'Name': name, 'Segment': ['Enterprise', 'Healthcare', 'Commercial'][i % 3]} for i, name in enumerate(NAMES)]
    case = make_workspace(tmp_path, master, deepcopy(master))
    store, project, _ = case
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Name', 'Segment'], 'source': ['Name', 'Segment']})
    # Reproduce bad persisted rows AND detection-cache entries under the same
    # policy stamp: ordinary detect/records previously treated these as current.
    payload = store.load(project)
    with store.connect() as db:
        for item in engine.records():
            language, status = OLD.get(item['original'], ('und', 'Uncertain'))
            bad = {**item, 'language': language, 'detected_language': language, 'status': status,
                   'translation_excluded': False, 'detection_outcome': 'old result'}
            db.execute('UPDATE translations SET payload=? WHERE project=? AND dataset=? AND revision=? AND row_id=? AND field=?',
                       (json.dumps(bad), project, bad['dataset'], bad['revision'], bad['row_id'], bad['field']))
            signature = detection_options(payload, item['dataset'], item['field'])[3]
            detection = {key: bad[key] for key in ('language', 'confidence', 'status', 'detection_outcome', 'reason',
                                                  'detector_language', 'detector_confidence')}
            db.execute('INSERT OR REPLACE INTO language_detection_cache VALUES (?,?,?)',
                       (project, digest([signature, item['original']]), json.dumps(detection)))
    return case


def assert_fixed(records):
    names = [item for item in records if item['field'] == 'Name']
    segments = [item for item in records if item['field'] == 'Segment']
    assert len(names) == len(segments) == 12
    assert all(item['language'] == item['detected_language'] == 'en' and item['status'] == 'Not Required' for item in names)
    assert all(item['status'] == 'Excluded' and item['translation_excluded'] for item in segments)
    assert not any(translation_eligible(item) for item in records)


def test_force_reanalysis_repairs_current_stamp_rows_and_poisoned_cache(stale_workspace):
    store, project, _ = stale_workspace
    engine = TranslationEngine(store, project)
    assert next(r for r in engine.records() if r['original'].strip() == 'HARBOR HEALTH')['language'] == 'vi'
    # Ordinary detection deliberately retains existing current-policy per-cell
    # state; force must explicitly bypass that path as well as the cache.
    engine.detect(store.load(project)['translation_fields'])
    assert next(r for r in engine.records() if r['original'].strip() == 'HARBOR HEALTH')['language'] == 'vi'
    report = engine.reanalyse()
    assert report['cells_analyzed'] == 24 and report['unique_values_analyzed'] == 9
    assert report['build'] == RULES_FINGERPRINT
    assert_fixed(engine.records())
    assert all(item['detection_build'] == RULES_FINGERPRINT and item['detection_analyzed_at'] == report['analyzed_at'] for item in engine.records())
    # Repaired detection results are now reusable, not just one-time UI output.
    with store.connect() as db:
        db.execute('DELETE FROM translations WHERE project=?', (project,))
    engine.detect(store.load(project)['translation_fields'])
    assert_fixed(engine.records())


def test_detect_button_force_mode_also_repairs_same_stamp(stale_workspace):
    store, project, _ = stale_workspace
    engine = TranslationEngine(store, project)
    engine.detect(store.load(project)['translation_fields'], force=True)
    assert_fixed(engine.records())


def test_force_bypasses_in_memory_detection_without_clearing_global_caches(stale_workspace, monkeypatch):
    import rta.language_detection as detector
    real = detector._detect
    mock = Mock(return_value={'language': 'vi', 'confidence': 1, 'status': 'Translation Needed'})
    mock.__wrapped__ = real.__wrapped__
    monkeypatch.setattr(detector, '_detect', mock)
    assert detector.detect_text('HARBOR HEALTH', 'account_name')['language'] == 'vi'
    fresh = detector.detect_text('HARBOR HEALTH', 'account_name', refresh=True)
    assert fresh['language'] == 'en'
    assert mock.call_count == 1
    raw = Mock(return_value=('vi', 1))
    raw.__wrapped__ = Mock(return_value=('en', 1))
    monkeypatch.setattr(detector, 'raw_detection', raw)
    detector.detect_text('HARBOR HEALTH', 'account_name', refresh=True)
    raw.assert_not_called()
    raw.__wrapped__.assert_called_once()


def test_code_fingerprint_invalidates_same_version_persistent_cache(stale_workspace, monkeypatch):
    store, project, _ = stale_workspace
    monkeypatch.setattr('rta.translation.RULES_FINGERPRINT', 'new-rules-build-with-same-version')
    engine = TranslationEngine(store, project)
    assert_fixed(engine.records())
    assert all(r['detection_version'] == VERSION and r['detection_build'] == 'new-rules-build-with-same-version' for r in engine.records())


def test_preserves_reviews_corrections_audit_decisions_usage_and_other_workspaces(stale_workspace):
    store, project, _ = stale_workspace
    engine = TranslationEngine(store, project)
    rows = engine.records()
    approved = next(r for r in rows if r['dataset'] == 'master' and r['row_id'] == 0 and r['field'] == 'Name')
    approved = engine.approve(approved, 'Human corrected Harbor Health', 'Translation Reviewer', True, 'vi')
    manually_confirmed = next(r for r in engine.records() if r['dataset'] == 'master' and r['row_id'] == 1 and r['field'] == 'Name')
    manually_confirmed = engine.confirm_language(manually_confirmed, 'fr', 'Language Reviewer', True)
    payload = store.load(project)
    payload['results'] = [{'row_id': 0, 'identifier': NAMES[0], 'original': 'Verify', 'proposed': 'Enterprise',
                           'final': 'Enterprise', 'decision': 'Approve', 'reviewer': 'Account Reviewer'}]
    store.save(project, payload)
    usage = GoogleTranslationUsage(store, project)
    usage.confirm_request(usage.start_request(16))
    other, _ = store.create()
    store.save(other, deepcopy(payload))
    with store.connect() as db:
        # An unrelated workspace's cache sentinel must not be invalidated.
        db.execute('INSERT INTO language_detection_cache VALUES (?,?,?)', (other, 'sentinel', '{"untouched":true}'))
        db.execute('INSERT INTO translation_cache VALUES (?,?,?)', (project, 'provider-sentinel', '{"english":"keep me"}'))
    results_before = deepcopy(store.load(project)['results'])
    audit_before = store.audit(project)
    usage_before = usage.summary()
    engine.reanalyse()
    after = engine.records()
    approved_after = next(r for r in after if r['dataset'] == 'master' and r['row_id'] == 0 and r['field'] == 'Name')
    assert approved_after['status'] == 'Approved Translation'
    assert approved_after['english'] == approved['english']
    assert approved_after['language'] == 'vi' and approved_after['language_confirmed']
    # Automatic detection is refreshed independently of the preserved human value.
    assert approved_after['detected_language'] == 'en'
    assert approved_after['automatic_detection_outcome'].startswith('Confident English')
    assert approved_after['reviewer'] == approved['reviewer'] and approved_after['reviewed_at'] == approved['reviewed_at']
    confirmed_after = next(r for r in after if r['dataset'] == 'master' and r['row_id'] == 1 and r['field'] == 'Name')
    assert confirmed_after['language'] == 'fr' and confirmed_after['language_reviewer'] == manually_confirmed['language_reviewer']
    assert confirmed_after['detected_language'] == 'en'
    assert store.audit(project) == audit_before
    assert store.load(project)['results'] == results_before
    assert usage.summary() == usage_before
    with store.connect() as db:
        assert db.execute('SELECT payload FROM language_detection_cache WHERE project=? AND key=?', (other, 'sentinel')).fetchone()[0] == '{"untouched":true}'
        assert db.execute('SELECT payload FROM translation_cache WHERE project=?', (project,)).fetchone()[0] == '{"english":"keep me"}'
    assert store.load(other) == payload


def test_refreshes_deselected_current_fields_but_leaves_old_revisions_and_selection(stale_workspace):
    store, project, _ = stale_workspace
    engine = TranslationEngine(store, project)
    payload = store.load(project)
    selection = {'master': ['Name'], 'source': ['Name']}
    payload['translation_fields'] = selection
    store.save(project, payload)
    with store.connect() as db:
        item = json.loads(db.execute('SELECT payload FROM translations WHERE project=? LIMIT 1', (project,)).fetchone()[0])
        historical = {**item, 'revision': 'previous-upload', 'language': 'vi', 'detected_language': 'vi'}
        db.execute('INSERT INTO translations VALUES (?,?,?,?,?,?)',
                   (project, historical['dataset'], historical['revision'], historical['row_id'], historical['field'], json.dumps(historical)))
    engine.reanalyse()
    assert store.load(project)['translation_fields'] == selection
    assert len(engine.records()) == 12 and all(item['field'] == 'Name' for item in engine.records())
    with store.connect() as db:
        current = [json.loads(raw) for (raw,) in db.execute('SELECT payload FROM translations WHERE project=? AND revision<>?', (project, 'previous-upload'))]
        assert_fixed(current)
        previous = json.loads(db.execute('SELECT payload FROM translations WHERE project=? AND revision=?', (project, 'previous-upload')).fetchone()[0])
    assert previous == historical


def test_current_csv_and_excel_omit_excluded_and_export_refreshed_language(stale_workspace):
    store, project, _ = stale_workspace
    engine = TranslationEngine(store, project)
    before = translation_csv(engine.export_records())
    assert b'vi' in before
    engine.reanalyse()
    selected = engine.export_records()
    assert len(selected) == 12 and all(r['field'] == 'Name' for r in selected)
    data = translation_csv(selected)
    assert data != before
    rows = list(csv.DictReader(StringIO(data.decode('utf-8-sig'))))
    assert len(rows) == 12
    assert all(r['Detected Language'] == 'en' and r['Translation Status'] == 'Not Required' for r in rows)
    assert all(r['Detection Policy'] == VERSION and r['Detection Build'] == RULES_FINGERPRINT and r['Analyzed At'] for r in rows)
    book = load_workbook(BytesIO(workbook([], store.audit(project), [], 6,
        translations=engine.records(), datasets=store.load(project))))
    values = list(book['Translation Details'].values)
    records = [dict(zip(values[0], row)) for row in values[1:]]
    assert len(records) == 12 and all(r['field'] == 'Name' and r['detected_language'] == 'en' for r in records)
    for dataset in ('master', 'source'):
        headers = [cell.value for cell in book[f'{dataset.title()} Translations'][1]]
        assert 'Segment' in headers  # original uploads remain intact
        assert 'Segment (English)' not in headers


def test_legacy_missing_policy_is_upgraded_on_export_without_manual_reanalysis(stale_workspace):
    store, project, _ = stale_workspace
    with store.connect() as db:
        rows = db.execute('SELECT rowid,payload FROM translations WHERE project=?', (project,)).fetchall()
        for rowid, raw in rows:
            item = json.loads(raw)
            item.pop('detection_policy')
            db.execute('UPDATE translations SET payload=? WHERE rowid=?', (json.dumps(item), rowid))
        # Old build cache entries must not be reused by a different policy build.
        db.execute('UPDATE language_detection_cache SET key=? || key WHERE project=?', ('legacy-', project))
    current = TranslationEngine(store, project).export_records()
    assert len(current) == 12
    assert all(r['detected_language'] == 'en' and r['status'] == 'Not Required' for r in current)


def test_legacy_approved_without_reviewer_fields_is_preserved(stale_workspace):
    store, project, _ = stale_workspace
    with store.connect() as db:
        rowid, raw = db.execute('SELECT rowid,payload FROM translations WHERE project=? AND field=? LIMIT 1', (project, 'Name')).fetchone()
        item = json.loads(raw)
        item.update(status='Approved Translation', english='Legacy manual text', language='vi', reviewer='', reviewed_at='')
        db.execute('UPDATE translations SET payload=? WHERE rowid=?', (json.dumps(item), rowid))
    engine = TranslationEngine(store, project)
    engine.reanalyse()
    item = next(r for r in engine.records() if r['english'] == 'Legacy manual text')
    assert item['status'] == 'Approved Translation' and item['language'] == 'vi'
    assert item['detected_language'] == 'en'


def test_runtime_detects_disk_rules_changed_without_restarting(monkeypatch):
    assert not runtime_info()['restart_required']
    monkeypatch.setattr('rta.language_detection.Path.read_bytes', lambda path: b'changed on disk')
    assert runtime_info()['restart_required']


def test_reanalysis_refuses_to_write_with_old_loaded_rules(stale_workspace, monkeypatch):
    store, project, _ = stale_workspace
    monkeypatch.setattr('rta.translation.runtime_info', lambda: {'restart_required': True})
    with pytest.raises(ValueError, match='Restart Streamlit'):
        TranslationEngine(store, project).reanalyse()
    assert 'translation_detection_refresh' not in store.load(project)


def test_ui_reanalysis_and_download_use_new_current_records(stale_workspace, monkeypatch):
    store, project, _ = stale_workspace
    monkeypatch.setattr('rta.translation_ui.ArgosProvider', Offline)
    from google.cloud import translate_v3
    google = Mock(side_effect=AssertionError('Google client must never be created for reanalysis'))
    monkeypatch.setattr(translate_v3, 'TranslationServiceClient', google)
    # Offline rule reanalysis remains available to backend callers. The new UI
    # reads Google cache and offers only Google paid actions.
    engine = TranslationEngine(store, project)
    engine.reanalyse()
    assert_fixed(engine.records())
    from tests.translation_workspace_helpers import seed_google
    seed_google(engine, {r['original']: 'en' for r in engine.records() if r['field'] == 'Name'})
    app = app_for(stale_workspace, monkeypatch)
    navigate(app, 'Translation Intelligence')
    assert not app.exception
    assert set(app.dataframe[0].value['Detected Language']) == {'en'}
    assert set(app.dataframe[0].value['Original Company Name']) == {r['original'] for r in engine.records() if r['field'] == 'Name'}
    google.assert_not_called()
    assert GoogleTranslationUsage(store, project).summary()['requests'] == 0
    payload = store.load(project)
    payload['results'] = [Result(i, row['Name'], row['Segment'], proposed=row['Segment']).dict()
                          for i, row in enumerate(payload['master'])]
    store.save(project, payload)
    import rta.export as export
    original_workbook, excel = export.workbook, []
    def capture_excel(*args, **kwargs):
        assert len(kwargs['translations']) == 12
        assert all(r['field'] == 'Name' and r['detected_language'] == 'en' for r in kwargs['translations'])
        data = original_workbook(*args, **kwargs)
        excel.append(data)
        return data
    monkeypatch.setattr(export, 'workbook', capture_excel)
    navigate(app, 'Export Results')
    assert not app.exception
    book = load_workbook(BytesIO(excel[-1]))
    assert book['Translation Details'].max_row == 13
    google.assert_not_called()


def test_unapproved_manual_correction_survives_reanalysis(stale_workspace):
    store, project, _ = stale_workspace
    with store.connect() as db:
        rowid, raw = db.execute('SELECT rowid,payload FROM translations WHERE project=? AND field=? LIMIT 1', (project, 'Name')).fetchone()
        previous = json.loads(raw)
        previous.update(status='Translated — Review', english='Manually corrected English draft', correction=True)
        db.execute('UPDATE translations SET payload=? WHERE rowid=?', (json.dumps(previous), rowid))
    engine = TranslationEngine(store, project)
    engine.reanalyse()
    item = next(r for r in engine.records() if r.get('correction'))
    assert item['english'] == previous['english']
    assert item['language'] == previous['language']
    assert item['status'] == 'Translated — Review' and item['detected_language'] == 'en'
    assert not store.audit(project)


def test_current_csv_keeps_formula_like_values_as_literal_data(stale_workspace):
    store, project, _ = stale_workspace
    item = TranslationEngine(store, project).records()[0]
    item.update(original='=fictional()', english='@fictional')
    rows = list(csv.DictReader(StringIO(translation_csv([item]).decode('utf-8-sig'))))
    assert rows[0]['Original text'] == "'=fictional()"
    assert rows[0]['English Translation'] == "'@fictional"
