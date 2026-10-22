"""The approved CSV header is the sole visible schema; engines/exports stay full."""
from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from rta.workflow_table import COLUMNS, visible_table
from tests.test_app import navigate, by_label
from tests.test_workflow_translation_display import setup, block_live_services
from tests.workflow_display_helpers import csv_columns, rendered_table


def test_configuration_matches_independent_windows_1252_csv_header():
    assert list(COLUMNS) == csv_columns()
    assert len(COLUMNS) == len(set(COLUMNS)) == 31


def test_projection_maps_original_value_and_keeps_missing_columns_blank():
    raw = {'Original Value': 'Original segment', 'Master Identifier': 'Fictional company',
           'Final Value': '', 'Record ID': 42, 'Internal only': {'retain': True},
           'Master Translation Status': 'Approved Translation'}
    before = deepcopy(raw)
    displayed = visible_table([raw])
    assert list(displayed) == csv_columns() and len(displayed) == 31
    assert displayed['Master Original Value'] == ['Original segment']
    assert displayed['Master Approved English Translation'] == ['']
    assert displayed['Record ID'] == [42]
    assert 'Internal only' not in displayed and 'Master Translation Status' not in displayed
    assert raw == before
    assert visible_table([{'Master Original Value': 'Already mapped', 'Original Value': 'Wrong fallback'}])['Master Original Value'] == ['Already mapped']
    assert list(visible_table([])) == csv_columns() and all(not v for v in visible_table([]).values())


@pytest.mark.parametrize('page', ['Run Reconciliation', 'AI Research', 'Review & Approvals'])
def test_real_streamlit_table_is_exactly_31_columns_and_backend_unchanged(tmp_path, monkeypatch, page):
    monkeypatch.setenv('TAVILY_API_KEY', 'mock-only')
    monkeypatch.setenv('OPENAI_API_KEY', 'mock-only')
    case, engine, app = setup(tmp_path, monkeypatch)
    before = deepcopy(case[0].load(case[1]))
    audit = case[0].audit(case[1])
    navigate(app, page)
    assert not app.exception
    table = rendered_table(app)
    assert len(table.columns) == 31 and list(table.columns) == csv_columns()
    assert len(table) == 1
    row = table.iloc[0]
    assert row['Master Approved English Translation'] == 'The Gardens of Provence'
    assert row['Matched Source Original Name'] == 'Jardines de Provenza'
    assert row['Master Original Value'] == 'Verify'
    assert row['Proposed Value'] == 'SMB' and row['Final Value'] == '' and row['Record ID'] == 0
    assert 'Match Basis' not in table and 'Translation Match Provenance' not in table
    assert case[0].load(case[1]) == before and case[0].audit(case[1]) == audit
    saved = before['results'][0]
    assert saved['translation_provenance'] and saved['match_basis']
    assert saved['account_context']['matched_source']['approved_english'] == 'The Gardens of Provence'


def test_export_preview_and_excel_retain_full_existing_schema(tmp_path, monkeypatch):
    case, engine, app = setup(tmp_path, monkeypatch)
    import rta.export
    original, downloads = rta.export.workbook, []
    def capture(*args, **kwargs):
        data = original(*args, **kwargs);downloads.append(data)
        return data
    monkeypatch.setattr(rta.export, 'workbook', capture)
    navigate(app, 'Export Results')
    preview = rendered_table(app)
    assert len(preview.columns) > 31
    assert 'Original Value' in preview and 'Master Original Value' not in preview
    assert 'Match Basis' in preview and 'Matched Source Approved English Translation' in preview
    book = load_workbook(BytesIO(downloads[-1]))
    headers = [c.value for c in book['Reconciliation Results'][1]]
    assert headers == list(preview.columns)
    assert 'Translation Match Provenance' in headers and 'Master Translation Status' in headers
    assert {'Audit Trail', 'Master Translations', 'Source Translations'} <= set(book.sheetnames)


def test_review_empty_filter_keeps_exact_31_column_header(tmp_path, monkeypatch):
    case, engine, app = setup(tmp_path, monkeypatch)
    navigate(app, 'Review & Approvals')
    by_label(app.text_input, 'Search accounts').set_value('no-such-company-829').run()
    table = rendered_table(app)
    assert len(table) == 0 and list(table.columns) == csv_columns()
    assert by_label(app.selectbox, 'Account to review').value == 0
    assert case[0].load(case[1])['results'][0]['final'] == ''


def test_translation_intelligence_table_is_not_restricted(tmp_path, monkeypatch):
    case, engine, app = setup(tmp_path, monkeypatch)
    navigate(app, 'Translation Intelligence')
    assert not app.exception
    columns = list(app.dataframe[0].value.columns)
    assert columns != csv_columns()
    assert 'Original Company Name' in columns and 'English Translation' in columns
