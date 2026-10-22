"""Table-first UI and bounded SQL queries; external services always mocked."""
import json
from copy import deepcopy
from io import BytesIO
from unittest.mock import Mock
import pytest
from openpyxl import load_workbook
from rta.translation import TranslationEngine, cell_key
from rta.translation_workspace import TranslationWorkspace, row_id
from rta.translation_providers import GoogleTranslationProvider
from rta.translation_usage import GoogleTranslationUsage
from tests.test_google_detection import google_case
from tests.test_google_translation import response
from tests.test_review_workflow import app_for
from tests.test_app import by_label, navigate
from tests.translation_workspace_helpers import seed_google, select_visible, TRANSLATE_CONSENT, DECISION_CONSENT


@pytest.fixture
def table_case(tmp_path):
    master = [{'Name': 'Les Jardins de Provence'}, {'Name': 'Cedar Labs'}, {'Name': 'Les Jardins de Provence'}]
    source = [{'Name': 'Servicios de Construcción'}, {'Name': 'Unknown Brand'}]
    case, engine = google_case(tmp_path, master, source)
    seed_google(engine, {'Les Jardins de Provence': 'fr', 'Cedar Labs': 'en',
                        'Servicios de Construcción': 'es', 'Unknown Brand': 'und'},
                {'master': ['Name'], 'source': ['Name']})
    return case, engine


def test_query_search_filters_pagination_and_workspace_isolation(table_case):
    case, engine = table_case
    workspace = TranslationWorkspace(engine)
    rows, total, page = workspace.page()
    assert total == 5 and len(rows) == 5 and page == 1
    assert workspace.page(search='jardins')[1] == 2
    assert workspace.page(search='%')[1] == 0  # Literal search, not SQL wildcard.
    assert workspace.page('Source', status='Uncertain')[1] == 1
    assert workspace.page('Master', status='Not Required')[1] == 1
    assert len({row_id(r) for r in rows}) == 5
    other, _ = case[0].create()
    payload = deepcopy(case[0].load(case[1]))
    case[0].save(other, payload)
    assert TranslationWorkspace(TranslationEngine(case[0], other)).page()[1] == 0


def test_selected_only_duplicate_cache_stale_and_draft_preservation(table_case):
    case, engine = table_case
    rows = engine.records()
    duplicate = [r for r in rows if r['language'] == 'fr']
    client = Mock()
    client.translate_text.return_value = response('Gardens of Provence')
    provider = lambda: GoogleTranslationProvider(consent=True, client=client)
    plan = engine.plan(['master', 'source'], provider(), True, selection=duplicate[:1])
    assert plan['unique_texts'] == 1
    report = engine.process(['master', 'source'], provider(), allow_names=True, selection=duplicate[:1])
    assert report['drafts'] == 1 and client.translate_text.call_count == 1
    now = engine.records()
    assert next(r for r in now if cell_key(r) == cell_key(duplicate[1]))['english'] == ''
    with pytest.raises(ValueError, match='changed'):
        engine.process(['master'], provider(), allow_names=True, selection=duplicate[:1])
    assert client.translate_text.call_count == 1
    engine.process(['master'], provider(), allow_names=True, selection=duplicate[1:])
    assert client.translate_text.call_count == 1
    drafts = [r for r in engine.records() if r['english']]
    approved = engine.approve(drafts[0], drafts[0]['english'], 'Reviewer', True)
    engine.process(['master'], provider(), allow_names=True, selection=[approved, drafts[1]])
    assert client.translate_text.call_count == 1
    assert GoogleTranslationUsage(case[0], case[1]).summary()['requests'] == 1
    assert case[0].load(case[1])['results'] == []


def test_bulk_reject_atomic_stale_and_audited(table_case):
    case, engine = table_case
    rows = [r for r in engine.records() if r['language'] == 'fr']
    client = Mock(); client.translate_text.return_value = response('Gardens of Provence')
    engine.process(['master'], GoogleTranslationProvider(consent=True, client=client), allow_names=True, selection=rows)
    drafts = [r for r in engine.records() if r['english']]
    with pytest.raises(ValueError):
        engine.bulk_reject(drafts, 'Reviewer', False)
    engine.approve(drafts[1], drafts[1]['english'], 'Reviewer', True)
    with pytest.raises(ValueError, match='changed'):
        engine.bulk_reject(drafts, 'Reviewer', True)
    assert next(r for r in engine.records() if cell_key(r) == cell_key(drafts[0]))['status'] == 'Translated — Review'
    assert len(case[0].audit(case[1])) == 1
    engine.bulk_reject(drafts[:1], 'Reviewer', True)
    assert len(case[0].audit(case[1])) == 2
    assert next(r for r in engine.records() if cell_key(r) == cell_key(drafts[0]))['draft_english']


def test_layout_selection_refresh_consent_export_and_individual_review(table_case, monkeypatch):
    case, engine = table_case
    client = Mock(); client.translate_text.return_value = response('Gardens of Provence')
    from google.cloud import translate_v3
    factory = Mock(return_value=client)
    monkeypatch.setattr(translate_v3, 'TranslationServiceClient', factory)
    app = app_for(case, monkeypatch); navigate(app, 'Translation Intelligence')
    assert [h.value for h in app.subheader][-4:] == ['Account translations', 'Detect Languages', 'Translate Selected Accounts', 'Review and Approval']
    assert app.dataframe[0].value.shape[0] == 5
    assert 'Original Company Name' in app.dataframe[0].value.columns
    assert all(e.label == 'Save / resume workspace' for e in app.expander) and not app.metric
    assert not any('provider' in w.label.lower() for w in [*app.radio, *app.selectbox])
    assert not by_label(app.button, 'Translate Selected').disabled
    by_label(app.text_input, 'Search accounts').set_value('Jardins').run()
    select_visible(app)
    prefix = f'translation_workspace_{case[1]}'
    assert len(app.session_state[prefix + '_selection']) == 2
    by_label(app.text_input, 'Search accounts').set_value('Cedar').run()
    assert len(app.session_state[prefix + '_selection']) == 2
    by_label(app.button, 'Translate Selected').click().run()
    factory.assert_not_called()
    assert app.error
    by_label(app.checkbox, TRANSLATE_CONSENT).check()
    by_label(app.button, 'Translate Selected').click().run()
    assert not app.exception and client.translate_text.call_count == 1
    assert by_label(app.text_input, 'Search accounts').value == 'Cedar'
    assert len(app.session_state[prefix + '_selection']) == 2
    assert not by_label(app.button, 'Translate Selected').disabled
    by_label(app.text_input, 'Translation reviewer').set_value('Bulk reviewer')
    by_label(app.button, 'Approve Selected').click().run()
    assert app.error
    by_label(app.checkbox, DECISION_CONSENT).check()
    by_label(app.button, 'Approve Selected').click().run()
    assert not app.exception
    assert sum(r['status'] == 'Approved Translation' for r in engine.records()) == 2
    by_label(app.text_input, 'Search accounts').set_value('').run()
    assert (app.dataframe[0].value['Translation Status'] == 'Approved Translation').sum() == 2
    by_label(app.button, 'Export to Excel').click().run()
    data = app.session_state[prefix + '_export']
    book = load_workbook(BytesIO(data))
    assert book['Translation Details'].max_row == 6
    assert book['Master Translations'].max_row == 4 and book['Source Translations'].max_row == 3
    assert book['Audit Trail'].max_row == 3
    client.detect_language.assert_not_called()
    assert not case[0].load(case[1])['results']


def test_large_dataset_query_only_decodes_page(tmp_path, monkeypatch):
    case, engine = google_case(tmp_path)
    seed_google(engine, {}, {'master': ['Name'], 'source': ['Name']})
    template = engine.records()[0]
    # Real 52k persisted rows; no giant editable frame or engine.records read.
    with case[0].connect() as db:
        db.executemany('INSERT OR REPLACE INTO translations VALUES (?,?,?,?,?,?)',
            [(case[1], template['dataset'], template['revision'], i, template['field'],
              json.dumps({**template, 'row_id': i, 'original': f'Business {i}'})) for i in range(52001)])
    monkeypatch.setattr(engine, 'records', Mock(side_effect=AssertionError('Unbounded record read')))
    workspace = TranslationWorkspace(engine)
    rows, total, page = workspace.page('Master', page=1000, size=25)
    assert total == 52001 and len(rows) == 25 and page == 1000
    assert rows[0]['row_id'] == 24975
    assert workspace.page(search='Business 52000')[1] == 1


def test_table_individual_edits_approve_reject_and_stale_rollback(table_case, monkeypatch):
    case, engine = table_case
    drafts = [r for r in engine.records() if r['language'] == 'fr']
    client = Mock(); client.translate_text.return_value = response('Gardens of Provence')
    engine.process(['master'], GoogleTranslationProvider(consent=True, client=client), allow_names=True, selection=drafts)
    app = app_for(case, monkeypatch); navigate(app, 'Translation Intelligence')
    key = app.dataframe[0].key
    assert key
    app.session_state[key] = {'edited_rows': {0: {'English Translation': 'Verified Gardens', 'Individual action': 'Approve'}},
                              'added_rows': [], 'deleted_rows': []}
    by_label(app.text_input, 'Translation reviewer').set_value('Table reviewer')
    by_label(app.checkbox, DECISION_CONSENT).check()
    by_label(app.button, 'Apply individual actions and edits').click().run()
    assert not app.exception
    approved = next(r for r in engine.records() if r['row_id'] == 0 and r['dataset'] == 'master')
    assert approved['status'] == 'Approved Translation' and approved['english'] == 'Verified Gardens'
    assert app.dataframe[0].value.iloc[0]['English Translation'] == 'Verified Gardens'
    key = app.dataframe[0].key
    app.session_state[key] = {'edited_rows': {2: {'Individual action': 'Reject'}}, 'added_rows': [], 'deleted_rows': []}
    by_label(app.text_input, 'Translation reviewer').set_value('Table reviewer')
    by_label(app.checkbox, DECISION_CONSENT).check()
    by_label(app.button, 'Apply individual actions and edits').click().run()
    assert not app.exception
    rejected = next(r for r in engine.records() if r['row_id'] == 2 and r['dataset'] == 'master')
    assert rejected['status'] == 'Rejected Translation' and not rejected['english']
    assert len(case[0].audit(case[1])) == 2 and not case[0].load(case[1])['results']


def test_pagination_preserves_selection_and_select_all_is_visible_only(tmp_path, monkeypatch):
    master = [{'Name': f'Fictional Company {i}'} for i in range(60)]
    case, engine = google_case(tmp_path, master, [{'Name': 'Source Company'}])
    seed_google(engine, {r['Name']: 'en' for r in master}, {'master': ['Name'], 'source': ['Name']})
    app = app_for(case, monkeypatch); navigate(app, 'Translation Intelligence')
    select_visible(app)
    prefix = f'translation_workspace_{case[1]}'
    assert len(app.session_state[prefix + '_selection']) == 25
    by_label(app.number_input, 'Translation page').set_value(2).run()
    assert len(app.dataframe[0].value) == 25 and not app.dataframe[0].value['Select'].any()
    select_visible(app)
    assert len(app.session_state[prefix + '_selection']) == 50
    by_label(app.number_input, 'Translation page').set_value(1).run()
    assert app.dataframe[0].value['Select'].all()
    by_label(app.text_input, 'Search accounts').set_value('Company 59').run()
    assert len(app.dataframe[0].value) == 1
    select_visible(app)
    assert len(app.session_state[prefix + '_selection']) == 51
    by_label(app.button, 'Clear selection').click().run()
    assert not app.session_state[prefix + '_selection']


def test_cached_detection_applies_to_existing_pending_record_without_policy_change(table_case):
    case, engine = table_case
    records = engine.records()
    # Existing pending rows share the current policy, but a raw cache arrives later.
    from rta.google_detection import raw_key
    item = next(r for r in records if r['original'] == 'Unknown Brand')
    item.update(google_language='', language='und', status='Uncertain')
    with case[0].connect() as db:
        db.execute('UPDATE translations SET payload=? WHERE project=? AND dataset=? AND revision=? AND row_id=? AND field=?',
                   (json.dumps(item), case[1], *cell_key(item)))
        db.execute('DELETE FROM language_detection_cache WHERE project=? AND key=?', (case[1], raw_key(item['original'])))
    from rta.google_detection import GOOGLE_POLICY
    with case[0].connect() as db:
        db.execute('INSERT INTO language_detection_cache VALUES (?,?,?)',
                   (case[1], raw_key(item['original']), json.dumps({'language_code': 'fr', 'confidence': .99, 'provider': 'google', 'policy': GOOGLE_POLICY})))
    # A full export/read must apply newly arrived cache entries without a paid call.
    updated = next(r for r in engine.records() if row_id(r) == row_id(item))
    assert updated['google_language'] == 'fr' and updated['status'] == 'Translation Needed'


def test_detection_consent_scope_rejects_reuploaded_data_before_dispatch(table_case, monkeypatch):
    case, engine = table_case
    workspace = TranslationWorkspace(engine)
    payload = case[0].load(case[1]); payload['master'][0]['Name'] = 'Changed name'
    case[0].save(case[1], payload)
    run = Mock(side_effect=AssertionError('No paid calls'))
    monkeypatch.setattr('rta.translation_workspace.run_detection', run)
    with pytest.raises(ValueError, match='changed'):
        workspace.run_detection(250, True)
    run.assert_not_called()


def test_unicode_search_and_existing_draft_cannot_be_requeued_implicitly(tmp_path, monkeypatch):
    case, engine = google_case(tmp_path, [{'Name': 'Société des Étoiles'}], [{'Name': 'Cedar Labs'}])
    seed_google(engine, {'Société des Étoiles': 'fr', 'Cedar Labs': 'en'}, {'master': ['Name'], 'source': ['Name']})
    workspace = TranslationWorkspace(engine)
    assert workspace.page(search='SOCIÉTÉ')[1] == 1
    client = Mock(); client.translate_text.return_value = response('Star Society')
    from google.cloud import translate_v3
    monkeypatch.setattr(translate_v3, 'TranslationServiceClient', Mock(return_value=client))
    selected = workspace.page('Master')[0]
    engine.process(['master'], GoogleTranslationProvider(consent=True, client=client), allow_names=True, selection=selected)
    draft = next(r for r in engine.records() if r['english'])
    app = app_for(case, monkeypatch); navigate(app, 'Translation Intelligence')
    by_label(app.text_input, 'Individual reviewer').set_value('Language reviewer')
    by_label(app.checkbox, 'I verified this translation preserves the correct entity, proper names, legal suffixes and identifiers.').check()
    by_label(app.button, 'Save confirmed language').click().run()
    assert app.error
    assert next(r for r in engine.records() if r['english']) == draft
    select_visible(app)
    assert not by_label(app.button, 'Translate Selected').disabled
    assert client.translate_text.call_count == 1


def test_bulk_ui_retains_stale_selected_snapshots(table_case, monkeypatch):
    case, engine = table_case
    drafts = [r for r in engine.records() if r['language'] == 'fr']
    client = Mock(); client.translate_text.return_value = response('Gardens of Provence')
    engine.process(['master'], GoogleTranslationProvider(consent=True, client=client), allow_names=True, selection=drafts)
    app = app_for(case, monkeypatch); navigate(app, 'Translation Intelligence')
    by_label(app.selectbox, 'Translation status').set_value('Translated — Review').run()
    select_visible(app)
    current = [r for r in engine.records() if r['english']]
    engine.approve(current[1], 'Other reviewer correction', 'Other reviewer', True)
    app.run()
    assert any('changed in another session' in warning.value for warning in app.warning)
    by_label(app.text_input, 'Translation reviewer').set_value('Bulk reviewer')
    by_label(app.checkbox, DECISION_CONSENT).check()
    by_label(app.button, 'Approve Selected').click().run()
    assert any('changed' in error.value for error in app.error)
    assert len(case[0].audit(case[1])) == 1
    assert next(r for r in engine.records() if row_id(r) == row_id(current[0]))['status'] == 'Translated — Review'


def test_legacy_offline_draft_survives_google_only_page_and_requires_language_review(tmp_path, monkeypatch):
    from tests.test_translation import make_workspace, Offline
    case = make_workspace(tmp_path)
    engine = TranslationEngine(case[0], case[1])
    engine.detect({'master': ['Name'], 'source': ['Name']})
    engine.process(['master'], Offline(), allow_names=True)
    original = next(r for r in engine.records() if r['english'])
    from google.cloud import translate_v3
    factory = Mock(side_effect=AssertionError('Migration must not call Google'))
    monkeypatch.setattr(translate_v3, 'TranslationServiceClient', factory)
    app = app_for(case, monkeypatch); navigate(app, 'Translation Intelligence')
    current = next(r for r in engine.records() if row_id(r) == row_id(original))
    assert current['english'] == original['english']
    assert current['status'] == 'Translated — Review' and current['provider'] == original['provider']
    assert current['google_language'] == '' and current['draft_language_review_required']
    select_visible(app)
    assert not by_label(app.button, 'Translate Selected').disabled
    with pytest.raises(ValueError):
        engine.bulk_approve([current], 'Reviewer', True)
    by_label(app.text_input, 'Individual reviewer').set_value('Language reviewer')
    by_label(app.checkbox, 'I verified this translation preserves the correct entity, proper names, legal suffixes and identifiers.').check()
    by_label(app.button, 'Save confirmed language').click().run()
    assert not app.exception
    current = next(r for r in engine.records() if row_id(r) == row_id(original))
    assert current['english'] == original['english'] and current['status'] == 'Translated — Review'
    assert current['language_confirmed'] and not current['draft_language_review_required']
    assert app.dataframe[0].value.iloc[0]['Review advice'] == 'Language confirmed by reviewer'
    assert engine.bulk_approve([current], 'Reviewer', True) == 1
    assert not case[0].load(case[1])['results']
    factory.assert_not_called()


def test_cached_google_refresh_keeps_existing_drafts_when_language_changes(table_case):
    case, engine = table_case
    drafts = [r for r in engine.records() if r['language'] == 'fr']
    client = Mock(); client.translate_text.return_value = response('Gardens of Provence')
    engine.process(['master'], GoogleTranslationProvider(consent=True, client=client), allow_names=True, selection=drafts)
    before = [r for r in engine.records() if r['english']]
    workspace = TranslationWorkspace(engine)
    workspace.detect(workspace.payload['translation_fields'], force=True,
                     google_results={'Les Jardins de Provence': {'language_code': 'en', 'confidence': .99}})
    after = [r for r in engine.records() if r['english']]
    assert [r['english'] for r in after] == [r['english'] for r in before]
    assert all(r['status'] == 'Translated — Review' and r['draft_language_review_required'] for r in after)
    assert all(r['detected_language'] == 'en' and r['google_language'] == 'en' for r in after)
    assert not case[0].audit(case[1]) and not case[0].load(case[1])['results']
    with pytest.raises(ValueError):
        engine.bulk_approve(after, 'Reviewer', True)


@pytest.mark.parametrize('empty', [None, ''])
def test_cleared_table_translation_cannot_be_approved_as_literal_none(table_case, monkeypatch, empty):
    case, engine = table_case
    rows = [r for r in engine.records() if r['language'] == 'fr'][:1]
    client = Mock(); client.translate_text.return_value = response('Gardens of Provence')
    engine.process(['master'], GoogleTranslationProvider(consent=True, client=client), allow_names=True, selection=rows)
    original = next(r for r in engine.records() if r['english'])
    app = app_for(case, monkeypatch); navigate(app, 'Translation Intelligence')
    app.session_state[app.dataframe[0].key] = {'edited_rows': {0: {'English Translation': empty, 'Individual action': 'Approve'}},
                                             'added_rows': [], 'deleted_rows': []}
    by_label(app.text_input, 'Translation reviewer').set_value('Reviewer')
    by_label(app.checkbox, DECISION_CONSENT).check()
    by_label(app.button, 'Apply individual actions and edits').click().run()
    assert any('empty' in error.value for error in app.error)
    assert next(r for r in engine.records() if r['english']) == original
    assert not case[0].audit(case[1])
