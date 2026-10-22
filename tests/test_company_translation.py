"""All selected foreign companies translate without any reconciliation run."""
from copy import deepcopy
from io import BytesIO
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from rta.google_detection import run_detection, plan_detection
from rta.translation import TranslationEngine, company_overview, translation_eligible
from rta.translation_providers import GoogleTranslationProvider
from rta.translation_usage import GoogleTranslationUsage
from rta.export import workbook
from rta.models import REVIEW_FILTERS
from tests.test_google_detection import google_case, detected
from tests.test_google_translation import response
from tests.test_translation import Offline
from tests.test_review_workflow import app_for
from tests.test_app import navigate, by_label

DATA = {
    'Servicios de Aurora': ('es', 'Aurora Services'),
    'Les Jardins de Provence': ('fr', 'Gardens of Provence'),
    'Global Tech Solutions': ('en', ''),
    'Empresa Sin Pareja': ('es', 'Unmatched Company'),
    'Werkstätten der Morgen': ('de', 'Morning Workshops'),
    'Élan': ('fr', 'Elan'),
    '北京科技有限公司': ('zh', 'Beijing Technology Company'),
}


@pytest.fixture
def companies(tmp_path):
    master = [{'Name': text, 'Segment': 'Enterprise'} for text in list(DATA)[:4]]
    source = [{'Name': text, 'Segment': 'Commercial'} for text in
              ['Werkstätten der Morgen', 'Les Jardins de Provence', 'Élan', '北京科技有限公司']]
    case, engine = google_case(tmp_path, master, source)
    engine.detect({'master': ['Name'], 'source': ['Name']})
    client = Mock()
    client.detect_language.side_effect = lambda **kw: detected(DATA[kw['request']['content']][0],
        0.56 if kw['request']['content'] == 'Élan' else 0.99)
    client.translate_text.side_effect = lambda **kw: response(*(DATA[text][1] for text in kw['request']['contents']))
    run_detection(engine, ['master', 'source'], consent=True, client=client)
    return case, engine, client


def process_all(engine, client):
    reports = []
    while engine.plan(['master', 'source'], GoogleTranslationProvider(), True, 2)['unique_texts']:
        reports.append(engine.process(['master', 'source'], GoogleTranslationProvider(consent=True, client=client),
                                      allow_names=True, limit=2))
        assert len(reports) <= 3
    return reports


def test_unmatched_master_only_source_only_dedup_batches_and_coverage(companies):
    case, engine, client = companies
    store, project, _ = case
    before = deepcopy(store.load(project))
    assert not before['results']
    assert client.detect_language.call_count == 7
    assert plan_detection(engine, ['master', 'source'])['requests'] == 0
    reports = process_all(engine, client)
    assert [r['remaining_unique_texts'] for r in reports] == [4, 2, 0]
    assert sum(r['drafts'] for r in reports) == 7
    submitted = [text for call in client.translate_text.call_args_list for text in call.kwargs['request']['contents']]
    assert len(submitted) == len(set(submitted)) == 6
    assert 'Global Tech Solutions' not in submitted and 'Élan' in submitted
    assert 'Empresa Sin Pareja' in submitted and 'Werkstätten der Morgen' in submitted
    assert all(not call.kwargs['retry'] for call in client.translate_text.call_args_list)
    items = engine.records()
    assert sum(r['status'] == 'Translated — Review' for r in items) == 7
    assert store.load(project)['master'] == before['master'] and store.load(project)['source'] == before['source']
    assert not store.load(project)['results']
    overview = company_overview(items)
    assert overview['Total unique company names'] == 7
    assert overview['English names'] == 1 and overview['Foreign-language names detected'] == 6
    assert overview['Pending translation review'] == 6
    assert overview['Translation coverage (%)'] == 100
    assert overview['Ambiguous names requiring manual language review'] == 1
    assert overview['Approved translations'] == 0
    assert overview['Remaining untranslated foreign names'] == 0
    calls = client.translate_text.call_count
    engine.process(['master', 'source'], GoogleTranslationProvider(consent=True, client=client), allow_names=True)
    assert client.translate_text.call_count == calls
    assert len(REVIEW_FILTERS) == 6


def test_manual_confirmation_rejection_correction_selected_bulk_and_cache(companies):
    case, engine, client = companies
    store, project, _ = case
    process_all(engine, client)
    ambiguous = next(r for r in engine.records() if r['original'] == 'Élan')
    confirmed = engine.confirm_language(ambiguous, 'es', 'Language reviewer', True)
    assert translation_eligible(confirmed)
    engine.process(['master', 'source'], GoogleTranslationProvider(consent=True, client=client), allow_names=True)
    french = [r for r in engine.records() if r['original'] == 'Les Jardins de Provence']
    with pytest.raises(ValueError):
        engine.bulk_approve(french, 'Reviewer', False)
    assert engine.bulk_approve(french, 'Reviewer', True) == 2
    assert company_overview(engine.records())['Approved translations'] == 1
    german = next(r for r in engine.records() if r['original'] == 'Werkstätten der Morgen')
    rejected = engine.reject(german, 'Reviewer', True)
    assert rejected['english'] == '' and rejected['draft_english'] == german['english']
    assert rejected['status'] == 'Rejected Translation' and not translation_eligible(rejected)
    corrected = engine.approve(rejected, 'Morning Technical Workshops', 'Reviewer', True, 'de')
    assert corrected['correction'] and corrected['status'] == 'Approved Translation'
    audit = store.audit(project)
    # A cached draft is reused for a newly confirmed cell in this workspace;
    # the reviewed human decision is explicitly made here, not by cache lookup.
    master = next(r for r in engine.records() if r['original'] == 'Servicios de Aurora')
    engine.reject(master, 'Reviewer', True)
    current = next(r for r in engine.records() if r['original'] == 'Servicios de Aurora')
    engine.confirm_language(current, 'es', 'Reviewer', True)
    count = client.translate_text.call_count
    engine.process(['master', 'source'], GoogleTranslationProvider(consent=True, client=client), allow_names=True)
    assert client.translate_text.call_count == count
    assert next(r for r in engine.records() if r['original'] == master['original'])['status'] == 'Translated — Review'
    assert len(store.audit(project)) > len(audit)
    assert not store.load(project)['results']


def test_bulk_stale_cell_rolls_back_every_approval_and_audit(companies):
    case, engine, client = companies
    process_all(engine, client)
    selected = [r for r in engine.records() if r['original'] == 'Les Jardins de Provence']
    engine.reject(selected[1], 'Other reviewer', True)
    before = case[0].audit(case[1])
    with pytest.raises(ValueError, match='changed'):
        engine.bulk_approve(selected, 'Bulk reviewer', True)
    assert next(r for r in engine.records() if r['dataset'] == selected[0]['dataset'] and
                r['original'] == selected[0]['original'])['status'] == 'Translated — Review'
    assert case[0].audit(case[1]) == before


def test_workspace_isolation(companies):
    case, engine, client = companies
    store, project, _ = case
    process_all(engine, client)
    other, _ = store.create()
    store.save(other, deepcopy(store.load(project)))
    other_engine = TranslationEngine(store, other)
    other_engine.detect(store.load(other)['translation_fields'])
    assert all(not r['english'] for r in other_engine.records())
    assert plan_detection(other_engine, ['master', 'source'])['cached_texts'] == 0
    assert GoogleTranslationUsage(store, other).summary()['requests'] == 0


def test_full_excel_and_ui_export_before_reconciliation(companies, monkeypatch):
    case, engine, client = companies
    store, project, _ = case
    process_all(engine, client)
    monkeypatch.setattr('rta.translation_ui.ArgosProvider', Offline)
    from google.cloud import translate_v3
    factory = Mock(side_effect=AssertionError('No Google calls during review/export'))
    monkeypatch.setattr(translate_v3, 'TranslationServiceClient', factory)
    import rta.export as export
    exports, original = [], export.workbook
    def capture(*args, **kwargs):
        data = original(*args, **kwargs)
        exports.append(data)
        return data
    monkeypatch.setattr(export, 'workbook', capture)
    app = app_for(case, monkeypatch)
    navigate(app, 'Translation Intelligence')
    assert not store.load(project)['results']
    from tests.translation_workspace_helpers import select_visible, DECISION_CONSENT
    by_label(app.text_input, 'Search accounts').set_value('Les Jardins').run()
    # Choose one stable draft snapshot, then exercise the real bulk form.
    draft = next(r for r in engine.records() if r['status'] == 'Translated — Review')
    from rta.translation_workspace import row_id
    app.session_state[f'translation_workspace_{project}_selection'] = {row_id(draft): draft}
    app.run()
    by_label(app.text_input, 'Translation reviewer').set_value('UI reviewer')
    by_label(app.button, 'Approve Selected').click().run()
    assert app.error
    by_label(app.checkbox, DECISION_CONSENT).check()
    by_label(app.button, 'Approve Selected').click().run()
    assert not app.exception
    assert sum(r['status'] == 'Approved Translation' for r in engine.records()) == 1
    navigate(app, 'Export Results')
    assert not app.exception and exports
    book = load_workbook(BytesIO(exports[-1]))
    assert book['Master Translations'].max_row == 5 and book['Source Translations'].max_row == 5
    for dataset in ('master', 'source'):
        raw = list(book[f'{dataset.title()} Translations'].values)
        exported = [dict(zip(raw[0], values)) for values in raw[1:]]
        for before, after in zip(store.load(project)[dataset], exported):
            assert after['Name'] == before['Name'] and after['Segment'] == before['Segment']
            assert 'Name (English)' in after and 'Name (English) status' in after
    values = list(book['Translation Details'].values)
    rows = [dict(zip(values[0], row)) for row in values[1:]]
    assert len(rows) == 8
    assert {'Original Company Name', 'Detected Language', 'English Translation', 'Translation Status', 'Translation Provider'} <= set(values[0])
    assert any(r['Translation Status'] == 'Translated — Review' and r['Original Company Name'] == 'Élan' and r['language_review_required'] for r in rows)
    assert any(r['Translation Status'] == 'Not Required' and r['Original Company Name'] == 'Global Tech Solutions' for r in rows)
    assert sum(bool(r['English Translation']) for r in rows) == 7
    assert not store.load(project)['results']
    factory.assert_not_called()
