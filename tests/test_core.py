from io import BytesIO
import json
import time
import polars as pl
import pytest
import requests
from openpyxl import Workbook, load_workbook
from rta.models import Mapping, Result, valid
from rta.matching import Matcher, CONFLICT_NOTE, RELATION_NOTE
from rta.ingest import parse_file, identifier_stats
from rta.review import decide
from rta.storage import Store, digest
from rta.research import Researcher, ResearchSettings, eligible
from rta.export import workbook

M = Mapping('Account', 'Company', 'Segment', 'Value')


def run(names, values, source_names, source_values, **kwargs):
    master = pl.DataFrame({'Account': names, 'Segment': values})
    source = pl.DataFrame({'Company': source_names, 'Value': source_values})
    matcher = Matcher(source, M, **kwargs)
    return matcher.run(master), matcher


def test_exact_precedence_and_three_suggestions():
    results, _ = run(['Alpha'], ['SMB'], ['Alfa', 'Alpha', 'Alphaa'], ['Other', 'SMB', 'Other'])
    r = results[0]
    assert r.match_status == 'Exact'
    assert r.proposed == 'SMB' and r.final == ''
    assert r.candidates[0]['account'] == 'Alpha'
    assert len(r.candidates) == 3


def test_normalized():
    results, _ = run([' ALPHA, Inc. '], ['Enterprise'], ['Alpha Inc'], ['Enterprise'])
    assert results[0].match_status == 'Normalized'
    assert results[0].proposed == 'Enterprise' and results[0].final == ''


def test_fuzzy_never_approves():
    results, _ = run(['Bluebird Studio'], ['SMB'], ['Bluebird Studios'], ['SMB'])
    assert results[0].candidates[0]['similarity'] > 80
    assert results[0].match_status == 'Fuzzy candidates'
    assert results[0].final == ''
    assert results[0].identity_confidence == 'Unverified'


def test_known_alias():
    results, _ = run(['Old Name'], ['SMB'], ['New Name'], ['SMB'], aliases={'Old Name': 'New Name'})
    assert results[0].match_status == 'Alias'
    assert results[0].proposed == 'SMB' and results[0].final == ''


@pytest.mark.parametrize('source_names,source_values,status', [
    (['Alpha', 'Alpha'], ['SMB', 'SMB'], 'Duplicate identifier'),
    (['Alpha', 'Alpha'], ['SMB', 'Enterprise'], 'Duplicate conflict'),
    (['Alpha', 'ALPHA'], ['SMB', 'Enterprise'], 'Duplicate conflict')])
def test_duplicate_source(source_names, source_values, status):
    results, matcher = run(['Alpha'], ['SMB'], source_names, source_values)
    assert results[0].match_status == status
    assert not results[0].final
    assert results[0].notes == CONFLICT_NOTE
    assert len(matcher.conflicts) == 2


def test_master_names_not_globally_unique():
    results, _ = run(['Alpha', 'Alpha'], ['SMB', 'SMB'], ['Alpha'], ['SMB'])
    assert all(r.match_status == 'Duplicate identifier' and not r.final for r in results)


def test_missing_identifier():
    results, _ = run(['', None], ['Null', 'SMB'], ['Alpha'], ['SMB'])
    assert all(r.match_status == 'Missing identifier' and not r.final and not r.candidates for r in results)


@pytest.mark.parametrize('placeholder', ['Verify', ' verified ', 'Unmapped', 'DUPLICATE', 'Blank', 'Null', 'N/A', 'N / A', 'None', '  ', 'N\t/A'])
def test_placeholders(placeholder):
    results, _ = run(['Alpha'], [placeholder], ['Alpha'], ['Enterprise'])
    r = results[0]
    assert not valid(placeholder)
    assert r.original == placeholder and r.final == '' and r.proposed == 'Enterprise'
    with pytest.raises(ValueError):
        decide(r, 'Manual value', 'Reviewer', manual_value=placeholder)


def test_unapproved_approved_rejected_and_retained():
    results, _ = run(['Alpha'], ['SMB'], ['Alpha'], ['Enterprise'])
    r = results[0]
    assert r.proposed == 'Enterprise' and not r.final
    event = decide(r, 'Approve proposal', 'Reviewer', 'Checked evidence')
    assert r.final == 'Enterprise' and r.review_status == 'Approved'
    assert event['approved_value'] == 'Enterprise' and event['timestamp']
    decide(r, 'Reject', 'Reviewer')
    assert not r.final and r.review_status == 'Rejected'
    decide(r, 'Retain original', 'Reviewer')
    assert r.final == 'SMB'


def test_parent_requires_review():
    results, _ = run(['Child'], ['Enterprise'], ['Parent'], ['Enterprise'], relationships={'Child': 'Parent'})
    r = results[0]
    assert r.notes == RELATION_NOTE and r.match_status == 'Relationship' and not r.final
    assert eligible(r)


def test_csv_strings_and_invalid_files():
    frame = parse_file(b'ID,Value\n00042,N/A\n', 'test.csv')
    assert frame['ID'][0] == '00042' and frame['Value'][0] == 'N/A'
    for data, name in [(b'', 'test.csv'), (b'ID,Value\n', 'test.csv'), (b'hello', 'test.zip'), (b'broken', 'test.xlsx')]:
        with pytest.raises(ValueError):
            parse_file(data, name)


def test_xlsx_identifier_number_format():
    book = Workbook()
    sheet = book.active
    sheet.append(['ID', 'Value'])
    sheet.append([42, 'N/A'])
    sheet['A2'].number_format = '00000'
    output = BytesIO()
    book.save(output)
    frame = parse_file(output.getvalue(), 'test.xlsx')
    assert frame['ID'][0] == '00042'


def test_identifier_stats():
    frame = pl.DataFrame({'ID': ['0001', '0001', '', None]})
    assert identifier_stats(frame, 'ID') == {'Missing identifiers': 2, 'Rows with duplicate identifiers': 2}


@pytest.fixture
def storage(tmp_path):
    store = Store(tmp_path / 'test.sqlite3')
    project, token = store.create()
    return store, project, token


def test_checkpoint_token_and_project_isolation(storage):
    store, project, token = storage
    assert store.authenticate(project, token)
    assert not store.authenticate(project, 'wrong')
    store.save(project, {'results': [Result(0, 'Alpha', 'SMB').dict()]})
    assert store.load(project)['results'][0]['identifier'] == 'Alpha'
    other, _ = store.create()
    store.put_cache(project, 'k', {'x': 1})
    assert store.get_cache(project, 'k') == {'x': 1}
    assert store.get_cache(other, 'k') is None
    store.put_cache(project, 'expired', {}, ttl_days=-1)
    assert store.get_cache(project, 'expired') is None


class Response:
    def __init__(self, payload=None, status=200):
        self.payload, self.status_code = payload, status

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError('Provider failure')

    def json(self):
        return self.payload


class Transport:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.calls = []

    def post(self, url, **kwargs):
        self.calls.append((url, kwargs))
        value = next(self.responses)
        if isinstance(value, Exception):
            raise value
        return value


def research_fixture(storage, selected=0, relationship='identity', urls=None):
    store, project, _ = storage
    source = pl.DataFrame({'Company': ['Alpha'], 'Value': ['Enterprise']})
    r = Result(0, 'Old Alpha', 'Verify', candidates=[{'source_row': 0, 'account': 'Alpha', 'value': 'Enterprise', 'similarity': 70}])
    transport = Transport([
        Response({'results': [{'url': 'https://example.com/evidence', 'content': 'Old Alpha renamed Alpha'}]}),
        Response({'choices': [{'message': {'content': json.dumps({'source_row': selected,
                 'relationship': {'identity': 'Direct', 'parent': 'Subsidiary'}.get(relationship, 'Unknown'),
                 'recommended_source_account': 'Alpha', 'proposed_value': 'Enterprise', 'parent_company': '',
                 'confidence': 'High', 'recommendation_status': 'Proposed', 'reasoning': 'Verified company identity.',
                 'identity_supported': True, 'inheritance_supported': False, 'inheritance_quote': '',
                 'supporting_quote': 'Old Alpha renamed Alpha',
                 'explanation': 'Company renamed.', 'evidence_urls': urls if urls is not None else ['https://example.com/evidence']})}}]})])
    researcher = Researcher(store, project, 't', 'o', ResearchSettings(interval=0), transport, lambda _: None)
    return r, source, transport, researcher


def test_research_cache_and_pending(storage):
    r, source, transport, researcher = research_fixture(storage)
    researcher.investigate(r, source, M, 'dataset-one')
    assert r.research_status == 'Completed' and r.proposed == 'Enterprise' and not r.final
    assert r.evidence == ['https://example.com/evidence']
    r.research_status = 'Not researched'
    researcher.investigate(r, source, M, 'dataset-one')
    assert r.research_status == 'Completed' and r.research_cache_hit and len(transport.calls) == 2
    assert researcher.cached(r, 'dataset-two') is None


@pytest.mark.parametrize('selected,urls', [(999, None), (0, ['https://invented.example'])])
def test_invented_candidates_or_evidence_rejected(storage, selected, urls):
    r, source, _, researcher = research_fixture(storage, selected, urls=urls)
    researcher.investigate(r, source, M, 'ctx')
    assert not r.proposed and not r.final


def test_research_parent_is_pending(storage):
    r, source, _, researcher = research_fixture(storage, relationship='parent')
    researcher.investigate(r, source, M, 'ctx')
    assert r.notes == RELATION_NOTE and r.review_status == 'Pending' and not r.final


def test_api_429_backoff_and_failure(storage):
    store, project, _ = storage
    delays = []
    transport = Transport([Response(status=429), Response(status=429), Response(status=429)])
    researcher = Researcher(store, project, 't', 'o', ResearchSettings(interval=0, retries=2), transport, delays.append)
    r = Result(0, 'Alpha', 'Verify')
    researcher.investigate(r, pl.DataFrame({'Company': ['Alpha'], 'Value': ['SMB']}), M, 'ctx')
    assert r.research_status.startswith('Failed') and not r.final
    assert len(transport.calls) == 3 and 1 in delays and 2 in delays


def test_timeout_and_offline(storage):
    store, project, _ = storage
    transport = Transport([requests.Timeout(), requests.Timeout()])
    researcher = Researcher(store, project, 't', 'o', ResearchSettings(interval=0, retries=1), transport, lambda _: None)
    r = Result(0, 'Alpha', 'Verify')
    source = pl.DataFrame({'Company': ['Alpha'], 'Value': ['SMB']})
    researcher.investigate(r, source, M, 'ctx')
    assert r.research_status.startswith('Failed')
    Researcher(store, project).investigate(r, source, M, 'ctx')
    assert r.research_status == 'Unavailable (offline)'


def test_reliable_matches_and_decisions_skip_research(storage):
    r, source, transport, researcher = research_fixture(storage)
    r.identity_confidence = 'Reliable local match'
    researcher.investigate(r, source, M, 'ctx')
    assert not transport.calls
    r.identity_confidence = 'Unverified'
    decide(r, 'Manual value', 'Reviewer', manual_value='Verified segment')
    researcher.investigate(r, source, M, 'ctx')
    assert not transport.calls and r.final == 'Verified segment'


def test_transactional_review_audit(storage):
    store, project, _ = storage
    r = Result(0, 'Alpha', 'Verify', proposed='Enterprise')
    event = decide(r, 'Approve proposal', 'Reviewer')
    store.commit_review(project, {'results': [r.dict()]}, event)
    assert store.load(project)['results'][0]['final'] == 'Enterprise'
    assert store.audit(project)[0]['reviewer'] == 'Reviewer'


def test_excel_integrity_and_formula_safety():
    r = Result(0, '=HYPERLINK("bad")', 'Verify', proposed='Enterprise', final='Enterprise',
               evidence=['https://example.com/evidence'])
    book = load_workbook(BytesIO(workbook([r], [], [], 100)))
    assert book.sheetnames == ['Reconciliation Results', 'Summary', 'Audit Trail', 'Source Conflicts']
    sheet = book['Reconciliation Results']
    columns = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, columns['Final Value']).value is None  # unapproved suggestion cannot leak to final
    assert sheet['A2'].data_type == 's'
    assert sheet.freeze_panes == 'A2' and sheet.auto_filter.ref
    assert sheet.cell(2, columns['Evidence URL 1']).hyperlink.target == 'https://example.com/evidence'
    decide(r, 'Approve proposal', 'Reviewer')
    book = load_workbook(BytesIO(workbook([r], [], [], 100)))
    assert book['Reconciliation Results'].cell(2, columns['Final Value']).value == 'Enterprise'
