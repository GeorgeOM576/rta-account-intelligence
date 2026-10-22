from io import BytesIO
from pathlib import Path

import polars as pl
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from rta.ingest import parse_file
from rta.matching import Matcher
from rta.models import Mapping, Result
from rta.review import decide, final_value, manual_options
from rta.storage import Store, digest
from rta.export import workbook, summary
from tests.test_app import by_label, navigate


@pytest.fixture
def workspace(tmp_path):
    store = Store(tmp_path / 'reviews.sqlite3')
    project, token = store.create()
    master = parse_file(Path('samples/master.csv').read_bytes(), 'master.csv')
    source = parse_file(Path('samples/source.csv').read_bytes(), 'source.csv')
    mapping = Mapping('Account', 'Company', 'Segment', 'Customer Segment')
    engine = Matcher(source, mapping)
    results = engine.run(master)
    payload = {'master': master.to_dicts(), 'source': source.to_dicts(),
               'master_schema': dict.fromkeys(master.columns, 'String'),
               'source_schema': dict.fromkeys(source.columns, 'String'),
               'master_name': 'master.csv', 'source_name': 'source.csv',
               'mapping': vars(mapping), 'results': [r.dict() for r in results],
               'context': digest(source.to_dicts()), 'conflicts': engine.conflicts}
    store.save(project, payload)
    return store, project, token


def record(store, project, row=1):
    return Result(**next(r for r in store.load(project)['results'] if r['row_id'] == row))


def test_manual_options_include_suggestions_exclude_placeholders():
    source = pl.DataFrame({'Segment': ['SMB', '', None, ' Verify ', 'N / A', 'Enterprise', 'SMB', ' smb ']})
    assert manual_options(source, 'Segment', [{'value': 'Enterprise'}, {'value': 'SMB'}]) == ['Enterprise', 'SMB']


@pytest.mark.parametrize('custom,expected,origin', [('', 'Healthcare', 'source dropdown'),
                                                  ('Custom verified segment', 'Custom verified segment', 'custom entry')])
def test_manual_decision_persists_and_exports(workspace, custom, expected, origin):
    store, project, _ = workspace
    store.review(project, 1, 'Manual', 'Reviewer', 'Verified account-specific segment', custom, 'Healthcare')
    reopened = Store(store.path)
    r = record(reopened, project)
    assert (r.final, r.review_status, r.decision, r.value_origin) == (expected, 'Manually Verified', 'Manual', origin)
    book = load_workbook(BytesIO(workbook([r], reopened.audit(project), [], 10)))
    sheet = book['Reconciliation Results']
    values = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert values['Final Value'] == expected
    assert values['Review Status'] == 'Manually Verified'
    assert values['Reviewer Decision'] == 'Manual'
    assert values['Reviewer Notes'] == 'Verified account-specific segment'
    assert values['Value Origin'] == origin
    assert len(sheet.conditional_formatting) > 0


@pytest.mark.parametrize('placeholder', ['', 'Verify', ' UNMAPPED ', 'Duplicate', 'N / A', 'Null', 'None'])
def test_invalid_manual_entries_do_not_commit(workspace, placeholder):
    store, project, _ = workspace
    # A nonempty placeholder custom entry cannot fall back to a valid dropdown.
    with pytest.raises(ValueError):
        store.review(project, 1, 'Manual', 'Reviewer', manual_value=placeholder,
                     source_value='Healthcare' if placeholder.strip() else '')
    assert not record(store, project).final
    assert store.audit(project) == []


def test_manual_dropdown_rejects_values_not_in_source(workspace):
    store, project, _ = workspace
    with pytest.raises(ValueError):
        store.review(project, 1, 'Manual', 'Reviewer', source_value='Invented dropdown value')


def test_approve_reject_pending_and_audit(workspace):
    store, project, _ = workspace
    store.review(project, 1, 'Approve', 'Reviewer', 'Accepted')
    r = record(store, project)
    assert r.final == 'Healthcare' and r.review_status == 'Approved'
    store.review(project, 1, 'Reject', 'Reviewer', 'New evidence')
    assert not record(store, project).final and record(store, project).review_status == 'Rejected'
    events = store.audit(project)
    assert events[1]['previous_final'] == 'Healthcare' and events[1]['new_final'] == ''
    assert events[1]['row_id'] == 1 and events[1]['review_status'] == 'Rejected'
    assert events[1]['timestamp'] and events[1]['notes'] == 'New evidence'
    store.review(project, 1, 'Pending', 'Reviewer')
    assert not record(store, project).final and record(store, project).status_label == 'Intelligence Proposed — Pending Approval'
    # A human-approved original can survive rejection; returning to pending clears Final.
    store.review(project, 0, 'Approve', 'Reviewer')
    for action, expected in [('Reject', 'Enterprise'), ('Pending', '')]:
        store.review(project, 0, action, 'Reviewer')
        assert final_value(record(store, project, 0)) == expected


def test_stale_checkpoint_cannot_overwrite_decision(workspace):
    store, project, _ = workspace
    stale = store.load(project)
    store.review(project, 1, 'Manual', 'Reviewer', manual_value='Custom segment')
    store.save(project, stale)
    assert record(store, project).final == 'Custom segment'
    assert stale['results'][1]['final'] == 'Custom segment'
    assert summary([record(store, project)])['Manually Verified'] == 1


def test_rejection_preserves_explicitly_verified_original(workspace):
    store, project, _ = workspace
    # Summit has conflicting Source values and a valid original Commercial value.
    store.review(project, 3, 'Retain original', 'Reviewer', 'Independently verified original')
    store.review(project, 3, 'Reject', 'Reviewer', 'Reject proposed replacement')
    r = record(store, project, 3)
    assert r.review_status == 'Rejected' and final_value(r) == 'Commercial'


def test_duplicate_names_use_stable_row_ids(workspace):
    store, project, _ = workspace
    payload = store.load(project)
    payload['results'][2]['identifier'] = payload['results'][1]['identifier']
    store.save(project, payload)
    store.review(project, 2, 'Manual', 'Reviewer', manual_value='Only second duplicate')
    assert not record(store, project, 1).final
    assert record(store, project, 2).final == 'Only second duplicate'


def test_migrate_legacy_manual_result():
    r = Result(5, 'Account', 'Verify', final='Enterprise', decision='Manual value', review_status='Approved')
    assert r.decision == 'Manual' and r.review_status == 'Manually Verified'
    assert final_value(r) == 'Enterprise'


def app_for(workspace, monkeypatch):
    store, project, token = workspace
    monkeypatch.setenv('RTA_DB_PATH', store.path)
    app = AppTest.from_file(Path('app.py').resolve(), default_timeout=30)
    app.session_state['project'], app.session_state['token'] = project, token
    return app.run()


@pytest.mark.parametrize('action,custom,expected,status', [
    ('Manual', '', 'Healthcare', 'Manually Verified'),
    ('Manual', 'Custom UI value', 'Custom UI value', 'Manually Verified'),
    ('Approve', '', 'Healthcare', 'Approved'),
    ('Reject', '', '', 'Rejected')])
def test_ui_immediate_updates_reload_and_export(workspace, monkeypatch, action, custom, expected, status):
    app = app_for(workspace, monkeypatch)
    navigate(app, 'Review & Approvals')
    assert by_label(app.selectbox, 'Account to review').value == 1
    by_label(app.selectbox, 'Decision').set_value(action).run()
    if action == 'Manual':
        dropdown = by_label(app.selectbox, 'Verified Source value')
        assert 'Healthcare' in dropdown.options and 'Enterprise' in dropdown.options
        dropdown.set_value('Healthcare')
        by_label(app.text_input, 'Custom verified value').set_value(custom)
    by_label(app.text_input, 'Reviewer name').set_value('UI Reviewer')
    by_label(app.text_area, 'Reviewer notes').set_value('UI evidence checked')
    by_label(app.checkbox, 'I have verified the evidence and the field value for this specific account.').check()
    by_label(app.button, 'Save review decision').click().run()
    assert not app.exception
    assert by_label(app.selectbox, 'Account to review').value == 1
    assert by_label(app.metric, 'Final value').value == (expected or '—')
    assert by_label(app.metric, 'Manually Verified').value == ('1' if action == 'Manual' else '0')
    assert by_label(app.metric, 'Pending Approvals').value == '9'
    assert any('Review decision saved' in s.value for s in app.success)
    assert any(status in m.value for m in app.markdown)
    assert app.dataframe[-1].value.iloc[1]['Final Value'] == expected
    assert app.dataframe[-1].value.iloc[1]['Review Status'] == status
    store, project, _ = workspace
    assert record(store, project).final == expected
    # A fresh Streamlit session loads only the saved project identity and SQLite state.
    reloaded = app_for(workspace, monkeypatch)
    navigate(reloaded, 'Review & Approvals')
    by_label(reloaded.selectbox, 'Review status').set_value('All').run()
    by_label(reloaded.selectbox, 'Account to review').set_value(1).run()
    assert by_label(reloaded.metric, 'Final value').value == (expected or '—')
    assert any(status in m.value for m in reloaded.markdown)
    # Capture the actual workbook produced by Export Results, not a reconstructed mock.
    import rta.export
    original = rta.export.workbook
    exported = []
    def capture(*args, **kwargs):
        data = original(*args, **kwargs)
        exported.append(data)
        return data
    monkeypatch.setattr(rta.export, 'workbook', capture)
    navigate(reloaded, 'Export Results')
    sheet = load_workbook(BytesIO(exported[-1]))['Reconciliation Results']
    headers = [c.value for c in sheet[1]]
    row = dict(zip(headers, [c.value for c in sheet[3]]))
    assert (row['Final Value'] or '') == expected and row['Review Status'] == status
    assert row['Reviewer Decision'] == action
    assert row['Reviewer Notes'] == 'UI evidence checked'
