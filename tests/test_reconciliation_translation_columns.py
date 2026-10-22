"""Exercise the rendered Streamlit reconciliation table, not just result serialization."""
import pytest

from tests.test_approved_translation_integration import fixture, approved
from tests.test_review_workflow import app_for
from tests.test_app import navigate, by_label
from tests.workflow_display_helpers import csv_columns, rendered_table, visible_and_backend_row

FRONT = ['Master Identifier', 'Master Approved English Translation', 'Master Translation Status',
         'Matched Source Approved English Translation', 'Matched Source Translation Status', 'Match Basis']


def test_actual_streamlit_table_shows_csv_schema_and_preserves_backend_statuses(tmp_path, monkeypatch):
    case, engine = fixture(tmp_path,
        [{'Name': 'Les Jardins de Provence', 'Segment': 'Verify'}],
        [{'Name': 'Jardines de Provenza', 'Segment': 'SMB'}])
    approved(engine, 'master', 'The Gardens of Provence')
    approved(engine, 'source', 'The Gardens of Provence')
    app = app_for(case, monkeypatch)
    navigate(app, 'Run Reconciliation')
    by_label(app.button, 'Run automatic local matching').click().run()
    assert not app.exception
    table = app.dataframe[-1].value
    assert list(table.columns) == csv_columns()
    row = visible_and_backend_row(app)
    assert row['Master Approved English Translation'] == 'The Gardens of Provence'
    assert row['Matched Source Approved English Translation'] == 'The Gardens of Provence'
    assert row['Master Translation Status'] == row['Matched Source Translation Status'] == 'Approved Translation'
    assert row['Match Basis'] == 'Master Translation ↔ Source Translation'
    assert row['Proposed Value'] == 'SMB' and row['Final Value'] == ''
    assert row['Suggested Match Basis 1'] == row['Match Basis']
    assert {'Master Original Value', 'Proposed Value', 'Final Value', 'Match Status', 'Review Status',
            'Suggested Account 1', 'Suggested Account 2', 'Suggested Account 3'} <= set(table.columns)


@pytest.mark.parametrize('side', ['master', 'source'])
@pytest.mark.parametrize('status', ['Translation Needed', 'Translated — Review', 'Rejected Translation'])
def test_actual_table_separates_status_from_approved_text(tmp_path, monkeypatch, side, status):
    from tests.test_translation import Offline
    from tests.translation_workspace_helpers import seed_google
    names = {'master': 'Bank of Mexico', 'source': 'Bank of Mexico'}
    names[side] = 'Banco de México'
    case, engine = fixture(tmp_path, [{'Name': names['master'], 'Segment': 'Verify'}],
                           [{'Name': names['source'], 'Segment': 'SMB'}])
    seed_google(engine, {'Banco de México': 'es', 'Bank of Mexico': 'en'})
    if status != 'Translation Needed':
        engine.process([side], Offline())
    if status == 'Rejected Translation':
        engine.reject(next(r for r in engine.records() if r['dataset'] == side), 'Reviewer', True)
    app = app_for(case, monkeypatch)
    navigate(app, 'Run Reconciliation')
    by_label(app.button, 'Run automatic local matching').click().run()
    assert not app.exception
    row = visible_and_backend_row(app)
    prefix = 'Master' if side == 'master' else 'Matched Source'
    # Unmatched fuzzy candidates are suggestions, not matched Source accounts.
    assert row[f'{prefix} Translation Status'] == (status if side == 'master' else '')
    if side == 'source':
        assert row['Suggested Account 1'] == 'Banco de México'
    assert row[f'{prefix} Approved English Translation'] == ''
    assert row['Proposed Value'] == row['Final Value'] == ''
    assert 'Translation' not in row['Match Basis']
    assert all('Translation' not in row[f'Suggested Match Basis {i}'] for i in range(1, 4))


def test_ui_approval_updates_columns_and_requires_safe_matching_rerun(tmp_path, monkeypatch):
    from tests.test_translation import Offline
    from tests.translation_workspace_helpers import seed_google
    case, engine = fixture(tmp_path, [{'Name': 'Banco de México', 'Segment': 'Verify'}],
                           [{'Name': 'Bank of Mexico', 'Segment': 'SMB'}])
    seed_google(engine, {'Banco de México': 'es', 'Bank of Mexico': 'en'})
    engine.process(['master'], Offline())
    app = app_for(case, monkeypatch)
    navigate(app, 'Run Reconciliation')
    by_label(app.button, 'Run automatic local matching').click().run()
    assert not any('results are stale' in w.value for w in app.warning)
    assert app.dataframe[-1].value['Master Approved English Translation'].iloc[0] == ''
    navigate(app, 'Translation Intelligence')
    by_label(app.selectbox, 'Translation dataset').set_value('Master').run()
    by_label(app.selectbox, 'Translation status').set_value('Translated — Review').run()
    by_label(app.text_input, 'Individual reviewer').set_value('Translation reviewer')
    by_label(app.checkbox, 'I verified this translation preserves the correct entity, proper names, legal suffixes and identifiers.').check()
    by_label(app.button, 'Approve translation').click().run()
    assert not app.exception
    navigate(app, 'Run Reconciliation')
    row = visible_and_backend_row(app)
    assert row['Master Approved English Translation'] == 'Bank of Mexico'
    assert row['Master Translation Status'] == 'Approved Translation'
    assert 'Translation' not in row['Match Basis']  # prior run never used the newly approved name
    assert any('results are stale' in w.value for w in app.warning)
    audit_before = case[0].audit(case[1])
    by_label(app.button, 'Run automatic local matching').click().run()
    row = visible_and_backend_row(app)
    assert row['Match Basis'] == row['Suggested Match Basis 1'] == 'Master Translation ↔ Source Original'
    assert row['Proposed Value'] == 'SMB' and row['Final Value'] == ''
    assert not any('results are stale' in w.value for w in app.warning)
    assert case[0].audit(case[1]) == audit_before
    app.run()  # signature survives the next normal Streamlit rerun
    assert not any('results are stale' in w.value for w in app.warning)


def test_research_page_uses_current_sqlite_approvals_not_saved_context(tmp_path, monkeypatch):
    from copy import deepcopy
    case, engine = fixture(tmp_path, [{'Name': 'Les Jardins de Provence', 'Segment': 'Verify'}],
                           [{'Name': 'The Gardens of Provence', 'Segment': 'SMB'}])
    monkeypatch.setenv('TAVILY_API_KEY', 'mock-only')
    monkeypatch.setenv('OPENAI_API_KEY', 'mock-only')
    seen = []
    def capture(self, result, *args, **kwargs):
        seen.append(deepcopy(result.account_context))
        return result
    monkeypatch.setattr('rta.research.Researcher.investigate', capture)
    app = app_for(case, monkeypatch)
    navigate(app, 'Run Reconciliation')
    by_label(app.button, 'Run automatic local matching').click().run()
    approved(engine, 'master', 'The Gardens of Provence')
    # Do not rerun reconciliation: the saved context still lacks this approval.
    assert not case[0].load(case[1])['results'][0]['account_context']['master']['approved_english']
    navigate(app, 'AI Research')
    assert by_label(app.button, 'Start / resume paid research').disabled
    assert not seen
    by_label(app.checkbox, 'I authorize external research for this batch and the associated API charges.').check().run()
    by_label(app.button, 'Start / resume paid research').click().run()
    assert not app.exception and len(seen) == 1
    assert seen[0]['master']['original_name'] == 'Les Jardins de Provence'
    assert seen[0]['master']['approved_english'] == 'The Gardens of Provence'


def test_excel_statuses_and_new_columns_match_streamlit(tmp_path, monkeypatch):
    from io import BytesIO
    from openpyxl import load_workbook
    from rta.export import workbook
    from rta.models import Result
    case, engine = fixture(tmp_path, [{'Name': 'Les Jardins de Provence', 'Segment': 'Verify'}],
                           [{'Name': 'Jardines de Provenza', 'Segment': 'SMB'}])
    approved(engine, 'master', 'The Gardens of Provence')
    approved(engine, 'source', 'The Gardens of Provence')
    app = app_for(case, monkeypatch)
    navigate(app, 'Run Reconciliation')
    by_label(app.button, 'Run automatic local matching').click().run()
    visible = visible_and_backend_row(app)
    payload = case[0].load(case[1])
    results = [Result(**r) for r in payload['results']]
    book = load_workbook(BytesIO(workbook(results, case[0].audit(case[1]), [], 1, engine.records(), payload)))
    sheet = book['Reconciliation Results']
    headers = [c.value for c in sheet[1]]
    exported = dict(zip(headers, [c.value for c in sheet[2]]))
    assert headers[:6] == FRONT
    for column in FRONT:
        assert exported[column] == visible[column]
    assert exported['Proposed Value'] == 'SMB' and exported['Final Value'] is None
    assert sheet.freeze_panes == 'A2' and sheet.auto_filter.ref
    assert {'Master Translations', 'Source Translations', 'Audit Trail'} <= set(book.sheetnames)


def test_legacy_results_show_no_invented_basis_and_request_rerun(tmp_path, monkeypatch):
    from rta.models import Result
    case, engine = fixture(tmp_path, [{'Name': 'Les Jardins de Provence', 'Segment': 'Verify'}],
                           [{'Name': 'The Gardens of Provence', 'Segment': 'SMB'}])
    approved(engine, 'master', 'The Gardens of Provence')
    payload = case[0].load(case[1])
    payload['results'] = [Result(0, 'Les Jardins de Provence', 'Verify',
        candidates=[{'source_row': 0, 'account': 'The Gardens of Provence', 'value': 'SMB', 'similarity': 20}]).dict()]
    case[0].save(case[1], payload)
    app = app_for(case, monkeypatch)
    navigate(app, 'Run Reconciliation')
    row = visible_and_backend_row(app)
    assert row['Master Approved English Translation'] == 'The Gardens of Provence'
    assert row['Match Basis'] == row['Suggested Match Basis 1'] == ''
    assert any('results are stale' in w.value for w in app.warning)
    by_label(app.button, 'Run automatic local matching').click().run()
    assert visible_and_backend_row(app)['Match Basis'] == 'Master Translation ↔ Source Original'
    assert not any('results are stale' in w.value for w in app.warning)
