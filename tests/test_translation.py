from copy import deepcopy
from io import BytesIO
import json

import polars as pl
import pytest
from openpyxl import load_workbook

from rta.models import Mapping, Result, INTELLIGENCE_PENDING
from rta.review import bulk_approval_error, final_value
from rta.storage import Store, digest
from rta.translation import TranslationEngine, detect_text, translation_metrics, translated_rows
from rta.translation_matching import reconcile
from rta.translation_providers import OpenAIProvider, ArgosProvider
from rta.export import workbook
from tests.test_app import by_label, navigate
from tests.test_review_workflow import app_for

SPANISH = 'Servicios financieros para empresas mexicanas'
ENGLISH = 'Financial services for Mexican businesses'
MAPPING = Mapping('Name', 'Name', 'Segment', 'Segment')


def make_workspace(tmp_path, master=None, source=None, filename='translation.sqlite3'):
    store = Store(tmp_path / filename)
    project, token = store.create()
    master = master or [{'Name': 'Banco de México', 'Segment': 'Verify', 'Description': SPANISH}]
    source = source or [{'Name': 'Bank of Mexico', 'Segment': 'SMB', 'Description': ENGLISH}]
    payload = {'master': master, 'source': source, 'mapping': vars(MAPPING), 'master_name': 'master.csv',
        'language_detection_provider': 'offline',  # This fixture explicitly exercises the retained local fallback.
        'source_name': 'source.csv', 'master_schema': dict.fromkeys(master[0], 'String'),
        'source_schema': dict.fromkeys(source[0], 'String'), 'context': digest(source), 'results': []}
    store.save(project, payload)
    return store, project, token


class Offline:
    name, model, languages, available = 'Fixture offline', 'test-model-1', ['es', 'ja', 'fr'], True

    def __init__(self, fail=False):
        self.calls = []
        self.fail = fail

    def translate(self, text, language):
        self.calls.append((text, language))
        if self.fail:
            raise ValueError('Do not expose this confidential provider error')
        return {'Banco de México': 'Bank of Mexico', SPANISH: ENGLISH,
                'メキシコ銀行': 'Bank of Mexico', '株式会社トヨタ': 'Toyota Corporation'}.get(text, 'Reviewed English draft')


@pytest.mark.parametrize('text,status,language', [
    (ENGLISH, 'Not Required', 'en'), (SPANISH, 'Translation Needed', 'es'),
    ('Service commercial pour les entreprises', 'Translation Needed', 'fr'),
    ('株式会社トヨタ', 'Translation Needed', 'ja'), ('Hello 世界', 'Uncertain', None),
    ('Café International', 'Uncertain', None), ('Microsoft Corporation', 'Uncertain', None),
    ('Acme', 'Uncertain', None), ('00042', 'Not Required', None), ('ABC-42', 'Not Required', None),
    ('abc42', 'Not Required', None), ('IBM', 'Not Required', None), ('2026-10-21', 'Not Required', None),
    ('21 Oct 2026', 'Not Required', None), ('https://example.com/es', 'Not Required', None),
    ('', 'Not Required', None), ('   ', 'Not Required', None),
])
def test_actual_offline_detection(text, status, language):
    detection = detect_text(text)
    assert detection['status'] == status
    if language:
        assert detection['language'] == language


@pytest.mark.parametrize('master_foreign,source_foreign', [(True, False), (False, True), (True, True), (False, False)])
def test_independent_dataset_detection(tmp_path, master_foreign, source_foreign):
    master = [{'Name': 'Master Corporation', 'Segment': 'SMB', 'Description': SPANISH if master_foreign else ENGLISH}]
    source = [{'Name': 'Source Corporation', 'Segment': 'SMB', 'Description': SPANISH if source_foreign else ENGLISH}]
    store, project, _ = make_workspace(tmp_path, master, source)
    engine = TranslationEngine(store, project)
    rows = engine.detect({'master': ['Description'], 'source': ['Description']})
    counts = translation_metrics(rows)
    assert counts[0]['Translation Needed'] == master_foreign
    assert counts[1]['Translation Needed'] == source_foreign
    assert counts[2]['Translation Needed'] == master_foreign + source_foreign
    assert store.load(project)['master'] == master and store.load(project)['source'] == source


def test_shared_draft_cache_never_shares_approvals_and_is_project_scoped(tmp_path):
    rows = [{'Name': 'Name', 'Segment': 'SMB', 'Description': SPANISH}] * 2
    store, project, _ = make_workspace(tmp_path, rows, rows[:1])
    engine = TranslationEngine(store, project)
    original = deepcopy(store.load(project))
    engine.detect({'master': ['Description'], 'source': ['Description']})
    provider = Offline()
    report = engine.process(['master', 'source'], provider)
    assert report['drafts'] == 3 and len(provider.calls) == 1
    records = engine.records()
    assert sum(r['cache_hit'] for r in records) == 2
    master = next(r for r in records if r['dataset'] == 'master' and r['row_id'] == 0)
    engine.approve(master, 'Corrected English', 'Reviewer', True)
    assert next(r for r in engine.records() if r['dataset'] == 'source')['english'] == ENGLISH
    assert next(r for r in engine.records() if r['dataset'] == 'source')['status'] == 'Translated — Review'
    engine.detect({'master': ['Description'], 'source': ['Description']})
    engine.process(['master', 'source'], provider)
    assert len(provider.calls) == 1
    assert store.load(project)['master'] == original['master'] and store.load(project)['source'] == original['source']
    # A different tenant/workspace never reads the first workspace's draft or approval.
    other, _ = store.create()
    store.save(other, original)
    other_engine = TranslationEngine(store, other)
    other_engine.detect({'master': ['Description'], 'source': ['Description']})
    other_engine.process(['master', 'source'], provider)
    assert len(provider.calls) == 2
    assert all(r['status'] == 'Translated — Review' for r in other_engine.records())
    assert len(store.audit(project)) == 1 and not store.audit(other)


def test_cached_draft_is_reused_in_a_later_batch(tmp_path):
    store, project, _ = make_workspace(tmp_path, source=[{'Name': 'Bank', 'Segment': 'SMB', 'Description': SPANISH}])
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Description'], 'source': ['Description']})
    provider = Offline()
    engine.process(['master'], provider)
    assert engine.records()[1]['status'] == 'Translation Needed'
    engine.process(['source'], provider)
    assert len(provider.calls) == 1 and engine.records()[1]['cache_hit']


def test_failure_retry_name_protection_and_uncertainty(tmp_path):
    store, project, _ = make_workspace(tmp_path)
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Name', 'Description'], 'source': ['Name', 'Description']})
    failed = Offline(fail=True)
    engine.process(['master'], failed, allow_names=False)
    name = next(r for r in engine.records() if r['dataset'] == 'master' and r['field'] == 'Name')
    assert name['status'] == 'Translation Needed' and not any(call[0] == name['original'] for call in failed.calls)
    description = next(r for r in engine.records() if r['dataset'] == 'master' and r['field'] == 'Description')
    assert description['status'] == 'Failed' and 'confidential' not in description['reason']
    working = Offline()
    engine.process(['master'], working, allow_names=True, include_uncertain=True)
    assert all(r['status'] == 'Translated — Review' for r in engine.records() if r['dataset'] == 'master')


def test_unavailable_offline_runtime_and_unsupported_models(monkeypatch):
    def info(self, payload):
        assert payload['action'] == 'info'
        return {'model': 'fixture-installed-es-1', 'languages': ['es']}
    monkeypatch.setattr(ArgosProvider, '_call', info)
    provider = ArgosProvider()
    assert provider.available and provider.languages == ['es']
    with pytest.raises(ValueError, match='model'):
        provider.translate('Unknown language', 'xx')
    def missing(self, payload):
        raise ValueError('Missing runtime')
    monkeypatch.setattr(ArgosProvider, '_call', missing)
    unavailable = ArgosProvider()
    assert not unavailable.available
    with pytest.raises(ValueError):
        unavailable.translate(SPANISH, 'es')


def test_model_versions_do_not_share_draft_entries(tmp_path):
    store, project, _ = make_workspace(tmp_path, source=[{'Name': 'Source', 'Segment': 'SMB', 'Description': SPANISH}])
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Description'], 'source': ['Description']})
    first, second = Offline(), Offline()
    second.model = 'test-model-2'
    engine.process(['master'], first)
    engine.process(['source'], second)
    assert len(first.calls) == len(second.calls) == 1
    assert next(r for r in engine.records() if r['dataset'] == 'source')['model'] == 'test-model-2'


def test_identifiers_cannot_be_silently_removed_by_providers(tmp_path):
    store, project, _ = make_workspace(tmp_path, master=[{'Name': 'Name', 'Segment': 'SMB', 'Description': SPANISH + ' ID42'}])
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Description'], 'source': []})
    engine.process(['master'], Offline())
    assert engine.records()[0]['status'] == 'Failed'
    assert store.load(project)['master'][0]['Description'].endswith('ID42')


def test_comparison_field_translations_are_context_not_invented_source_values(tmp_path):
    store, project, _ = make_workspace(tmp_path,
        [{'Name': 'Banco de México', 'Segment': 'Commercial', 'Description': ENGLISH}],
        [{'Name': 'Bank of Mexico', 'Segment': 'Comercial', 'Description': ENGLISH}])
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Name', 'Segment'], 'source': ['Name', 'Segment']})
    engine.configure_detection({'enabled_classifications': {'master': ['Segment'], 'source': ['Segment']}})
    for entry in engine.records():
        if entry['dataset'] == 'master' and entry['field'] == 'Name':
            engine.approve(entry, 'Bank of Mexico', 'R', True)
        if entry['dataset'] == 'source' and entry['field'] == 'Segment':
            engine.approve(entry, 'Commercial', 'R', True, 'es')
    payload = store.load(project)
    r = reconcile(pl.DataFrame(payload['master']), pl.DataFrame(payload['source']), MAPPING, engine.records())[0][0]
    assert r.proposed == 'Comercial' and r.final == ''
    checks = r.translation_provenance[0]['comparison_checks']
    assert any(c['direction'] == 'Original Master → Translated Source' and c['agrees'] for c in checks)
    assert any(c['field'] == 'Segment' for c in r.translation_provenance[0]['comparison_context'])


def test_translation_approvals_validate_stale_cells_and_preserve_history(tmp_path):
    store, project, _ = make_workspace(tmp_path)
    engine = TranslationEngine(store, project)
    entry = engine.detect({'master': ['Name'], 'source': []})[0]
    with pytest.raises(ValueError):
        engine.approve(entry, 'Bank of Mexico', '', True)
    with pytest.raises(ValueError):
        engine.approve(entry, 'Bank of Mexico', 'R', False)
    approved = engine.approve(entry, 'Bank of Mexico', 'First', True)
    with pytest.raises(ValueError, match='changed'):
        engine.approve(entry, 'Stale correction', 'Second', True)
    corrected = engine.approve(approved, 'Bank of Mexico Corporation', 'Second', True)
    events = store.audit(project)
    assert len(events) == 2 and events[1]['previous_translation']['english'] == 'Bank of Mexico'
    assert corrected['reviewer'] == 'Second' and corrected['reviewed_at']
    payload = store.load(project)
    payload['master'][0]['Name'] = 'New dataset identity'
    store.save(project, payload)
    assert engine.records() == []
    with pytest.raises(ValueError, match='changed'):
        engine.approve(corrected, 'Old entity', 'R', True)
    assert len(store.audit(project)) == 2


@pytest.mark.parametrize('master_name,source_name,master_english,source_english,direction', [
    ('Bank of Mexico', 'Bank of Mexico', None, None, 'Original Master → Original Source'),
    ('Banco de México', 'Bank of Mexico', 'Bank of Mexico', None, 'Translated Master → Original Source'),
    ('Bank of Mexico', 'Banco de México', None, 'Bank of Mexico', 'Original Master → Translated Source'),
    ('Banco de México', 'メキシコ銀行', 'Bank of Mexico', 'Bank of Mexico', 'Translated Master → Translated Source'),
])
def test_all_comparison_directions_never_auto_approve(tmp_path, master_name, source_name, master_english, source_english, direction):
    store, project, _ = make_workspace(tmp_path,
        [{'Name': master_name, 'Segment': 'SMB', 'Description': ENGLISH}],
        [{'Name': source_name, 'Segment': 'SMB', 'Description': ENGLISH}])
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Name'], 'source': ['Name']})
    for entry in engine.records():
        english = master_english if entry['dataset'] == 'master' else source_english
        if english:
            engine.approve(entry, english, 'Translator', True)
    payload = store.load(project)
    results, _ = reconcile(pl.DataFrame(payload['master']), pl.DataFrame(payload['source']), MAPPING, engine.records())
    r = results[0]
    assert r.proposed == 'SMB' and r.original == 'SMB' and r.final == final_value(r) == ''
    assert r.status_label == INTELLIGENCE_PENDING
    assert direction in [p['direction'] for p in r.translation_provenance]
    if master_english or source_english:
        assert r.identity_confidence != 'Reliable local match'
        assert bulk_approval_error(r, payload)
    payload['results'] = [r.dict()]
    store.save(project, payload)
    store.review(project, 0, 'Approve', 'Account reviewer', 'Identity and Source comparison independently verified')
    saved = Result(**store.load(project)['results'][0])
    assert saved.final == 'SMB' and saved.status_label == 'Approved'
    assert store.audit(project)[-1]['translation_provenance'] == r.translation_provenance


def test_unapproved_or_stale_translations_do_not_match(tmp_path):
    store, project, _ = make_workspace(tmp_path)
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Name'], 'source': ['Name']})
    engine.process(['master'], Offline(), allow_names=True)
    payload = store.load(project)
    master, source = pl.DataFrame(payload['master']), pl.DataFrame(payload['source'])
    assert not reconcile(master, source, MAPPING, engine.records())[0][0].proposed
    entry = next(r for r in engine.records() if r['dataset'] == 'master')
    approved = engine.approve(entry, 'Bank of Mexico', 'R', True)
    assert reconcile(master, source, MAPPING, [approved])[0][0].proposed == 'SMB'
    stale = {**approved, 'revision': 'wrong dataset'}
    assert not reconcile(master, source, MAPPING, [stale])[0][0].proposed


def test_translation_collisions_preserve_conflict_safeguards(tmp_path):
    store, project, _ = make_workspace(tmp_path,
        [{'Name': 'Bank of Mexico', 'Segment': 'Verify', 'Description': ENGLISH}],
        [{'Name': 'Banco de México', 'Segment': 'SMB', 'Description': SPANISH},
         {'Name': 'メキシコ銀行', 'Segment': 'Enterprise', 'Description': ENGLISH}])
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Name'], 'source': ['Name']})
    for item in engine.records():
        if item['dataset'] == 'source':
            engine.approve(item, 'Bank of Mexico', 'R', True)
    payload = store.load(project)
    r = reconcile(pl.DataFrame(payload['master']), pl.DataFrame(payload['source']), MAPPING, engine.records())[0][0]
    assert r.match_status == 'Duplicate conflict' and r.proposed == r.final == ''


def test_export_translation_fields_and_collision_safe_names(tmp_path):
    original = {'Name': 'Banco de México', 'Name (English)': 'Existing uploaded column', 'Segment': 'Verify', 'Description': SPANISH}
    store, project, _ = make_workspace(tmp_path, [original])
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Name'], 'source': ['Name']})
    item = next(r for r in engine.records() if r['dataset'] == 'master')
    engine.approve(item, 'Bank of Mexico', 'R', True)
    payload, records = store.load(project), engine.records()
    rows = translated_rows(payload, 'master', records)
    assert rows[0]['Name (English)'] == 'Existing uploaded column'
    assert rows[0]['Name (English) translation'] == 'Bank of Mexico'
    book = load_workbook(BytesIO(workbook([], store.audit(project), [], 1, translations=records, datasets=payload)))
    assert {'Translation Summary', 'Translation Details', 'Master Translations', 'Source Translations'} <= set(book.sheetnames)
    sheet = book['Master Translations']
    values = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert values['Name'] == 'Banco de México' and values['Name (English) translation'] == 'Bank of Mexico'
    assert values['Name (English) translation status'] == 'Approved Translation'
    assert values['Name (English) translation language'] == 'es'
    assert store.load(project)['master'] == [original]


class API:
    def __init__(self):
        self.calls = []

    def post(self, url, **kwargs):
        self.calls.append((url, kwargs))
        class Response:
            def raise_for_status(self):
                pass
            def json(self):
                return {'choices': [{'message': {'content': json.dumps({'english': ENGLISH})}}]}
        return Response()


@pytest.mark.parametrize('override', [{'consent': False}, {'max_requests': 0}, {'max_tokens': 0}, {'max_usd': 0}, {'input_rate': 0}])
def test_paid_fallback_requires_consent_and_all_limits(override):
    api = API()
    settings = dict(consent=True, max_requests=2, max_tokens=100000, max_usd=1, input_rate=1, output_rate=2)
    settings.update(override)
    provider = OpenAIProvider('secret', transport=api, **settings)
    with pytest.raises(ValueError):
        provider.translate(SPANISH, 'es')
    assert not api.calls


def test_paid_fallback_caches_unique_text_and_enforces_request_limit(tmp_path):
    store, project, _ = make_workspace(tmp_path, source=[{'Name': 'Source', 'Segment': 'SMB', 'Description': SPANISH}])
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Description'], 'source': ['Description']})
    api = API()
    fallback = OpenAIProvider('secret', transport=api, consent=True, max_requests=1, max_tokens=100000, max_usd=1, input_rate=1, output_rate=2)
    engine.process(['master', 'source'], Offline(fail=True), fallback=fallback)
    assert len(api.calls) == 1 and all(r['status'] == 'Translated — Review' for r in engine.records())
    body = api.calls[0][1]['json']
    assert 'master_schema' not in str(body) and body['max_completion_tokens'] == 1024
    with pytest.raises(ValueError):
        fallback.translate('Different input', 'es')
    assert len(api.calls) == 1


def test_batch_limit_treats_datasets_equally(tmp_path):
    master = [{'Name': 'Name', 'Segment': 'SMB', 'Description': SPANISH + f' {i}'} for i in range(10)]
    source = [{'Name': 'Source', 'Segment': 'SMB', 'Description': 'Service commercial pour les entreprises'}]
    store, project, _ = make_workspace(tmp_path, master, source)
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Description'], 'source': ['Description']})
    class PreserveNumbers(Offline):
        def translate(self, text, language):
            self.calls.append((text, language))
            return 'English ' + text
    engine.process(['master', 'source'], PreserveNumbers(), limit=2)
    records = engine.records()
    assert sum(r['status'] == 'Translated — Review' and r['dataset'] == 'master' for r in records) == 1
    assert sum(r['status'] == 'Translated — Review' and r['dataset'] == 'source' for r in records) == 1


def test_integrated_ui_translation_reconciliation_review_and_export(tmp_path, monkeypatch):
    case = make_workspace(tmp_path)
    store, project, _ = case
    # Exercise the retained offline provider through the engine; the page itself
    # is now Google-only. Seed equivalent mock Google metadata before opening it.
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Name', 'Description'], 'source': ['Name', 'Description']})
    from tests.translation_workspace_helpers import seed_google
    seed_google(engine, {'Banco de México': 'es', 'Bank of Mexico': 'en', SPANISH: 'es', ENGLISH: 'en'})
    engine.process(['master', 'source'], Offline(), allow_names=True)
    app = app_for(case, monkeypatch)
    navigate(app, 'Translation Intelligence')
    assert not app.exception
    by_label(app.selectbox, 'Translation dataset').set_value('Master').run()
    by_label(app.selectbox, 'Translation status').set_value('Translated — Review').run()
    selector = by_label(app.selectbox, 'Translation to review')
    index = next(i for i, label in enumerate(selector.options) if '· Name ·' in label)
    selector.set_value(index).run()
    by_label(app.text_area, 'Reviewed English translation').set_value('Bank of Mexico')
    by_label(app.text_input, 'Individual reviewer').set_value('Translation reviewer')
    by_label(app.checkbox, 'I verified this translation preserves the correct entity, proper names, legal suffixes and identifiers.').check()
    by_label(app.button, 'Approve translation').click().run()
    assert not app.exception and len(store.audit(project)) == 1
    navigate(app, 'Run Reconciliation')
    by_label(app.button, 'Run automatic local matching').click().run()
    assert not app.exception
    saved = store.load(project)['results'][0]
    assert saved['proposed'] == 'SMB' and saved['final'] == '' and saved['translation_provenance']
    navigate(app, 'Review & Approvals')
    assert by_label(app.metric, 'Final value').value == '—'
    by_label(app.checkbox, 'Select All filtered accounts').check().run()
    by_label(app.text_input, 'Bulk reviewer name').set_value('Bulk reviewer')
    by_label(app.checkbox, 'I have verified the evidence and field value for every selected account.').check()
    by_label(app.button, 'Save Review Decisions').click().run()
    assert any('Skipped 1' in warning.value for warning in app.warning)
    assert store.load(project)['results'][0]['final'] == '' and len(store.audit(project)) == 1
    by_label(app.text_input, 'Reviewer name').set_value('Account reviewer')
    by_label(app.checkbox, 'I have verified the evidence and the field value for this specific account.').check()
    by_label(app.button, 'Save review decision').click().run()
    assert not app.exception and by_label(app.metric, 'Final value').value == 'SMB'
    assert len(store.audit(project)) == 2
    reloaded = app_for(case, monkeypatch)
    navigate(reloaded, 'Run Reconciliation')
    by_label(reloaded.button, 'Run automatic local matching').click().run()
    assert store.load(project)['results'][0]['final'] == 'SMB'
    import rta.export
    original, exports = rta.export.workbook, []
    def capture(*args, **kwargs):
        data = original(*args, **kwargs)
        exports.append(data)
        return data
    monkeypatch.setattr(rta.export, 'workbook', capture)
    navigate(reloaded, 'Export Results')
    book = load_workbook(BytesIO(exports[-1]))
    assert 'Master Translations' in book.sheetnames
    assert book['Audit Trail'].max_row == 3
    assert store.load(project)['master'][0]['Name'] == 'Banco de México'
