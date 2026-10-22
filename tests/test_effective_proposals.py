from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from rta.export import result_row, summary, workbook
from rta.models import Result, INTELLIGENCE_PENDING, MANUAL_REQUIRED, review_status_label
from tests.test_app import by_label, navigate
from tests.test_review_workflow import workspace, app_for
from tests.test_source_selection import suggestions
from tests.test_recommendations import setup


def test_actual_local_pipeline_uses_effective_proposal(workspace):
    store, project, _ = workspace
    before = store.load(project)
    results = [Result(**r) for r in before['results']]
    assert [r.row_id for r in results if r.status_label == INTELLIGENCE_PENDING] == [0, 1, 4, 8]
    assert all(not r.ai_proposed for r in results)
    counts = summary(results)
    assert counts[INTELLIGENCE_PENDING] == 4 and counts[MANUAL_REQUIRED] == 6
    assert counts['Pending Approvals'] == counts[INTELLIGENCE_PENDING] + counts[MANUAL_REQUIRED] == 10
    for r in results:
        assert r.status_label == review_status_label(r) == result_row(r)['Review Status']
    book = load_workbook(BytesIO(workbook(results, [], [], 10)))
    sheet = book['Reconciliation Results']
    headers = [c.value for c in sheet[1]]
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    for index in [0, 1, 4, 8]:
        assert rows[index]['Review Status'] == INTELLIGENCE_PENDING
        assert rows[index]['Intelligence Proposed Value'] == before['results'][index]['proposed']
        assert rows[index]['AI Proposed Value'] is None
    assert store.load(project) == before and not store.audit(project)


@pytest.mark.parametrize('proposed', [None, '', ' ', '\t\n'])
def test_ai_metadata_never_fills_missing_effective_proposal(proposed):
    r = Result(1, 'Account', 'Original', proposed=proposed, ai_proposed='Enterprise',
               ai_confidence='High', research_status='Completed', recommendation_status='Proposed',
               identity_confidence='Reliable local match')
    before = r.dict()
    assert r.status_label == MANUAL_REQUIRED
    assert r.dict() == before


@pytest.mark.parametrize('ai', ['', None, 'Different segment'])
def test_local_proposal_does_not_depend_on_ai_metadata(ai):
    r = Result(1, 'Account', 'Original', proposed=' Enterprise ', ai_proposed=ai,
               ai_confidence='Low', research_status='Failed')
    assert r.status_label == INTELLIGENCE_PENDING
    assert r.proposed == ' Enterprise '


@pytest.mark.parametrize('decision,status', [('Approve', 'Approved'), ('Reject', 'Rejected'), ('Manual', 'Manually Verified')])
@pytest.mark.parametrize('proposed', [None, 'Enterprise'])
def test_saved_outcomes_precede_effective_proposal(decision, status, proposed):
    r = Result(1, 'Account', 'Original', proposed=proposed, decision=decision, review_status=status)
    assert r.status_label == status


@pytest.mark.parametrize('fixture', ['workspace', 'suggestions'])
def test_human_readable_values_without_raw_segment_block(request, fixture, monkeypatch):
    case = request.getfixturevalue(fixture)
    store, project, _ = case
    before = store.load(project)
    app = app_for(case, monkeypatch)
    for page in ['Review & Approvals', 'AI Research']:
        navigate(app, page)
        assert not app.exception
        assert not any('**Proposed Segment**' in m.value for m in app.markdown)
        assert not any('Supporting Source account' in element.value for element in app.json)
        assert by_label(app.metric, 'Original value')
        assert by_label(app.metric, 'Intelligence Proposed Value')
        assert by_label(app.metric, 'Final value')
        assert not any('AI Proposed — Pending Approval' in m.value for m in [*app.markdown, *app.info])
    assert store.load(project) == before and not store.audit(project)


def test_local_filtered_selection_and_status_bubble(workspace, monkeypatch):
    store, project, _ = workspace
    before = store.load(project)
    app = app_for(workspace, monkeypatch)
    navigate(app, 'Review & Approvals')
    by_label(app.selectbox, 'Review status').set_value(INTELLIGENCE_PENDING).run()
    assert list(app.dataframe[-1].value['Record ID']) == [0, 1, 4, 8]
    assert any('redthread-status' in m.value and INTELLIGENCE_PENDING in m.value for m in app.markdown)
    by_label(app.checkbox, 'Select All filtered accounts').check().run()
    assert any('Selected accounts: 4 of 4' in c.value for c in app.caption)
    by_label(app.text_input, 'Bulk reviewer name').set_value('Local reviewer')
    by_label(app.checkbox, 'I have verified the evidence and field value for every selected account.').check()
    by_label(app.button, 'Save Review Decisions').click().run()
    assert not app.exception
    assert [e['row_id'] for e in store.audit(project)] == [0, 1, 4, 8]
    persisted = store.load(project)
    for i, r in enumerate(persisted['results']):
        if i in [0, 1, 4, 8]:
            assert r['final'] == r['proposed'] and r['review_status'] == 'Approved'
            assert r['original'] == before['results'][i]['original'] and r['ai_proposed'] == ''
        else:
            assert r == before['results'][i]


def test_filtered_select_all_covers_pages_only_for_effective_proposals(workspace, monkeypatch):
    store, project, _ = workspace
    payload = store.load(project)
    local = payload['results'][1]
    records = []
    for index in range(123):
        r = deepcopy(local)
        r.update(row_id=index, identifier=f'Local {index:03}')
        records.append(r)
    ai = deepcopy(local)
    ai.update(row_id=123, identifier='AI proposal', ai_proposed='Healthcare')
    missing = deepcopy(local)
    missing.update(row_id=124, proposed=None, ai_proposed='AI metadata only')
    reviewed = deepcopy(local)
    reviewed.update(row_id=125, decision='Reject', review_status='Rejected')
    payload['results'] = records + [ai, missing, reviewed]
    store.save(project, payload)
    app = app_for(workspace, monkeypatch)
    navigate(app, 'Review & Approvals')
    by_label(app.selectbox, 'Review status').set_value(INTELLIGENCE_PENDING).run()
    by_label(app.number_input, 'Selection page').set_value(3).run()
    by_label(app.checkbox, 'Select All filtered accounts').check().run()
    assert any('Selected accounts: 124 of 124' in c.value for c in app.caption)
    next(c for c in app.checkbox if c.label.startswith('I confirm including ')).check()
    by_label(app.selectbox, 'Bulk decision').set_value('Reject')
    by_label(app.text_input, 'Bulk reviewer name').set_value('Reviewer')
    by_label(app.button, 'Save Review Decisions').click().run()
    assert not app.exception
    assert [e['row_id'] for e in store.audit(project)] == list(range(124))
    assert store.load(project)['results'][124:] == payload['results'][124:]
