"""Fictional XLSX → actual Streamlit workflow → verified Excel report.

AppTest lacks browser file selection, so only uploader return values are injected.
Import, mapping, detection, provider adapter, persistence, review, matching and
Excel generation all execute real application code. Google is always mocked.
"""
from collections import Counter
from copy import deepcopy
from io import BytesIO
import ast
import json
import os
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
from zipfile import ZipFile

from google.api_core.exceptions import DeadlineExceeded
from google.cloud import translate_v3
from openpyxl import load_workbook
import polars as pl
import pytest
import streamlit as st
from streamlit.delta_generator import DeltaGenerator
from streamlit.testing.v1 import AppTest

from rta.ingest import parse_file
from rta.models import Mapping, Result, INTELLIGENCE_PENDING, MANUAL_REQUIRED, REVIEW_FILTERS
from rta.review import final_value
from rta.storage import Store
from rta.translation import TranslationEngine, translation_metrics
from rta.translation_matching import reconcile
from rta.translation_providers import GoogleTranslationProvider
from rta.translation_usage import GoogleTranslationUsage
from tests.test_app import by_label, navigate
from tests.test_google_translation import response
from tests.test_review_workflow import app_for
from tests.test_translation import Offline

FIXTURES = Path(__file__).parent / 'fixtures' / 'translation_intelligence'
EXPECTED = json.loads((FIXTURES / 'expected.json').read_text())
ALLOW_NAMES = 'Allow drafts for organization-name fields; I will verify proper names, legal suffixes and identifiers'
GOOGLE_CONSENT = 'I authorize this translation batch, its estimated cost and retries.'
TRANSLATION_CONFIRM = 'I verified this translation preserves the correct entity, proper names, legal suffixes and identifiers.'
ACCOUNT_CONFIRM = 'I have verified the evidence and the field value for this specific account.'


@pytest.fixture
def workflow(tmp_path, monkeypatch):
    monkeypatch.setenv('RTA_DB_PATH', str(tmp_path / 'e2e.sqlite3'))
    monkeypatch.delenv('TAVILY_API_KEY', raising=False)
    monkeypatch.delenv('OPENAI_API_KEY', raising=False)
    monkeypatch.setattr(st, 'secrets', {})
    monkeypatch.setattr('rta.translation_ui.ArgosProvider', Offline)
    # Block ordinary network requests as a second guard; the Google SDK factory
    # itself is mocked, so its native gRPC transport can never be constructed.
    monkeypatch.setattr('requests.sessions.Session.request',
                        Mock(side_effect=AssertionError('Live HTTP is forbidden in this fictional workflow.')))
    uploaded = {}
    for dataset, label in [('master', 'Master file'), ('source', 'Source file')]:
        path = FIXTURES / f'fictional_{dataset}_accounts.xlsx'
        uploaded[label] = SimpleNamespace(name=path.name, getvalue=lambda path=path: path.read_bytes())
    original_uploader = DeltaGenerator.file_uploader
    def uploader(self, label, *args, **kwargs):
        actual = original_uploader(self, label, *args, **kwargs)
        return uploaded.get(label, actual)
    monkeypatch.setattr(DeltaGenerator, 'file_uploader', uploader)
    client = Mock()
    calls = []
    def translate(**kwargs):
        request = kwargs['request']
        assert kwargs['retry'] is None
        assert request['parent'] == 'projects/rta-translation/locations/global'
        assert request['target_language_code'] == 'en'
        assert request['mime_type'] == 'text/plain'
        calls.append(deepcopy(request))
        return response(*(EXPECTED['translations'][text] for text in request['contents']))
    client.translate_text.side_effect = translate
    factory = Mock(return_value=client)
    monkeypatch.setattr(translate_v3, 'TranslationServiceClient', factory)
    app = AppTest.from_file(Path('app.py').resolve(), default_timeout=30).run()
    assert not app.exception
    navigate(app, 'Upload Data')
    assert not by_label(app.button, 'Import uploaded datasets').disabled
    by_label(app.button, 'Import uploaded datasets').click().run()
    assert not app.exception
    store = Store(tmp_path / 'e2e.sqlite3')
    project, token = app.session_state['project'], app.session_state['token']
    payload = store.load(project)
    for dataset in ('master', 'source'):
        filename = f'fictional_{dataset}_accounts.xlsx'
        parsed = parse_file((FIXTURES / filename).read_bytes(), filename)
        assert parsed.height == 12
        assert payload[dataset] == parsed.to_dicts()
        assert payload[f'{dataset}_name'] == filename
    navigate(app, 'Configure Reconciliation')
    for label, value in [('Master identifier', 'Account Name'), ('Source identifier', 'Account Name'),
                         ('Master comparison field', 'Segment'), ('Source reference field', 'Segment')]:
        by_label(app.selectbox, label).set_value(value)
    by_label(app.button, 'Save column mapping').click().run()
    assert not app.exception
    assert store.load(project)['mapping'] == vars(Mapping('Account Name', 'Account Name', 'Segment', 'Segment'))
    # Annotated Google responses replace the retired offline UI path. They are
    # synthetic contract fixtures, not a claim about live detection accuracy.
    from tests.translation_workspace_helpers import seed_google, select_visible
    languages = ['fr', 'es', 'de', 'fr', 'es', 'de', 'fr', 'es', 'fr']
    labels = {row['Account Name']: 'en' for d in ('master', 'source') for row in payload[d]}
    labels.update(dict(zip(EXPECTED['translations'], languages)))
    labels.update({'Cobalt Lantarn Supply': 'und', 'Unmatched Velvet Acorn Atelier': 'und'})
    seed_google(TranslationEngine(store, project), labels,
                {'master': ['Account Name'], 'source': ['Account Name']})
    navigate(app, 'Translation Intelligence')
    assert TranslationEngine(store, project).plan(['master', 'source'], GoogleTranslationProvider())['unique_texts'] == 9
    select_visible(app)
    assert not app.exception
    factory.assert_not_called()
    return SimpleNamespace(app=app, store=store, project=project, token=token,
                           client=client, factory=factory, calls=calls, translate=translate)


def run_reconciliation(workflow):
    app = workflow.app
    navigate(app, 'Run Reconciliation')
    by_label(app.button, 'Run automatic local matching').click().run()
    assert not app.exception
    return [Result(**row) for row in workflow.store.load(workflow.project)['results']]


def approve_translation(workflow, dataset, row_id):
    app = workflow.app
    navigate(app, 'Translation Intelligence')
    by_label(app.selectbox, 'Translation dataset').set_value(dataset.title()).run()
    by_label(app.selectbox, 'Translation status').set_value('Translated — Review').run()
    selector = by_label(app.selectbox, 'Translation to review')
    index = next(i for i, label in enumerate(selector.options)
                 if label.startswith(f'{dataset.title()} row {row_id + 1} · Account Name ·'))
    selector.set_value(index).run()
    original = workflow.store.load(workflow.project)[dataset][row_id]['Account Name']
    by_label(app.text_area, 'Reviewed English translation').set_value(EXPECTED['translations'][original])
    by_label(app.text_input, 'Individual reviewer').set_value('Fictional Translation Reviewer')
    by_label(app.checkbox, TRANSLATION_CONFIRM).check()
    by_label(app.button, 'Approve translation').click().run()
    assert not app.exception


def sheet_records(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def test_fictional_excel_translation_review_reconciliation_export(workflow, monkeypatch, tmp_path):
    w, app = workflow, workflow.app
    engine = TranslationEngine(w.store, w.project)
    initial_payload = deepcopy(w.store.load(w.project))
    records = engine.records()
    # Independent Master/Source classification from annotated mock Google responses.
    for dataset in ('master', 'source'):
        rows = [r for r in records if r['dataset'] == dataset]
        assert len(rows) == 12
        assert [r['row_id'] for r in rows if r['status'] == 'Translation Needed'] == EXPECTED[f'{dataset}_flagged_rows']
        assert {r['language'] for r in rows if r['status'] == 'Translation Needed'} == {'fr', 'es', 'de'}
    assert [r['row_id'] for r in records if r['dataset'] == 'master' and r['status'] == 'Uncertain'] == []
    assert [r['row_id'] for r in records if r['dataset'] == 'source' and r['status'] == 'Uncertain'] == [8, 10]
    assert translation_metrics(records)[2]['Translation Needed'] == 10
    plan = engine.plan(['master', 'source'], GoogleTranslationProvider(), allow_names=True)
    characters = sum(len(text) for text in EXPECTED['translations'])
    assert plan['unique_texts'] == 9  # exact French text shared by both files
    assert plan['cached_texts'] == 0
    assert plan['characters'] == characters
    assert plan['estimated_usd'] == pytest.approx(characters * 20 / 1_000_000)
    assert any(f"Estimated cost: ${plan['estimated_usd']:.6f}" in c.value for c in app.markdown)
    assert plan['characters'] == characters  # Accounting retained in backend, removed from visible UI.
    by_label(app.button, 'Translate Selected').click().run()
    assert any('Authorize' in error.value for error in app.error)
    w.factory.assert_not_called()
    assert GoogleTranslationUsage(w.store, w.project).summary()['requests'] == 0
    by_label(app.checkbox, GOOGLE_CONSENT).check()
    by_label(app.button, 'Translate Selected').click().run()
    assert not app.exception
    assert w.factory.call_count == 1
    assert w.client.translate_text.call_count == 3
    assert {call['source_language_code'] for call in w.calls} == {'fr', 'es', 'de'}
    sent = [text for call in w.calls for text in call['contents']]
    assert set(sent) == set(EXPECTED['translations'])
    assert len(sent) == len(set(sent)) == 9
    drafts = [r for r in engine.records() if r['status'] == 'Translated — Review']
    assert len(drafts) == 10
    assert all(r['provider'] == GoogleTranslationProvider.name and not r['reviewer'] for r in drafts)
    assert sum(r['cache_hit'] for r in drafts) == 1
    usage = GoogleTranslationUsage(w.store, w.project)
    monthly = usage.summary()
    assert monthly['confirmed_characters'] == characters
    assert monthly['uncertain_characters'] == 0
    assert monthly['requests'] == 3
    assert monthly['remaining_characters'] == 500000 - characters
    assert monthly['estimated_before_credits'] == pytest.approx(plan['estimated_usd'])
    assert monthly['estimated_after_credits'] == 0
    assert monthly['progress'] == characters / 500000
    assert usage.history() == [{'Month': monthly['month'], 'Confirmed characters': characters,
                                'Uncertain characters': 0, 'API request attempts': 3}]
    assert w.store.load(w.project)['master'] == initial_payload['master']
    assert w.store.load(w.project)['source'] == initial_payload['source']
    baseline, _ = reconcile(pl.DataFrame(initial_payload['master']), pl.DataFrame(initial_payload['source']),
                            Mapping(**initial_payload['mapping']))
    unapproved = run_reconciliation(w)
    assert [(r.match_status, r.proposed, final_value(r)) for r in unapproved] == [
        (r.match_status, r.proposed, final_value(r)) for r in baseline]
    assert all(not p['translations'] for r in unapproved for p in r.translation_provenance)
    assert all(not r.final for r in unapproved)
    for row_id in EXPECTED['approve_master_rows']:
        approve_translation(w, 'master', row_id)
    for row_id in [4, 5, 6]:
        approve_translation(w, 'source', row_id)
    partial = run_reconciliation(w)
    assert all(partial[i].match_status == 'Translation-assisted' for i in range(1, 7))
    assert partial[7].match_status != 'Translation-assisted' and not partial[7].proposed
    approve_translation(w, 'source', 7)
    translated = run_reconciliation(w)
    assert all(translated[i].match_status == 'Translation-assisted' for i in EXPECTED['translation_assisted_rows'])
    assert all(translated[i].proposed == EXPECTED['segments'][i] for i in EXPECTED['translation_assisted_rows'])
    assert {p['direction'] for r in translated[1:8] for p in r.translation_provenance} == {
        'Translated Master → Original Source', 'Original Master → Translated Source',
        'Translated Master → Translated Source'}
    assert all(r.status_label == INTELLIGENCE_PENDING for r in translated if r.proposed)
    assert all(not r.final and not final_value(r) for r in translated)
    assert translated[8].match_status == 'Fuzzy candidates' and not translated[8].proposed
    assert not translated[10].proposed  # unmatched business never receives an invented value
    assert all(translated[i].match_status == 'Exact' for i in [0, 9, 11])
    # Select a fresh field with repeated values to test actual persisted cache reuse.
    engine.detect({'master': ['Account Name', 'Trading Name'], 'source': ['Account Name', 'Trading Name']})
    navigate(app, 'Translation Intelligence')
    by_label(app.selectbox, 'Translation dataset').set_value('All').run()
    by_label(app.selectbox, 'Translation status').set_value('Translation Needed').run()
    from tests.translation_workspace_helpers import select_visible
    select_visible(app)
    cached_plan = engine.plan(['master', 'source'], GoogleTranslationProvider(), allow_names=True)
    assert cached_plan['unique_texts'] == cached_plan['cached_texts'] == 9
    assert cached_plan['characters'] == cached_plan['estimated_usd'] == 0
    assert any('Estimated cost: $0.000000' in c.value for c in app.markdown)
    assert not by_label(app.checkbox, GOOGLE_CONSENT).value
    by_label(app.checkbox, GOOGLE_CONSENT).check()
    by_label(app.button, 'Translate Selected').click().run()
    assert not app.exception
    assert w.client.translate_text.call_count == 3 and w.factory.call_count == 1
    assert usage.summary() == monthly
    trading_drafts = [r for r in engine.records() if r['field'] == 'Trading Name' and r['english']]
    assert len(trading_drafts) == 10
    assert all(r['status'] == 'Translated — Review' and r['cache_hit'] for r in trading_drafts)
    assert len([r for r in engine.records() if r['status'] == 'Approved Translation']) == 8
    # Settings and usage remain private to each workspace, like translation cache and reviews.
    usage.set_allowance(1000)  # Accounting controls intentionally removed from this workspace UI.
    assert usage.allowance() == 1000
    other_project, _ = w.store.create()
    w.store.save(other_project, deepcopy(initial_payload))
    other_engine = TranslationEngine(w.store, other_project)
    from tests.translation_workspace_helpers import seed_google
    seed_google(other_engine, {r['original']: r['google_language'] for r in records},
                {'master': ['Account Name'], 'source': ['Account Name']})
    other_usage = GoogleTranslationUsage(w.store, other_project)
    assert other_usage.summary()['requests'] == 0 and other_usage.history() == []
    assert other_usage.allowance() == 500000
    other_plan = other_engine.plan(['master', 'source'], GoogleTranslationProvider(), allow_names=True)
    assert other_plan['cached_texts'] == 0 and other_plan['characters'] == characters
    assert not any(r['status'] == 'Approved Translation' for r in other_engine.records())
    assert w.store.audit(other_project) == []
    # Human account approval is separate from translation approval.
    navigate(app, 'Review & Approvals')
    filters = ['All', INTELLIGENCE_PENDING, MANUAL_REQUIRED, 'Approved', 'Rejected', 'Manually Verified']
    assert REVIEW_FILTERS == filters
    assert by_label(app.selectbox, 'Review status').options == filters
    by_label(app.selectbox, 'Account to review').set_value(1).run()
    assert by_label(app.metric, 'Final value').value == '—'
    by_label(app.text_input, 'Reviewer name').set_value('Fictional Account Reviewer')
    by_label(app.button, 'Save review decision').click().run()
    assert any('Confirm' in error.value for error in app.error)
    assert not w.store.load(w.project)['results'][1]['final']
    by_label(app.checkbox, ACCOUNT_CONFIRM).check()
    by_label(app.button, 'Save review decision').click().run()
    assert not app.exception
    assert by_label(app.metric, 'Final value').value == 'Enterprise'
    by_label(app.selectbox, 'Account to review').set_value(2).run()
    by_label(app.selectbox, 'Decision').set_value('Reject').run()
    by_label(app.text_input, 'Reviewer name').set_value('Fictional Account Reviewer')
    by_label(app.text_area, 'Reviewer notes').set_value('Fictional UAT: reject proposed classification.')
    by_label(app.button, 'Save review decision').click().run()
    assert not app.exception and by_label(app.metric, 'Final value').value == '—'
    by_label(app.selectbox, 'Account to review').set_value(8).run()
    by_label(app.selectbox, 'Decision').set_value('Manual').run()
    by_label(app.selectbox, 'Verified Source value').set_value('SMB')
    by_label(app.text_input, 'Reviewer name').set_value('Fictional Account Reviewer')
    by_label(app.text_area, 'Reviewer notes').set_value('Fictional UAT: verify spelling variant manually.')
    by_label(app.checkbox, ACCOUNT_CONFIRM).check()
    by_label(app.button, 'Save review decision').click().run()
    assert not app.exception and by_label(app.metric, 'Final value').value == 'SMB'
    reviewed = [Result(**r) for r in w.store.load(w.project)['results']]
    assert reviewed[1].status_label == 'Approved' and final_value(reviewed[1]) == 'Enterprise'
    assert reviewed[2].status_label == 'Rejected' and not final_value(reviewed[2])
    assert reviewed[8].status_label == 'Manually Verified' and final_value(reviewed[8]) == 'SMB'
    assert all(not final_value(r) for r in reviewed if r.row_id not in (1, 8))
    for status in filters:
        by_label(app.selectbox, 'Review status').set_value(status).run()
        assert not app.exception
        expected_rows = [r for r in reviewed if status == 'All' or r.status_label == status]
        options = by_label(app.selectbox, 'Account to review').options
        # Filtering scopes the table and bulk decisions; individual review stays on row 8.
        table = next(frame.value for frame in app.dataframe if 'Master Identifier' in frame.value.columns)
        assert list(table['Record ID']) == [r.row_id for r in expected_rows]
        assert by_label(app.selectbox, 'Account to review').value == 8
        assert len(options) == len({r.row_id for r in expected_rows} | {8})
        assert all(any(label.startswith(f'Row {r.row_id + 1} · ') for label in options) for r in expected_rows)
    rerun = run_reconciliation(w)
    assert [r.dict() for r in rerun if r.decision] == [r.dict() for r in reviewed if r.decision]
    assert len(w.store.audit(w.project)) == 11  # 8 translation approvals + 3 account decisions
    # Fresh session must recover persisted review, cache and usage without more Google calls.
    reloaded = app_for((w.store, w.project, w.token), monkeypatch)
    navigate(reloaded, 'Review & Approvals')
    by_label(reloaded.selectbox, 'Review status').set_value('All').run()
    by_label(reloaded.selectbox, 'Account to review').set_value(1).run()
    assert by_label(reloaded.metric, 'Final value').value == 'Enterprise'
    navigate(reloaded, 'Translation Intelligence')
    assert not reloaded.metric  # The monthly dashboard is deliberately hidden.
    recovered_usage = GoogleTranslationUsage(w.store, w.project).summary()
    assert recovered_usage['confirmed_characters'] == characters and recovered_usage['requests'] == 3
    import rta.export
    original_export, exports = rta.export.workbook, []
    def capture(*args, **kwargs):
        data = original_export(*args, **kwargs)
        exports.append(data)
        return data
    monkeypatch.setattr(rta.export, 'workbook', capture)
    navigate(reloaded, 'Export Results')
    data = exports[-1]
    book = load_workbook(BytesIO(data))
    assert set(book.sheetnames) == {'Reconciliation Results', 'Summary', 'Audit Trail', 'Source Conflicts',
                                   'Translation Summary', 'Translation Details', 'Master Translations', 'Source Translations'}
    output = sheet_records(book['Reconciliation Results'])
    assert len(output) == 12
    assert output[1]['Final Value'] == 'Enterprise' and output[1]['Review Status'] == 'Approved'
    assert output[2]['Final Value'] is None and output[2]['Review Status'] == 'Rejected'
    assert output[8]['Final Value'] == 'SMB' and output[8]['Review Status'] == 'Manually Verified'
    for index in [0, 3, 4, 5, 6, 7, 9, 11]:
        assert output[index]['Review Status'] == INTELLIGENCE_PENDING
        assert output[index]['Proposed Value'] == EXPECTED['segments'][index]
        assert output[index]['Final Value'] is None
    assert output[10]['Review Status'] == MANUAL_REQUIRED and output[10]['Final Value'] is None
    for index in range(1, 8):
        assert json.loads(output[index]['Translation Match Provenance'])[0]['translations']
    details = sheet_records(book['Translation Details'])
    assert len(details) == 48
    counts = Counter(r['status'] for r in details)
    assert counts['Approved Translation'] == 8
    assert counts['Translated — Review'] == 12
    for dataset in ('master', 'source'):
        sheet_rows = sheet_records(book[f'{dataset.title()} Translations'])
        assert len(sheet_rows) == 12
        for index, original in enumerate(initial_payload[dataset]):
            assert all(sheet_rows[index][field] == value for field, value in original.items())
        index = EXPECTED[f'approve_{dataset}_rows'][0]
        assert sheet_rows[index]['Account Name (English)'] == EXPECTED['translations'][initial_payload[dataset][index]['Account Name']]
        assert sheet_rows[index]['Account Name (English) status'] == 'Approved Translation'
        assert sheet_rows[index]['Trading Name (English) status'] == 'Translated — Review'
    audit = sheet_records(book['Audit Trail'])
    assert len(audit) == 11
    assert sum(row.get('event_type') == 'translation_review' for row in audit) == 8
    assert {row['decision'] for row in audit} == {'Approve translation', 'Approve', 'Reject', 'Manual'}
    assert all(row['reviewer'] in ('Fictional Translation Reviewer', 'Fictional Account Reviewer') for row in audit)
    for event in audit:
        assert event['timestamp']
        if event.get('event_type') == 'translation_review':
            previous = ast.literal_eval(event['previous_translation'])
            approved = ast.literal_eval(event['translation'])
            assert previous['status'] == 'Translated — Review'
            assert approved['status'] == 'Approved Translation'
            assert approved['english'] == EXPECTED['translations'][approved['original']]
            assert approved['reviewer'] == 'Fictional Translation Reviewer'
        else:
            assert event['previous_decision'] is None and event['previous_final'] is None
            assert event['new_final'] == {'Approve': 'Enterprise', 'Reject': None, 'Manual': 'SMB'}[event['decision']]
    assert all(sheet.freeze_panes == 'A2' and sheet.auto_filter.ref for sheet in book)
    with ZipFile(BytesIO(data)) as archive:
        xml = '\n'.join(archive.read(name).decode() for name in archive.namelist() if name.endswith('.xml'))
    assert w.token not in xml
    assert 'BEGIN PRIVATE KEY' not in xml and 'access_token' not in xml
    assert w.client.translate_text.call_count == 3
    report_path = tmp_path / 'verified-fictional-results.xlsx'
    report_path.write_bytes(data)
    artifact_dir = os.environ.get('RTA_E2E_ARTIFACT_DIR')
    if artifact_dir:
        target = Path(artifact_dir)
        target.mkdir(parents=True, exist_ok=True)
        (target / report_path.name).write_bytes(data)
        (target / 'verification.json').write_text(json.dumps({
            'master_rows': 12, 'source_rows': 12, 'unique_translated_texts': 9,
            'mocked_api_requests': 3, 'confirmed_input_characters': characters,
            'translated_cells_before_reviews': 10, 'approved_translation_cells': 8,
            'translation_assisted_accounts': 7, 'exported_translation_records': len(details),
            'audit_events': len(audit), 'export_sheets': book.sheetnames,
            'review_status_filters': filters, 'live_api_calls': 0,
        }, indent=2) + '\n')


def test_fictional_batch_timeout_usage_and_explicit_retry(workflow):
    w, app = workflow, workflow.app
    first = True
    def timeout_once(**kwargs):
        nonlocal first
        if first:
            first = False
            raise DeadlineExceeded('Fictional timeout; no external request was made.')
        return w.translate(**kwargs)
    w.client.translate_text.side_effect = timeout_once
    by_label(app.checkbox, GOOGLE_CONSENT).check()
    by_label(app.button, 'Translate Selected').click().run()
    assert not app.exception
    assert any('2 failed' in notice.value for notice in app.warning)
    engine = TranslationEngine(w.store, w.project)
    records = engine.records()
    assert len([r for r in records if r['status'] == 'Failed']) == 2
    assert len([r for r in records if r['status'] == 'Translated — Review']) == 8
    german = [text for text in EXPECTED['translations'] if text in {
        'Werkstätten am Silbernen Morgen', 'Systeme der Goldenen Wiese'}]
    german_characters = sum(map(len, german))
    all_characters = sum(map(len, EXPECTED['translations']))
    usage = GoogleTranslationUsage(w.store, w.project)
    failed_month = usage.summary()
    assert failed_month['requests'] == w.client.translate_text.call_count == 3
    assert failed_month['uncertain_characters'] == german_characters
    assert failed_month['confirmed_characters'] == all_characters - german_characters
    assert failed_month['total_characters'] == all_characters
    retry_plan = engine.plan(['master', 'source'], GoogleTranslationProvider(), allow_names=True)
    assert retry_plan['unique_texts'] == 2 and retry_plan['characters'] == german_characters
    assert not by_label(app.checkbox, GOOGLE_CONSENT).value
    by_label(app.button, 'Translate Selected').click().run()
    assert w.client.translate_text.call_count == 3
    by_label(app.checkbox, GOOGLE_CONSENT).check()
    by_label(app.button, 'Translate Selected').click().run()
    assert not app.exception
    assert w.client.translate_text.call_count == 4
    assert not [r for r in engine.records() if r['status'] == 'Failed']
    assert len([r for r in engine.records() if r['status'] == 'Translated — Review']) == 10
    monthly = usage.summary()
    assert monthly['requests'] == 4
    assert monthly['confirmed_characters'] == all_characters
    assert monthly['uncertain_characters'] == german_characters
    assert monthly['total_characters'] == all_characters + german_characters
    assert not w.store.audit(w.project)  # translation execution never invents approval events
