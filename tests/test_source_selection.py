from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from rta.export import workbook
from rta.models import Result
from rta.review import source_suggestion, suggestion_version
from rta.storage import Store, digest
from tests.test_app import by_label, navigate
from tests.test_recommendations import setup, research, M
from tests.test_review_workflow import app_for


@pytest.fixture
def suggestions(setup):
    research(setup)
    store, project, token, source, r = setup
    other = deepcopy(r.dict())
    other.update(row_id=1, identifier='Other Beta')
    store.save(project, {
        'source': source.to_dicts(), 'mapping': vars(M), 'context': 'source-context',
        'source_schema': dict.fromkeys(source.columns, 'String'),
        'master': [{'Account': r.identifier, 'Segment': r.original},
                   {'Account': 'Other Beta', 'Segment': 'Verify'}],
        'master_schema': {'Account': 'String', 'Segment': 'String'},
        'results': [r.dict(), other],
    })
    return store, project, token


def approve(store, project, source_row, **kwargs):
    payload = store.load(project)
    result = Result(**payload['results'][0])
    return store.review(project, 0, 'Approve', 'Reviewer', 'Independently checked Source identity and segment',
        selected_source_row=source_row, selection_version=suggestion_version(result, payload),
        confirmed=True, **kwargs)


@pytest.mark.parametrize('source_row,expected', [(0, 'Enterprise'), (1, 'SMB'), (2, 'Public Sector')])
def test_each_suggestion_saves_actual_source_value(suggestions, source_row, expected):
    store, project, _ = suggestions
    before = store.load(project)
    saved = approve(store, project, source_row)
    reopened = Store(store.path)
    assert reopened.load(project) == saved
    r = Result(**saved['results'][0])
    assert r.final == expected and r.review_status == 'Approved' and r.decision == 'Approve'
    assert r.original == 'Verify' and r.ai_proposed == r.proposed == 'Enterprise'
    assert r.value_origin == 'reviewer-selected source'
    assert r.selected_source == {
        'source_row': source_row, 'account': before['source'][source_row]['Company'],
        'value': expected, 'similarity': before['results'][0]['candidates'][source_row]['similarity'],
        'dataset': digest({'source': before['source'], 'mapping': before['mapping']}),
    }
    for field in ['identifier', 'row_id', 'candidates', 'research_snapshot', 'evidence', 'match_status']:
        assert saved['results'][0][field] == before['results'][0][field]
    assert saved['results'][1] == before['results'][1]
    event = reopened.audit(project)[0]
    assert event['selected_source'] == r.selected_source
    assert event['reviewer'] == r.reviewer == 'Reviewer'
    assert event['timestamp'] == r.reviewed_at and event['decision'] == 'Approve'
    assert event['approved_value'] == expected and event['ai_proposed'] == 'Enterprise'
    book = load_workbook(BytesIO(workbook([r], reopened.audit(project), [], 3)))
    sheet = book['Reconciliation Results']
    exported = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert exported['Final Value'] == exported['Selected Source Comparison Value'] == expected
    assert exported['AI Proposed Value'] == 'Enterprise'
    assert exported['Selected Source Record ID'] == source_row
    assert exported['Selected Source Similarity'] == r.selected_source['similarity']


@pytest.mark.parametrize('value', [None, '', '  ', 'N/A', 'Verify'])
def test_missing_source_values_never_approve(suggestions, value):
    store, project, _ = suggestions
    payload = store.load(project)
    payload['source'][1]['Value'] = value
    store.save(project, payload)
    before = store.load(project)
    with pytest.raises(ValueError, match='no valid comparison value'):
        approve(store, project, 1)
    assert store.load(project) == before and not store.audit(project)


def test_cached_candidate_value_is_not_authoritative(suggestions):
    store, project, _ = suggestions
    payload = store.load(project)
    payload['results'][0]['candidates'][1]['value'] = 'Invented segment'
    store.save(project, payload)
    assert approve(store, project, 1)['results'][0]['final'] == 'SMB'


@pytest.mark.parametrize('source_row', [-1, 3, 999, '1', True])
def test_unavailable_or_invalid_source_rows(suggestions, source_row):
    store, project, _ = suggestions
    with pytest.raises(ValueError, match='Source record'):
        approve(store, project, source_row)
    assert not store.audit(project)


def test_confirmation_and_stale_versions(suggestions):
    store, project, _ = suggestions
    payload = store.load(project)
    version = suggestion_version(Result(**payload['results'][0]), payload)
    with pytest.raises(ValueError, match='Confirm'):
        store.review(project, 0, 'Approve', 'R', selected_source_row=1, selection_version=version)
    payload['source'][1]['Value'] = 'Changed segment'
    store.save(project, payload)
    with pytest.raises(ValueError, match='dataset changed'):
        store.review(project, 0, 'Approve', 'R', selected_source_row=1, selection_version=version, confirmed=True)
    with pytest.raises(ValueError, match='dataset changed'):
        store.review(project, 1, 'Approve', 'R', selected_source_row=1, selection_version=version, confirmed=True)
    assert not store.audit(project)


@pytest.mark.parametrize('change', ['insufficient', 'invalid_evidence', 'conflict', 'identity'])
def test_selection_preserves_evidence_and_identity_safeguards(suggestions, change):
    store, project, _ = suggestions
    payload = store.load(project)
    if change == 'insufficient':
        payload['results'][0]['recommendation_status'] = 'Insufficient Evidence'
    elif change == 'invalid_evidence':
        payload['results'][0]['research_snapshot']['retrieved_evidence'][0]['content'] = 'Unrelated'
    elif change == 'conflict':
        payload['source'].append({'Company': 'Beta', 'Value': 'Different segment'})
    else:
        payload['source'][1]['Company'] = 'Replaced identity'
    store.save(project, payload)
    before = store.load(project)
    with pytest.raises(ValueError):
        approve(store, project, 1)
    assert store.load(project) == before and not store.audit(project)


@pytest.mark.parametrize('action,expected', [('Reject', ''), ('Manual', 'Custom')])
def test_reject_and_manual_ignore_suggestion(suggestions, action, expected):
    store, project, _ = suggestions
    approve(store, project, 1)
    # These actions must not validate or reuse a stale selected suggestion.
    saved = store.review(project, 0, action, 'R', manual_value='Custom', selected_source_row=999)
    assert saved['results'][0]['final'] == expected
    assert saved['results'][0]['selected_source'] == {}
    assert store.audit(project)[1]['previous_selected_source']['account'] == 'Beta'


def test_legacy_and_bulk_approval_do_not_use_individual_selection(suggestions):
    store, project, _ = suggestions
    approve(store, project, 1)
    payload = store.load(project)
    version = digest(Result(**payload['results'][0]).dict())
    payload, report = store.bulk_review(project, [0], 'Approve', 'Bulk reviewer', confirmed=True,
                                       overwrite_versions={0: version})
    assert report['saved'] == [0]
    assert payload['results'][0]['final'] == 'Enterprise'
    assert payload['results'][0]['selected_source'] == {}
    store.review(project, 0, 'Approve', 'Legacy reviewer')
    assert store.load(project)['results'][0]['final'] == 'Enterprise'
    legacy = Result(9, 'Legacy', 'SMB', final='SMB', review_status='Approved')
    assert legacy.selected_source == {}


@pytest.mark.parametrize('choice,expected', [(0, 'Enterprise'), (1, 'SMB'), (2, 'Public Sector')])
def test_ui_selection_sync_and_reload(suggestions, monkeypatch, choice, expected):
    app = app_for(suggestions, monkeypatch)
    navigate(app, 'Review & Approvals')
    radio = by_label(app.radio, 'Top 3 Suggested Source Matches')
    assert radio.value is None and len(radio.options) == 3
    radio.set_value(choice).run()
    assert by_label(app.metric, 'Selected Source comparison value').value == expected
    assert by_label(app.metric, 'Final value preview').value == expected
    assert by_label(app.metric, 'Final value').value == '—'
    by_label(app.text_input, 'Reviewer name').set_value('UI reviewer')
    by_label(app.checkbox, 'I have verified the evidence and the field value for this specific account.').check()
    by_label(app.button, 'Save review decision').click().run()
    assert not app.exception
    assert by_label(app.metric, 'Final value').value == expected
    assert by_label(app.metric, 'Pending Approvals').value == '1'
    assert by_label(app.metric, 'Approved Corrections').value == '1'
    assert any('redthread-status' in m.value and 'Approved' in m.value for m in app.markdown)
    assert app.dataframe[-1].value.iloc[0]['Final Value'] == expected
    assert app.dataframe[-1].value.iloc[0]['Review Status'] == 'Approved'
    reloaded = app_for(suggestions, monkeypatch)
    navigate(reloaded, 'Review & Approvals')
    by_label(reloaded.selectbox, 'Review status').set_value('All').run()
    by_label(reloaded.selectbox, 'Account to review').set_value(0).run()
    assert by_label(reloaded.metric, 'Final value').value == expected
    # Capture the report actually produced by the Export Results page.
    import rta.export
    original = rta.export.workbook
    exports = []
    def capture(*args, **kwargs):
        data = original(*args, **kwargs)
        exports.append(data)
        return data
    monkeypatch.setattr(rta.export, 'workbook', capture)
    navigate(reloaded, 'Export Results')
    sheet = load_workbook(BytesIO(exports[-1]))['Reconciliation Results']
    row = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert row['Final Value'] == row['Selected Source Comparison Value'] == expected
    assert row['AI Proposed Value'] == 'Enterprise'
    navigate(reloaded, 'Run Reconciliation')
    assert reloaded.dataframe[-1].value.iloc[0]['Final Value'] == expected


def test_ui_clears_selection_on_account_and_dataset_changes(suggestions, monkeypatch):
    app = app_for(suggestions, monkeypatch)
    navigate(app, 'Review & Approvals')
    by_label(app.radio, 'Top 3 Suggested Source Matches').set_value(1).run()
    by_label(app.button, 'Clear suggested Source selection').click().run()
    assert by_label(app.radio, 'Top 3 Suggested Source Matches').value is None
    by_label(app.radio, 'Top 3 Suggested Source Matches').set_value(1).run()
    by_label(app.checkbox, 'I have verified the evidence and the field value for this specific account.').check().run()
    by_label(app.radio, 'Top 3 Suggested Source Matches').set_value(2).run()
    assert not by_label(app.checkbox, 'I have verified the evidence and the field value for this specific account.').value
    by_label(app.selectbox, 'Account to review').set_value(1).run()
    assert by_label(app.radio, 'Top 3 Suggested Source Matches').value is None
    by_label(app.selectbox, 'Account to review').set_value(0).run()
    assert by_label(app.radio, 'Top 3 Suggested Source Matches').value is None
    by_label(app.radio, 'Top 3 Suggested Source Matches').set_value(1).run()
    store, project, _ = suggestions
    payload = store.load(project)
    payload['source'][1]['Value'] = 'New dataset value'
    store.save(project, payload)
    app.run()
    assert not app.exception and by_label(app.radio, 'Top 3 Suggested Source Matches').value is None


def test_ui_missing_value_preview_is_blocked(suggestions, monkeypatch):
    store, project, _ = suggestions
    payload = store.load(project)
    payload['source'][1]['Value'] = None
    store.save(project, payload)
    app = app_for(suggestions, monkeypatch)
    navigate(app, 'Review & Approvals')
    by_label(app.radio, 'Top 3 Suggested Source Matches').set_value(1).run()
    assert by_label(app.metric, 'Final value preview').value == 'Blocked'
    by_label(app.text_input, 'Reviewer name').set_value('R')
    by_label(app.checkbox, 'I have verified the evidence and the field value for this specific account.').check()
    by_label(app.button, 'Save review decision').click().run()
    assert any('no valid comparison value' in e.value for e in app.error)
    assert not store.audit(project)
