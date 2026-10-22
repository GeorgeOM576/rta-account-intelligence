"""Google-primary detection, consent, cache isolation and integration; mocks only."""
from copy import deepcopy
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
import sqlite3

import polars as pl
import pytest
from google.api_core.exceptions import DeadlineExceeded, PermissionDenied

from rta.google_detection import (GoogleDetectionProvider, GOOGLE_POLICY, GOOGLE_SOURCE, interpret,
                                  cached_result, plan_detection, run_detection)
from rta.language_detection import FOREIGN, AMBIGUOUS
from rta.models import Mapping, Result, INTELLIGENCE_PENDING, REVIEW_FILTERS
from rta.storage import Store
from rta.translation import TranslationEngine, translation_eligible
from rta.translation_matching import reconcile
from rta.translation_providers import GoogleTranslationProvider
from rta.translation_usage import GoogleTranslationUsage
from tests.test_app import by_label, navigate
from tests.test_google_translation import response as translation_response
from tests.test_review_workflow import app_for
from tests.test_translation import make_workspace, Offline, SPANISH, ENGLISH

CASES = json.loads((Path(__file__).parent / 'fixtures/language_detection/labelled_cases.json').read_text())


def detected(language='es', confidence=0.99, alternatives=()):
    return SimpleNamespace(languages=[SimpleNamespace(language_code=language, confidence=confidence),
        *[SimpleNamespace(language_code=code, confidence=score) for code, score in alternatives]])


def google_case(tmp_path, master=None, source=None):
    case = make_workspace(tmp_path, master, source)
    engine = TranslationEngine(case[0], case[1])
    engine.select_detection_provider('google')
    return case, engine


def test_google_is_default_and_no_local_or_cloud_work_on_upload_or_read(tmp_path, monkeypatch):
    case = make_workspace(tmp_path)
    store, project, _ = case
    payload = store.load(project)
    payload.pop('language_detection_provider')
    store.save(project, payload)
    local = Mock(side_effect=AssertionError('No automatic local detection'))
    monkeypatch.setattr('rta.translation.detect_text', local)
    from google.cloud import translate_v3
    cloud = Mock(side_effect=AssertionError('No cloud calls without consent'))
    monkeypatch.setattr(translate_v3, 'TranslationServiceClient', cloud)
    engine = TranslationEngine(store, project)
    records = engine.detect({'master': ['Name', 'Segment'], 'source': ['Name', 'Segment']})
    assert all(r['detection_provider'] == 'google' and r['google_language'] == '' for r in records)
    assert all(r['status'] == 'Excluded' for r in records if r['field'] == 'Segment')
    assert all(r['status'] == 'Uncertain' and r['language'] == 'und' for r in records if r['field'] == 'Name')
    engine.records(); engine.export_records(); engine.reanalyse()
    local.assert_not_called(); cloud.assert_not_called()
    assert GoogleTranslationUsage(store, project).summary()['requests'] == 0


def test_detection_provider_consent_and_adc_project(monkeypatch):
    from google.cloud import translate_v3
    client = Mock()
    client.detect_language.return_value = detected()
    factory = Mock(return_value=client)
    monkeypatch.setattr(translate_v3, 'TranslationServiceClient', factory)
    with pytest.raises(ValueError, match='confirmation'):
        GoogleDetectionProvider().detect(SPANISH)
    factory.assert_not_called()
    value = GoogleDetectionProvider(consent=True).detect(SPANISH)
    assert value['language_code'] == 'es' and value['confidence'] == 0.99
    assert factory.call_args.kwargs['client_options'].quota_project_id == 'rta-translation'
    assert 'credentials' not in factory.call_args.kwargs
    kwargs = client.detect_language.call_args.kwargs
    assert kwargs['request'] == {'parent': 'projects/rta-translation/locations/global', 'content': SPANISH, 'mime_type': 'text/plain'}
    assert kwargs['retry'] is None


def test_cache_deduplicates_across_datasets_fields_and_repeated_runs(tmp_path):
    row = {'Name': 'Les Ateliers de la Lune Violette', 'Description': 'Les Ateliers de la Lune Violette', 'Segment': 'Enterprise'}
    case, engine = google_case(tmp_path, [row] * 3, [row] * 2)
    engine.detect({'master': ['Name', 'Description', 'Segment'], 'source': ['Name', 'Description', 'Segment']})
    plan = plan_detection(engine, ['master', 'source'])
    assert plan['requests'] == plan['unique_texts'] == 1 and plan['characters'] == len(row['Name'])
    client = Mock()
    client.detect_language.return_value = detected('fr')
    before = deepcopy(case[0].load(case[1]))
    with pytest.raises(ValueError):
        run_detection(engine, ['master', 'source'], client=client)
    client.detect_language.assert_not_called()
    run_detection(engine, ['master', 'source'], consent=True, client=client)
    assert client.detect_language.call_count == 1
    assert all(r['google_language'] == 'fr' and r['detection_source'] == GOOGLE_SOURCE for r in engine.records() if r['field'] != 'Segment')
    assert all(r['status'] == 'Excluded' for r in engine.records() if r['field'] == 'Segment')
    assert plan_detection(engine, ['master', 'source'])['requests'] == 0
    usage = GoogleTranslationUsage(case[0], case[1]).summary()
    run_detection(engine, ['master', 'source'], consent=True, client=client)
    assert client.detect_language.call_count == 1
    assert GoogleTranslationUsage(case[0], case[1]).summary() == usage
    assert case[0].load(case[1])['master'] == before['master']
    assert not case[0].audit(case[1])


def test_local_cache_cannot_be_presented_as_google_and_workspace_isolation(tmp_path):
    case = make_workspace(tmp_path)
    store, project, _ = case
    engine = TranslationEngine(store, project)
    engine.detect({'master': ['Description'], 'source': ['Description']})
    assert engine.records()[0]['language'] == 'es'
    assert 'offline' in engine.records()[0]['detection_source']
    engine.select_detection_provider('google')
    records = engine.records()
    assert all(r['google_language'] == '' and 'pending' in r['detection_source'] for r in records)
    assert plan_detection(engine, ['master', 'source'])['requests'] == 2
    client = Mock()
    client.detect_language.side_effect = lambda **kw: detected('es' if kw['request']['content'] == SPANISH else 'en')
    run_detection(engine, ['master', 'source'], consent=True, client=client)
    assert client.detect_language.call_count == 2
    other, _ = store.create()
    store.save(other, deepcopy(store.load(project)))
    other_engine = TranslationEngine(store, other)
    other_engine.detect({'master': ['Description'], 'source': ['Description']})
    assert plan_detection(other_engine, ['master', 'source'])['requests'] == 2
    assert cached_result(store, other, SPANISH) is None
    assert GoogleTranslationUsage(store, other).summary()['requests'] == 0
    # Switching back requires explicit offline selection and never relabels it Google.
    engine.select_detection_provider('offline')
    assert all(r['google_language'] == '' and 'offline' in r['detection_source'] for r in engine.records())


@pytest.mark.parametrize('text,language,confidence,alternatives', [
    ('Müller GmbH', 'de', 0.99, ()), ('Data Studio', 'en', 0.99, ()),
    ('Nova', 'es', 0.99, ()), ('Banco Azul', 'es', 0.99, ()),
    ('Les Ateliers de la Lune Violette', 'fr', 0.6, ()),
    ('Les Ateliers de la Lune Violette', 'fr', 0.95, (('es', 0.9),)),
])
def test_valid_google_language_is_authoritative_with_advisory_review(text, language, confidence, alternatives):
    raw = {'language_code': language, 'confidence': confidence,
           'alternatives': [{'language_code': code, 'confidence': score} for code, score in alternatives]}
    item = interpret(text, 'account_name', raw=raw)
    assert item['status'] == ('Not Required' if language == 'en' else 'Translation Needed') and item['language'] == language
    assert item['language_review_required'] == (confidence < .90 or bool(alternatives))
    assert item['google_language'] == language and item['google_confidence'] == confidence
    assert translation_eligible(item) == (language != 'en')


@pytest.mark.parametrize('error', [DeadlineExceeded, PermissionDenied])
def test_failure_usage_uncertain_and_no_automatic_retry(tmp_path, error):
    case, engine = google_case(tmp_path)
    engine.detect({'master': ['Description'], 'source': []})
    client = Mock()
    client.detect_language.side_effect = error('secret credentials must never be displayed')
    report = run_detection(engine, ['master'], consent=True, client=client)
    assert report['failed'] == 1 and client.detect_language.call_count == 1
    assert 'secret' not in engine.records()[0]['reason']
    usage = GoogleTranslationUsage(case[0], case[1]).summary()
    assert usage['operations']['detection']['uncertain_characters'] == len(SPANISH)
    assert usage['operations']['translation']['requests'] == 0
    assert plan_detection(engine, ['master'])['requests'] == 1


def test_detection_translation_usage_separate_shared_credit_and_manual_approval(tmp_path):
    case, engine = google_case(tmp_path)
    store, project, _ = case
    engine.detect({'master': ['Description'], 'source': []})
    client = Mock()
    client.detect_language.return_value = detected()
    run_detection(engine, ['master'], consent=True, client=client)
    client.translate_text.return_value = translation_response(ENGLISH)
    # Detection consent does not authorize subsequent translation.
    with pytest.raises(ValueError):
        engine.process(['master'], GoogleTranslationProvider(client=client))
    client.translate_text.assert_not_called()
    engine.process(['master'], GoogleTranslationProvider(consent=True, client=client))
    usage = GoogleTranslationUsage(store, project)
    usage.set_allowance(len(SPANISH))
    summary = usage.summary()
    assert summary['operations']['detection']['confirmed_characters'] == len(SPANISH)
    assert summary['operations']['translation']['confirmed_characters'] == len(SPANISH)
    assert summary['operations']['detection']['requests'] == summary['operations']['translation']['requests'] == 1
    assert summary['total_characters'] == len(SPANISH) * 2
    assert summary['estimated_after_credits'] == pytest.approx(len(SPANISH) * 20 / 1e6)
    assert engine.records()[0]['status'] == 'Translated — Review'
    assert not store.load(project)['results']


def test_manual_language_and_translation_correction_never_overwritten(tmp_path):
    case, engine = google_case(tmp_path)
    engine.detect({'master': ['Name'], 'source': []})
    original = engine.records()[0]
    manual = engine.confirm_language(original, 'fr', 'Human language reviewer', True)
    client = Mock()
    run_detection(engine, ['master'], consent=True, client=client)
    client.detect_language.assert_not_called()
    approved = engine.approve(engine.records()[0], 'Reviewed entity name', 'Human translator', True, 'fr')
    audit = case[0].audit(case[1])
    engine.reanalyse()
    item = engine.records()[0]
    assert item['language'] == manual['language'] and item['english'] == approved['english']
    assert item['status'] == 'Approved Translation'
    assert case[0].audit(case[1]) == audit
    assert item['google_language'] == ''


def test_legacy_usage_schema_migration_preserves_rows_as_translation(tmp_path):
    path = tmp_path / 'legacy.sqlite3'
    with sqlite3.connect(path) as db:
        db.execute('CREATE TABLE google_translation_usage(id INTEGER PRIMARY KEY, project TEXT, month TEXT, submitted_at TEXT, characters INTEGER, status TEXT)')
        db.execute("INSERT INTO google_translation_usage VALUES (1,'workspace','2026-10','2026-10-21T12:00:00Z',16,'confirmed')")
    store = Store(path)
    with store.connect() as db:
        assert db.execute('SELECT operation,characters,status FROM google_translation_usage').fetchone() == ('translation', 16, 'confirmed')
    summary = GoogleTranslationUsage(store, 'workspace').summary('2026-10')
    assert summary['operations']['translation']['confirmed_characters'] == 16
    assert summary['operations']['detection']['requests'] == 0


def test_detection_bookkeeping_failure_stops_dispatch(tmp_path, monkeypatch):
    case, engine = google_case(tmp_path)
    engine.detect({'master': ['Name', 'Description'], 'source': []})
    client = Mock()
    monkeypatch.setattr(GoogleTranslationUsage, 'start_request', Mock(side_effect=sqlite3.OperationalError('private')))
    with pytest.raises(RuntimeError, match='No further requests'):
        run_detection(engine, ['master'], consent=True, client=client)
    client.detect_language.assert_not_called()


def test_google_primary_ui_detection_translation_reconciliation_export(tmp_path, monkeypatch):
    master = [{'Name': 'Les Ateliers de la Lune Violette', 'Segment': 'Verify', 'Description': ''}]
    source = [{'Name': 'Violet Moon Workshops', 'Segment': 'Enterprise', 'Description': ''}]
    case, engine = google_case(tmp_path, master, source)
    engine.detect({'master': ['Name'], 'source': ['Name']})
    monkeypatch.setattr('rta.translation_ui.ArgosProvider', Offline)
    from google.cloud import translate_v3
    client = Mock()
    client.detect_language.side_effect = lambda **kw: detected('fr' if kw['request']['content'] == master[0]['Name'] else 'en')
    client.translate_text.return_value = translation_response('Violet Moon Workshops')
    factory = Mock(return_value=client)
    monkeypatch.setattr(translate_v3, 'TranslationServiceClient', factory)
    app = app_for(case, monkeypatch)
    navigate(app, 'Translation Intelligence')
    assert not any(r.label == 'Language detection provider' for r in app.radio)
    factory.assert_not_called()
    assert any('2 names' in c.value and 'Estimated cost' in c.value for c in app.markdown)
    by_label(app.button, 'Detect Languages').click().run()
    factory.assert_not_called()
    assert app.error
    consent_label = 'I authorize this detection batch and its estimated cost.'
    by_label(app.checkbox, consent_label).check()
    by_label(app.button, 'Detect Languages').click().run()
    assert not app.exception and client.detect_language.call_count == 2
    assert by_label(app.button, 'Detect Languages').disabled
    assert client.detect_language.call_count == 2  # all cached; no duplicate requests
    from tests.translation_workspace_helpers import select_visible
    select_visible(app)
    translation_consent = 'I authorize this translation batch, its estimated cost and retries.'
    by_label(app.button, 'Translate Selected').click().run()
    client.translate_text.assert_not_called()
    by_label(app.checkbox, translation_consent).check()
    by_label(app.button, 'Translate Selected').click().run()
    assert not app.exception and client.translate_text.call_count == 1
    payload = case[0].load(case[1])
    assert not reconcile(pl.DataFrame(master), pl.DataFrame(source), Mapping(**payload['mapping']), engine.records())[0][0].proposed
    by_label(app.selectbox, 'Translation dataset').set_value('Master').run()
    by_label(app.selectbox, 'Translation status').set_value('Translated — Review').run()
    by_label(app.text_input, 'Individual reviewer').set_value('Google mock reviewer')
    by_label(app.checkbox, 'I verified this translation preserves the correct entity, proper names, legal suffixes and identifiers.').check()
    by_label(app.button, 'Approve translation').click().run()
    navigate(app, 'Run Reconciliation')
    by_label(app.button, 'Run automatic local matching').click().run()
    result = Result(**case[0].load(case[1])['results'][0])
    assert result.proposed == 'Enterprise' and result.status_label == INTELLIGENCE_PENDING
    assert not result.final and result.match_status == 'Translation-assisted'
    navigate(app, 'Review & Approvals')
    assert by_label(app.selectbox, 'Review status').options == REVIEW_FILTERS
    by_label(app.text_input, 'Reviewer name').set_value('Account reviewer')
    by_label(app.checkbox, 'I have verified the evidence and the field value for this specific account.').check()
    by_label(app.button, 'Save review decision').click().run()
    assert case[0].load(case[1])['results'][0]['final'] == 'Enterprise'
    import rta.export
    actual, exports = rta.export.workbook, []
    def capture(*args, **kwargs):
        data = actual(*args, **kwargs); exports.append(data); return data
    monkeypatch.setattr(rta.export, 'workbook', capture)
    navigate(app, 'Export Results')
    from openpyxl import load_workbook
    from io import BytesIO
    book = load_workbook(BytesIO(exports[-1]))
    assert 'Translation Details' in book.sheetnames and book['Audit Trail'].max_row == 3
    summary = GoogleTranslationUsage(case[0], case[1]).summary()
    assert summary['operations']['detection']['requests'] == 2
    assert summary['operations']['translation']['requests'] == 1
    assert client.detect_language.call_count == 2 and client.translate_text.call_count == 1


def test_68_case_routing_comparison_is_mock_only(tmp_path):
    """An annotated-response contract test, explicitly NOT measured Google accuracy."""
    client = Mock()
    actual = []
    for case in CASES:
        # Synthetic responses use hand labels; they are not captured Google data.
        code = case['language'] or 'und'
        client.detect_language.return_value = detected(code, 0.99 if code != 'und' else 0.4)
        from rta.google_detection import preflight
        if preflight(case['text'], case['field_kind']):
            result = interpret(case['text'], case['field_kind'])
        else:
            raw = GoogleDetectionProvider(consent=True, client=client).detect(case['text'])
            result = interpret(case['text'], case['field_kind'], raw=raw)
        actual.append({**case, 'prediction': result})
    positive = [r for r in actual if r['prediction']['status'] == 'Translation Needed']
    assert len(actual) == 68 and all(r['label'] == 'foreign' for r in positive)
    assert len(positive) == 28
    assert sum(r['label'] == 'foreign' for r in actual) == 28
    assert all(r['prediction']['google_language'] == r['language'] for r in positive)
    assert all(r['prediction']['status'] == 'Excluded' for r in actual if r['label'] == 'excluded')
    assert not any(r['prediction']['status'] == 'Translation Needed' for r in actual if r['label'] == 'english')


@pytest.mark.parametrize('confidence', [float('nan'), float('inf'), -0.1, 1.1])
def test_invalid_google_confidence_is_not_cached_or_trusted(tmp_path, confidence):
    case, engine = google_case(tmp_path)
    engine.detect({'master': ['Description'], 'source': []})
    client = Mock()
    client.detect_language.return_value = detected('es', confidence)
    report = run_detection(engine, ['master'], consent=True, client=client)
    assert report['failed'] == 1
    assert cached_result(case[0], case[1], SPANISH) is None
    assert engine.records()[0]['language'] == 'und'
    assert GoogleTranslationUsage(case[0], case[1]).summary()['operations']['detection']['confirmed_characters'] == len(SPANISH)


def test_google_english_result_is_not_overridden_by_script_heuristics():
    result = interpret('Hello 世界', 'account_name', raw={'language_code': 'en', 'confidence': 0.99})
    assert result['status'] == 'Not Required' and result['language'] == 'en'
    assert result['google_language'] == 'en'


def test_setup_failure_reports_zero_api_attempts(tmp_path, monkeypatch):
    case, engine = google_case(tmp_path)
    engine.detect({'master': ['Description'], 'source': []})
    from google.cloud import translate_v3
    monkeypatch.setattr(translate_v3, 'TranslationServiceClient', Mock(side_effect=ValueError('private ADC details')))
    report = run_detection(engine, ['master'], consent=True)
    assert report['planned_requests'] == 1 and report['requests'] == 0 and report['failed'] == 1
    assert GoogleTranslationUsage(case[0], case[1]).summary()['requests'] == 0


def test_malformed_google_cache_entry_requires_new_consent_not_local_fallback(tmp_path):
    case, engine = google_case(tmp_path)
    engine.detect({'master': ['Description'], 'source': []})
    from rta.google_detection import raw_key
    with case[0].connect() as db:
        db.execute('INSERT INTO language_detection_cache VALUES (?,?,?)',
                   (case[1], raw_key(SPANISH), json.dumps({'provider': 'google', 'policy': GOOGLE_POLICY, 'language': 'es'})))
    assert cached_result(case[0], case[1], SPANISH) is None
    assert plan_detection(engine, ['master'])['requests'] == 1
    assert engine.records()[0]['google_language'] == ''
