from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from rta.export import result_row, summary, workbook
from rta.models import Result, AI_PENDING, MANUAL_REQUIRED
from tests.test_app import by_label, navigate
from tests.test_review_workflow import workspace, app_for
from tests.test_source_selection import suggestions
from tests.test_recommendations import setup


@pytest.mark.parametrize('proposed,expected', [
    ('', MANUAL_REQUIRED), (None, MANUAL_REQUIRED), ('  \t\n', MANUAL_REQUIRED),
    ('Enterprise', AI_PENDING), (' N/A ', AI_PENDING), ('0', AI_PENDING),
])
@pytest.mark.parametrize('raw', ['Pending', 'Verified unchanged'])
def test_workflow_statuses_are_display_only(proposed, expected, raw):
    r = Result(1, 'Account', 'Original', ai_proposed='Unrelated AI metadata', proposed=proposed,
               review_status=raw, final='Existing final')
    before = r.dict()
    assert r.status_label == expected
    assert result_row(r)['Review Status'] == expected
    assert summary([r])[expected] == 1
    assert r.dict() == before


@pytest.mark.parametrize('decision,status', [
    ('Approve', 'Approved'), ('Approve', 'Approved — AI Recommendation'),
    ('Reject', 'Rejected'), ('Reject', 'Rejected — AI Recommendation'),
    ('Manual', 'Manually Verified'), ('Pending', 'Pending'),
])
@pytest.mark.parametrize('ai', ['', 'Enterprise'])
def test_saved_decisions_take_precedence(decision, status, ai):
    r = Result(0, 'Account', 'Original', decision=decision, review_status=status, ai_proposed=ai)
    expected = {'Approve': 'Approved', 'Reject': 'Rejected', 'Manual': 'Manually Verified',
                'Pending': MANUAL_REQUIRED}[decision]
    assert r.status_label == expected
    assert summary([r])[AI_PENDING] == 0
    assert summary([r])[MANUAL_REQUIRED] == (1 if decision == 'Pending' else 0)
    # A raw outcome without a saved human decision is not an approval.
    r.decision = ''
    assert r.status_label == MANUAL_REQUIRED


@pytest.fixture
def large_review(workspace):
    store, project, token = workspace
    payload = store.load(project)
    blocked = payload['results'][2]
    records = []
    for row_id in range(125):
        r = deepcopy(blocked)
        r.update(row_id=row_id, identifier=f'Blocked {row_id:03}')
        records.append(r)
    eligible = deepcopy(payload['results'][1])
    eligible['row_id'] = 125
    records.append(eligible)
    payload['results'] = records
    store.save(project, payload)
    return store, project, token


def save_bulk(app, action='Approve'):
    by_label(app.checkbox, 'Select All filtered accounts').check().run()
    by_label(app.selectbox, 'Bulk decision').set_value(action)
    by_label(app.text_input, 'Bulk reviewer name').set_value('UX reviewer')
    by_label(app.checkbox, 'I have verified the evidence and field value for every selected account.').check()
    for checkbox in app.checkbox:
        if checkbox.label.startswith('I confirm including '):
            checkbox.check()
    by_label(app.button, 'Save Review Decisions').click().run()
    assert not app.exception


def notice(app):
    return [w for w in app.warning if w.value.startswith('Bulk review:')]


def test_large_warning_summary_pagination_clear_and_repeat(large_review, monkeypatch):
    store, project, _ = large_review
    app = app_for(large_review, monkeypatch)
    navigate(app, 'Review & Approvals')
    assert not [c for c in app.checkbox if c.label.startswith('Select row')]
    by_label(app.button, 'Show Account List').click().run()
    assert len([c for c in app.checkbox if c.label.startswith('Select row')]) == 50
    save_bulk(app)
    assert len(notice(app)) == 1
    assert 'Selected 126 · Saved 1 · Skipped 125 · Warnings 125' in notice(app)[0].value
    assert len(notice(app)) == 1
    assert by_label(app.expander, 'View Warnings')
    detail = next(d.value for d in app.dataframe if 'Account identifier' in d.value.columns)
    assert len(detail) == 50 and detail.iloc[0]['Account identifier'] == 'Blocked 000'
    grouped = next(d.value for d in app.dataframe if 'Accounts' in d.value.columns)
    assert grouped['Accounts'].sum() == 125 and len(grouped) == 1
    by_label(app.number_input, 'Warning page').set_value(3).run()
    detail = next(d.value for d in app.dataframe if 'Account identifier' in d.value.columns)
    assert len(detail) == 25 and detail.iloc[-1]['Account identifier'] == 'Blocked 124'
    assert len(notice(app)) == 1  # Survives reruns.
    before, audit = store.load(project), store.audit(project)
    assert len(audit) == 1
    by_label(app.button, 'Clear All Warnings').click().run()
    assert not notice(app) and not any(e.label == 'View Warnings' for e in app.expander)
    assert store.load(project) == before and store.audit(project) == audit
    app.run()
    assert not notice(app)
    by_label(app.text_input, 'Search accounts').set_value('Blocked').run()
    save_bulk(app)
    assert len(notice(app)) == 1 and 'Selected 125 · Saved 0 · Skipped 125 · Warnings 125' in notice(app)[0].value
    assert by_label(app.number_input, 'Warning page').value == 1
    assert store.load(project) == before and store.audit(project) == audit
    # A fresh operation replaces the warning report, rather than appending to it.
    save_bulk(app, 'Reject')
    assert not notice(app)
    assert any('Selected 125 · Saved 125 · Skipped 0 · Warnings 0' in s.value for s in app.success)
    assert not any(e.label == 'View Warnings' for e in app.expander)
    assert len(store.audit(project)) == 126
    assert all(r['decision'] for r in store.load(project)['results'])


def test_selection_survives_pagination_and_select_all_is_filtered(large_review, monkeypatch):
    app = app_for(large_review, monkeypatch)
    navigate(app, 'Review & Approvals')
    by_label(app.button, 'Show Account List').click().run()
    next(c for c in app.checkbox if c.label.startswith('Select row 1 ')).check().run()
    by_label(app.number_input, 'Selection page').set_value(3).run()
    next(c for c in app.checkbox if c.label.startswith('Select row 126 ')).check().run()
    assert any('Selected accounts: 2 of 126' in c.value for c in app.caption)
    by_label(app.number_input, 'Selection page').set_value(1).run()
    assert next(c for c in app.checkbox if c.label.startswith('Select row 1 ')).value
    by_label(app.text_input, 'Search accounts').set_value('Blocked 00').run()
    by_label(app.checkbox, 'Select All filtered accounts').check().run()
    assert any('Selected accounts: 10 of 10' in c.value for c in app.caption)
    by_label(app.selectbox, 'Bulk decision').set_value('Reject')
    by_label(app.text_input, 'Bulk reviewer name').set_value('R')
    by_label(app.button, 'Save Review Decisions').click().run()
    store, project, _ = large_review
    assert [e['row_id'] for e in store.audit(project)] == list(range(10))


def test_critical_validation_errors_are_still_visible(large_review, monkeypatch):
    app = app_for(large_review, monkeypatch)
    navigate(app, 'Review & Approvals')
    by_label(app.checkbox, 'Select All filtered accounts').check().run()
    by_label(app.button, 'Save Review Decisions').click().run()
    assert any('Enter a reviewer name' in e.value for e in app.error)
    store, project, _ = large_review
    assert not store.audit(project) and not notice(app)


def test_ai_workflow_filter_keeps_eligibility_checks_and_export_consistent(suggestions, monkeypatch):
    store, project, _ = suggestions
    payload = store.load(project)
    # Both display AI Proposed, but only the first has evidence for its own identity.
    manual = deepcopy(payload['results'][1])
    manual.update(row_id=2, proposed=' \t ', ai_proposed=' \t ', recommendation_status='')
    rejected = deepcopy(payload['results'][0])
    rejected.update(row_id=3, decision='Reject', review_status='Rejected')
    payload['results'] += [manual, rejected]
    store.save(project, payload)
    before = store.load(project)
    app = app_for(suggestions, monkeypatch)
    navigate(app, 'Review & Approvals')
    assert by_label(app.metric, AI_PENDING).value == '2'
    assert by_label(app.metric, MANUAL_REQUIRED).value == '1'
    by_label(app.selectbox, 'Review status').set_value(AI_PENDING).run()
    assert list(app.dataframe[-1].value['Review Status']) == [AI_PENDING, AI_PENDING]
    assert any('redthread-status' in m.value and AI_PENDING in m.value for m in app.markdown)
    save_bulk(app)
    assert 'Selected 2 · Saved 1 · Skipped 1 · Warnings 1' in notice(app)[0].value
    persisted = store.load(project)
    assert persisted['results'][0]['review_status'] == 'Approved — AI Recommendation'
    assert persisted['results'][1] == before['results'][1]
    assert persisted['results'][2:] == before['results'][2:]
    assert len(store.audit(project)) == 1
    by_label(app.selectbox, 'Review status').set_value(MANUAL_REQUIRED).run()
    assert len(app.dataframe[-1].value) == 1
    # Filtering changes the table, while the selected individual account remains stable.
    assert by_label(app.selectbox, 'Account to review').value == 0
    # Actual Export Results output uses exactly the same derived statuses and counts.
    import rta.export
    original, exports = rta.export.workbook, []
    def capture(*args, **kwargs):
        data = original(*args, **kwargs)
        exports.append(data)
        return data
    monkeypatch.setattr(rta.export, 'workbook', capture)
    navigate(app, 'Export Results')
    book = load_workbook(BytesIO(exports[-1]))
    sheet = book['Reconciliation Results']
    headers = [c.value for c in sheet[1]]
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert [r['Review Status'] for r in rows] == ['Approved', AI_PENDING, MANUAL_REQUIRED, 'Rejected']
    assert rows[0]['Final Value'] == rows[0]['AI Proposed Value'] == 'Enterprise'
    assert rows[1]['AI Proposed Value'] == 'Enterprise' and rows[1]['Final Value'] is None
    counts = dict(book['Summary'].iter_rows(min_row=2, values_only=True))
    assert counts[AI_PENDING] == counts[MANUAL_REQUIRED] == 1
    navigate(app, 'Run Reconciliation')
    assert list(app.dataframe[-1].value['Review Status']) == [row['Review Status'] for row in rows]
    reloaded = app_for(suggestions, monkeypatch)
    navigate(reloaded, 'Review & Approvals')
    assert by_label(reloaded.metric, AI_PENDING).value == '1'
    assert by_label(reloaded.metric, MANUAL_REQUIRED).value == '1'
