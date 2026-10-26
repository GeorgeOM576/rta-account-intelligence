from collections import Counter
import csv
import json
import hashlib
from pathlib import Path
from io import BytesIO, StringIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from .models import valid, AI_PENDING, MANUAL_REQUIRED
from .review import final_value, legacy_auto_finalized

LOADED_BUILD = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()


def translation_csv(records):
    """Explicit current-translation export; no raw detector hints or audit snapshots."""
    rows = []
    for item in records:
        if item.get('translation_excluded') or item['status'] == 'Excluded':
            continue
        rows.append({'Dataset': item['dataset'].title(), 'Row': item['row_id'] + 1, 'Field': item['field'],
            'Record Scope': 'Current selected enabled field',
            'Original text': item['original'], 'English Translation': item['english'],
            'Detected Language': item.get('detected_language', item['language']),
            'Google Detected Language': item.get('google_language', ''), 'Google Confidence': item.get('google_confidence'),
            'Language Classification': item.get('language_classification', ''),
            'Language Review Recommended': item.get('language_review_required', False),
            'Language Review Reason': item.get('language_review_reason', ''),
            'Detection Source': item.get('detection_source', ''),
            'Detection Provider': item.get('detection_provider', ''),
            'Translation Provider': item.get('provider', ''),
            'Confirmed Language': item['language'] if item.get('language_confirmed') or item['status'] == 'Approved Translation' else '',
            'Translation Status': item['status'],
            'Detection Outcome': (item.get('automatic_detection_outcome', item.get('detection_outcome', '')).replace('Confident English', 'English').replace('Confident Non-English', 'Non-English')
                                  if item.get('detection_provider') == 'google' else item.get('automatic_detection_outcome', item.get('detection_outcome', ''))),
            'Detection Policy': item.get('detection_version', ''), 'Detection Build': item.get('detection_build', ''),
            'Analyzed At': item.get('detection_analyzed_at', '')})
    output = StringIO(newline='')
    if rows:
        writer = csv.DictWriter(output, fieldnames=list(rows[0]))
        writer.writeheader()
        # CSV is often opened in Excel: keep formula-like uploads as literal data.
        for row in rows:
            writer.writerow({key: "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@', '\t', '\r')) else value
                             for key, value in row.items()})
    return output.getvalue().encode('utf-8-sig')


def summary(results):
    counts = Counter(r.match_status for r in results)
    return {'Total Master Records': len(results), 'Exact Matches': counts['Exact'],
            'Normalized Matches': counts['Normalized'], 'Alias Matches': counts['Alias'],
            'Fuzzy Candidates': counts['Fuzzy candidates'], 'Duplicate Conflicts': counts['Duplicate conflict'],
            'Duplicate Identifiers': counts['Duplicate identifier'],
            'Accounts Requiring Research': sum(r.identity_confidence != 'Reliable local match' and not r.decision for r in results),
            'Accounts Researched': sum(r.research_status in ('Completed', 'Cached', 'Insufficient Evidence') for r in results),
            'AI Recommendations': sum(r.recommendation_status == 'Proposed' for r in results),
            'Insufficient Evidence': sum(r.recommendation_status == 'Insufficient Evidence' for r in results),
            'Pending Approvals': sum(r.status_label in (AI_PENDING, MANUAL_REQUIRED) for r in results),
            AI_PENDING: sum(r.status_label == AI_PENDING for r in results),
            MANUAL_REQUIRED: sum(r.status_label == MANUAL_REQUIRED for r in results),
            'Approved Corrections': sum(r.status_label in ('Approved', 'Manually Verified') and final_value(r) != r.original for r in results),
            'Manually Verified': sum(r.status_label == 'Manually Verified' for r in results),
            'Unresolved Accounts': sum(not final_value(r) for r in results),
            'Legacy Automatically Finalized Matches': sum(legacy_auto_finalized(r) for r in results)}


def translation_columns(r):
    """Shared leading columns. A suggested candidate is not a matched Source."""
    master = r.account_context.get('master', {})
    source = r.account_context.get('matched_source', {})
    return {'Master Identifier': r.identifier,
            'Master Approved English Translation': master.get('approved_english', ''),
            'Master Translation Status': master.get('translation_status', ''),
            'Matched Source Approved English Translation': source.get('approved_english', ''),
            'Matched Source Translation Status': source.get('translation_status', ''),
            'Match Basis': r.match_basis,
            'Master Original Name': master.get('original_name', r.identifier),
            'Matched Source Original Name': source.get('original_name', '')}


def result_row(r):
    # Enforce the final-value invariant at the export boundary as well as in the UI.
    final = final_value(r)
    master = r.account_context.get('master', {})
    source = r.account_context.get('matched_source', {})
    row = {**translation_columns(r), 'Original Value': r.original, 'Proposed Value': r.proposed,
           'Final Value': final, 'Match Status': r.match_status, 'Review Status': r.status_label}
    for i in range(3):
        c = r.candidates[i] if i < len(r.candidates) else {}
        row.update({f'Suggested Account {i+1}': c.get('account', ''),
                    f'Similarity Score {i+1}': c.get('similarity'), f'Suggested Value {i+1}': c.get('value', ''),
                    f'Suggested Match Basis {i+1}': c.get('match_basis', '')})
    row.update({'Relationship Type': r.relationship, 'Research Status': r.research_status,
                'Research Explanation': r.explanation, 'Evidence URL 1': r.evidence[0] if r.evidence else '',
                'Evidence URL 2': r.evidence[1] if len(r.evidence) > 1 else '', 'Reviewer Decision': r.decision,
                'Reviewer Notes': r.notes, 'Identity Confidence': r.identity_confidence,
                'Reviewer': r.reviewer, 'Reviewed At': r.reviewed_at, 'Value Origin': r.value_origin,
                'Record ID': r.row_id, 'Master Account': r.identifier,
                'AI Recommended Source Account': r.recommended_source_account,
                'Parent Company': r.parent_company, 'AI Proposed Value': r.ai_proposed,
                'AI Confidence': r.ai_confidence, 'AI Recommendation Status': r.recommendation_status,
                'Recommendation Reasoning': r.recommendation_reasoning, 'Research Timestamp': r.researched_at,
                'Verified Facts': json.dumps(r.verified_facts, ensure_ascii=False),
                'Suggested Matches': json.dumps(r.suggested_matches, ensure_ascii=False),
                'Proposed Segment': r.proposed, 'Recommendation Decision': r.recommendation_status,
                'Organization Mapping Policy': r.mapping_policy})
    row.update({'Selected Source Record ID': r.selected_source.get('source_row'),
                'Selected Source Account': r.selected_source.get('account', ''),
                'Selected Source Comparison Value': r.selected_source.get('value', ''),
                'Selected Source Similarity': r.selected_source.get('similarity'),
                'Selected Source Dataset': r.selected_source.get('dataset', '')})
    for i, url in enumerate(r.evidence[2:], 3):
        row[f'Evidence URL {i}'] = url
    row['Intelligence Proposed Value'] = r.proposed
    row['Stored Review Status'] = r.review_status
    row['Legacy Automatic Final Requires Review'] = legacy_auto_finalized(r)
    row['Translation Match Provenance'] = json.dumps(r.translation_provenance, ensure_ascii=False)
    row.update({'Original Account Name': master.get('original_name', r.identifier),
        'Approved English Translation': master.get('approved_english', ''),
        'Matched Account Name': source.get('original_name', ''),
        'Matched Account Approved English Translation': source.get('approved_english', ''),
        'Master Original Name': master.get('original_name', r.identifier),
        'Master Approved English Translation': master.get('approved_english', ''),
        'Source Original Name': source.get('original_name', ''),
        'Source Approved English Translation': source.get('approved_english', ''),
        'Match Method': r.match_method or r.match_status,
        'Match Basis': r.match_basis, 'Match Score': r.match_score,
        'Secondary Validation Status': r.secondary_validation.get('status', ''),
        'Secondary Validation Confidence': r.secondary_validation.get('confidence', ''),
        'Secondary Validation Reason': r.secondary_validation.get('reason', ''),
        'Secondary Validation Evidence': json.dumps(r.secondary_validation.get('evidence', []), ensure_ascii=False),
        'Local Match Confidence': r.local_confidence, 'Local Recommendation Reason': r.local_reason,
        'Research Recommendation': r.recommendation_reasoning,
        'Research Evidence / Sources': '\n'.join(r.evidence)})
    return row


def workbook(results, audit, conflicts, source_count, translations=None, datasets=None, *, current_translations=None):
    context_records = current_translations if current_translations is not None else translations
    if datasets and context_records is not None and datasets.get('mapping'):
        from .approved_names import ApprovedNames
        from .models import Mapping
        ApprovedNames(datasets, Mapping(**datasets['mapping']), context_records).attach(results)
    if translations is not None:
        translations = [item for item in translations if not item.get('translation_excluded') and item['status'] != 'Excluded']
    book = Workbook()
    book.remove(book.active)
    tables = [('Reconciliation Results', [result_row(r) for r in results]),
              ('Summary', [{'Metric': k, 'Count': v} for k, v in
                           {**summary(results), 'Total Source Records': source_count}.items()]),
              ('Audit Trail', audit), ('Source Conflicts', conflicts)]
    if translations is not None:
        from .translation import translation_metrics, translated_rows
        tables.append(('Translation Summary', translation_metrics(translations)))
        tables.append(('Translation Details', [{**item,
            'Original Company Name': item['original'] if item.get('field_kind') == 'account_name' else '',
            'Detected Language': item.get('detected_language', item['language']),
            'English Translation': item['english'], 'Translation Status': item['status'],
            'Translation Provider': item['provider']} for item in translations]))
        if datasets:
            for dataset in ('master', 'source'):
                tables.append((f'{dataset.title()} Translations', translated_rows(datasets, dataset, translations)))
    for title, records in tables:
        sheet = book.create_sheet(title)
        headers = list(dict.fromkeys(k for r in records for k in r)) or ['No records']
        sheet.append(headers)
        for record in records:
            sheet.append([str(record.get(h, '')) if isinstance(record.get(h), (list, dict))
                          else record.get(h, '') for h in headers])
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='17324D')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 32
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                # Force uploaded formula-like content to remain literal strings.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if headers[cell.column-1].startswith('Evidence URL') and str(cell.value).startswith(('https://', 'http://')):
                    cell.hyperlink = cell.value
                    cell.font = Font(color='146C94', underline='single')
        for i, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(i)].width = 48 if 'Explanation' in header or 'Notes' in header else 25
        if 'Review Status' in headers and sheet.max_row > 1:
            col = get_column_letter(headers.index('Review Status') + 1)
            for status, color in [(AI_PENDING, 'FFF0C2'), (MANUAL_REQUIRED, 'FFF0C2'),
                                  ('Approved', 'D7F0E2'), ('Rejected', 'FBE0E0'),
                                  ('Manually Verified', 'DCEBFF')]:
                sheet.conditional_formatting.add(f'{col}2:{col}{sheet.max_row}',
                    FormulaRule(formula=[f'{col}2="{status}"'], fill=PatternFill('solid', fgColor=color)))
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
