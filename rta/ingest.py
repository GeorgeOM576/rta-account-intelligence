from io import BytesIO
from pathlib import Path
import polars as pl
from openpyxl import load_workbook


def parse_file(data: bytes, filename: str) -> pl.DataFrame:
    if not data:
        raise ValueError('The uploaded file is empty.')
    try:
        suffix = Path(filename).suffix.lower()
        if suffix == '.csv':
            frame = pl.read_csv(BytesIO(data), infer_schema=False, try_parse_dates=False)
        elif suffix == '.xlsx':
            book = load_workbook(BytesIO(data), read_only=True, data_only=True)
            sheet = book.active
            rows = sheet.iter_rows()
            headers = [str(c.value or '').strip() for c in next(rows)]
            if any(not x for x in headers) or len(set(headers)) != len(headers):
                raise ValueError('Column headers must be nonempty and unique.')
            def text(cell):
                if cell.value is None:
                    return None
                if isinstance(cell.value, int) and cell.number_format and set(cell.number_format) == {'0'}:
                    return str(cell.value).zfill(len(cell.number_format))
                return str(cell.value)
            values = [[text(c) for c in row] for row in rows]
            frame = pl.DataFrame(values, schema={h: pl.String for h in headers}, orient='row')
            book.close()
        else:
            raise ValueError('Choose a CSV or XLSX file.')
        if frame.height == 0 or frame.width == 0:
            raise ValueError('The dataset has no records.')
        return frame
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Cannot read this file. Check its format, encoding, and column headers.') from exc


def identifier_stats(frame, column):
    # Blank columns may have Null dtype, including frames reconstructed from rows.
    values = pl.col(column).cast(pl.Utf8).fill_null('').str.strip_chars()
    return frame.select(
        (values == '').sum().alias('Missing identifiers'),
        values.filter(values != '').is_duplicated().sum().alias('Rows with duplicate identifiers'),
    ).row(0, named=True)
