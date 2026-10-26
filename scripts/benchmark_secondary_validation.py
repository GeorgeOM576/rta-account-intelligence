"""Offline challenge scoring. Answer keys are used only after reconciliation completes."""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import polars as pl
from openpyxl import load_workbook
from rta.approved_names import ApprovedNames
from rta.export import workbook
from rta.models import Mapping, Result
from rta.secondary_validation import SecondaryValidator
from rta.storage import digest
from rta.translation_matching import reconcile


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream))


def score(results, key, mapping):
    metrics = dict(correct_first_candidates=0, correct_accepted_matches=0, proposals=0,
                   correct_proposals=0, false_positive_matches=0, incorrect_proposals=0,
                   unresolved=0, unmatched_left_unmatched=0, unresolved_true_matches=0)
    for result, expected in zip(results, key):
        account, value = expected['Expected Source Account'] or '', expected['Expected Segment'] or ''
        if result.candidates and account and result.candidates[0]['account'] == account:
            metrics['correct_first_candidates'] += 1
        linked = result.account_context.get('matched_source', {}).get('fields', {}).get(mapping.source_id, '')
        if linked:
            metrics['correct_accepted_matches' if linked == account else 'false_positive_matches'] += 1
        else:
            metrics['unresolved'] += 1
            metrics['unresolved_true_matches' if account else 'unmatched_left_unmatched'] += 1
        if result.proposed:
            metrics['proposals'] += 1
            correct = bool(account and linked == account and result.proposed == value)
            metrics['correct_proposals' if correct else 'incorrect_proposals'] += 1
    matchable = sum(bool(k['Expected Source Account']) for k in key)
    metrics['proposal_coverage_all_records'] = metrics['proposals'] / len(key)
    metrics['proposal_coverage_matchable_records'] = metrics['proposals'] / matchable
    metrics['proposal_precision'] = metrics['correct_proposals'] / metrics['proposals'] if metrics['proposals'] else None
    metrics['pending_human_decisions'] = sum(not r.decision for r in results)
    metrics['final_values'] = sum(bool(r.final) for r in results)
    return metrics


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for arg in ('master', 'source', 'baseline-export', 'answer-key', 'before-results', 'before-raw-results', 'output'):
        parser.add_argument('--' + arg, type=Path, required=True)
    args = parser.parse_args()
    master, source, saved = (read_csv(path) for path in (args.master, args.source, args.baseline_export))
    assert len(master) == len(saved)
    assert all(m['Account Name'] == b['Master Identifier'] and m['Segment'] == b['Master Original Value']
               for m, b in zip(master, saved))
    mapping = Mapping('Account Name', 'Account Name', 'Segment', 'Segment')
    # Reuse only fields explicitly labeled approved in the saved export; no new approval is inferred.
    records = [dict(dataset='master', row_id=i, field='Account Name', original=m['Account Name'],
                    revision=digest(master), english=b['Master Approved English Translation'], status='Approved Translation')
               for i, (m, b) in enumerate(zip(master, saved)) if b['Master Approved English Translation']]
    frame_m, frame_s = pl.DataFrame(master), pl.DataFrame(source)
    after, conflicts = reconcile(frame_m, frame_s, mapping, records)
    names = ApprovedNames({'master': master, 'source': source}, mapping, records)
    validator = SecondaryValidator(master, source, mapping, names)
    # CSV summaries/URLs are retained as context, without fabricating quoted research snapshots.
    for r, b in zip(after, saved):
        r.research_status = b['Research Status']
        r.recommendation_status = b['AI Recommendation Status']
        r.explanation = b['Research Explanation']
        r.ai_confidence = b['AI Confidence']
        r.evidence = [b[k] for k in ('Evidence URL 1', 'Evidence URL 2') if b[k]]
        validator.evaluate(r)
    names.attach(after)
    raw, _ = reconcile(frame_m, frame_s, mapping)
    before = [Result(**r) for r in json.loads(args.before_results.read_text())]
    before_raw = [Result(**r) for r in json.loads(args.before_raw_results.read_text())]
    # Do not load expected identities until all engine calls above have finished.
    reference = load_workbook(args.answer_key, data_only=True)
    values = list(reference['Answer Key - Do Not Upload'].values)
    key = [dict(zip(values[0], row)) for row in values[1:]]
    assert len(master) == len(source) == len(key) == len(before) == len(before_raw)
    assert len({r['Account Name'] for r in source}) == len(source)
    source_lookup = {r['Account Name']: r['Segment'] for r in source}
    assert all(k['Master Account'] == m['Account Name'] for k, m in zip(key, master))
    assert all(not k['Expected Source Account'] or source_lookup[k['Expected Source Account']] == k['Expected Segment'] for k in key)
    for sheet, uploaded in [('Master', master), ('Source', source)]:
        values = list(reference[sheet].values)
        expected_rows = [{field: str(value or '') for field, value in zip(values[0], row)} for row in values[1:]]
        assert expected_rows == uploaded
    provided_before = []
    source_rows = {row['Account Name']: i for i, row in enumerate(source)}
    for i, row in enumerate(saved):
        result = Result(i, row['Master Identifier'], row['Master Original Value'],
                        proposed=row['Proposed Value'], final=row['Final Value'],
                        decision=row['Reviewer Decision'], match_status=row['Match Status'])
        result.candidates = [{'source_row': source_rows[row[f'Suggested Account {n}']],
                              'account': row[f'Suggested Account {n}'],
                              'similarity': float(row[f'Similarity Score {n}'])}
                             for n in range(1, 4) if row[f'Suggested Account {n}']]
        matched = row['Matched Source Original Name']
        if matched:
            assert matched in source_rows
            result.account_context = {'matched_source': {'fields': source[source_rows[matched]], 'original_name': matched}}
        provided_before.append(result)
    assert all(not r.final and not r.decision for r in after + raw)
    args.output.mkdir(parents=True, exist_ok=True)
    paths = [args.master, args.source, args.baseline_export, args.answer_key, args.before_results, args.before_raw_results]
    summary = {'before': score(provided_before, key, mapping), 'local_before': score(before, key, mapping),
               'after': score(after, key, mapping),
               'raw_before': score(before_raw, key, mapping), 'raw_after': score(raw, key, mapping),
               'approved_translation_inputs': len(records),
               'sha256': {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}}
    (args.output / 'summary.json').write_text(json.dumps(summary, indent=2))
    (args.output / 'after-results.json').write_text(json.dumps([r.dict() for r in after], ensure_ascii=False, indent=2))
    rows = []
    for old, local_old, new, expected in zip(provided_before, before, after, key):
        old_link = old.account_context.get('matched_source', {}).get('original_name', '')
        new_link = new.account_context.get('matched_source', {}).get('original_name', '')
        rows.append({'row': new.row_id + 1, 'master': new.identifier,
                     'expected_source': expected['Expected Source Account'] or '', 'expected_segment': expected['Expected Segment'] or '',
                     'before_first': old.candidates[0]['account'] if old.candidates else '',
                     'local_before_first': local_old.candidates[0]['account'] if local_old.candidates else '',
                     'after_first': new.candidates[0]['account'] if new.candidates else '',
                     'before_match': old_link, 'after_match': new_link,
                     'before_proposal': old.proposed, 'after_proposal': new.proposed,
                     'validation_status': new.secondary_validation.get('status', ''),
                     'validation_confidence': new.secondary_validation.get('confidence', ''),
                     'reason': new.secondary_validation.get('reason', ''),
                     'evidence': json.dumps(new.secondary_validation.get('evidence', []), ensure_ascii=False)})
    with (args.output / 'comparison.csv').open('w', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(rows)
    output_book = workbook(after, [], conflicts, len(source), datasets={'master': master, 'source': source, 'mapping': vars(mapping)}, current_translations=records)
    (args.output / 'updated-results.xlsx').write_bytes(output_book)
    # Verify persisted Excel approval boundaries, not only in-memory result objects.
    from io import BytesIO
    sheet = load_workbook(BytesIO(output_book))['Reconciliation Results']
    headers = [c.value for c in sheet[1]]
    exported = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert len(exported) == len(master)
    assert sum(bool(r['Proposed Value']) for r in exported) == summary['after']['proposals']
    assert all(not r['Final Value'] and not r['Reviewer Decision'] for r in exported)
    print(json.dumps({k: v for k, v in summary.items() if k != 'sha256'}, indent=2))


if __name__ == '__main__':
    main()
